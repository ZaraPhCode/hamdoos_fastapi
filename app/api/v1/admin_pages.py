"""Admin page routes — renders Jinja2 templates for the admin panel."""

from __future__ import annotations

import os
import uuid
from datetime import datetime, timezone
from typing import Optional
from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, Request, UploadFile
from fastapi.responses import HTMLResponse, JSONResponse, RedirectResponse
from fastapi.templating import Jinja2Templates
from sqlalchemy import select, func
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.database import get_db
from app.core.dependencies import get_admin_user, require_any_role
from app.models.identity import User, Role, IdentityInformation, RoleClaim, UserRole
from app.models.product import Product, Category, Brand, ProductType, ProductUnit, Currency, Tag, CategoryOption, PriceHistory, ProductTag, RelatedProduct, SimilarProduct, ProductImage, MenuDatasheet, Variety, ProductVariety
from app.models.product_features import TechnicalFeature, TechnicalTable, CategoryTechnicalFeature, TechnicalTableProduct, TechnicalFeatureEnum, TechnicalFeatureValue
from app.models.order import OrderModel as Order, PayMethod, PostType, Discount, OrderProduct, OrderStatusRecord
from app.models.invoice import Invoice, InvoiceProduct, Supplier, SupplierProduct, PurchaseOrder
from app.models.finance import Receipt, Wallet, CurrencyDetail, WarehouseMovement, PaymentRequest
from app.models.customer_content import Comment, Media, NotifiedProduct
from app.models.support import Ticket, Chat
from app.models.common import Log, AdminParameter, SmsCode, MobileNumber, BankInfo, SiteSetting
from app.models.manufacturer import Manufacturer
from app.models.log_enums import (
    LOG_TABLE_ORDER as TABLE_OPTIONS,
    LOG_TYPE_INT as LOG_TYPE_OPTIONS,
    LOG_TYPE_NAME as LOG_TYPE_NAME_MAP,
    LOG_TABLE_NAME as LOG_TABLE_NAME_MAP,
    resolve_table_int,
    resolve_type_int,
)
from app.schemas.product import CategoryCreate, CategoryUpdate
from app.utils.common_works import generate_slug
from app.utils.persian_tools import to_farsi, to_farsi_full, from_farsi_date, is_phone_number, is_email
from app.services import admin_service, product_service, order_service, invoice_service, warehouse_service, finance_service, support_service, identity_service
from app.config.site_config import register_template_globals

templates = register_template_globals(Jinja2Templates(directory="app/templates"))
router = APIRouter(prefix="/administration", tags=["Admin Pages"])


# ── Admin Login ──
# The platform has a single login page at /login. Visiting /administration/login
# redirects there; the same JWT cookie grants access to store + admin routes.

@router.get("/login", response_class=HTMLResponse)
async def admin_login_page(request: Request):
    from fastapi.responses import RedirectResponse
    from urllib.parse import quote
    next_url = request.query_params.get("next", "")
    target = f"/login?next={quote(next_url)}" if next_url else "/login"
    return RedirectResponse(url=target, status_code=302)


@router.post("/login", response_class=HTMLResponse)
async def admin_login_submit(request: Request):
    from fastapi.responses import RedirectResponse
    next_url = (await request.form()).get("next", "")
    target = f"/login?next={next_url}" if next_url else "/login"
    return RedirectResponse(url=target, status_code=302)


# ── Dashboard ──

@router.get("", response_class=HTMLResponse)
@router.get("/dashboard", response_class=HTMLResponse)
async def admin_dashboard(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Product Manager", "Orders Manager", "Financial Manager", "Warehouse Keeper")),
    db: AsyncSession = Depends(get_db),
):
    stats = await admin_service.get_dashboard_stats(db)
    order_distribution = await admin_service.get_order_status_distribution(db)
    recent_orders = await admin_service.get_recent_orders(db, 10)
    low_stock = await admin_service.get_low_stock_products(db, 5, 10)
    return templates.TemplateResponse(
        "admin/dashboard.html",
        {
            "request": request,
            "current_user": current_user,
            "stats": stats,
            "order_distribution": order_distribution,
            "recent_orders": recent_orders,
            "low_stock_products": low_stock,
        },
    )


# ── Products ──

@router.get("/products", response_class=HTMLResponse)
async def admin_products(
    request: Request,
    page: int = Query(1),
    q: str = Query(None),
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    from app.schemas.product import ProductSearchParams
    params = ProductSearchParams(query=q, page=page, page_size=20)
    products, total = await product_service.search_products(db, params)
    # Load admin users for filter dropdown
    users_result = await db.execute(
        select(User).where(User.is_removed == False).order_by(User.first_name)
    )
    users = users_result.scalars().all()
    tree = await product_service.get_admin_category_tree(db)
    options = _flatten_category_options(tree)
    return templates.TemplateResponse("admin/products.html", {
        "request": request,
        "current_user": current_user,
        "products": [product_service._build_product_list_response(p) for p in products],
        "total": total, "page": page, "total_pages": (total + 19) // 20, "query": q,
        "users": users, "categories": options,
    })


@router.get("/products/new", response_class=HTMLResponse)
async def admin_product_create(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    tree = await product_service.get_admin_category_tree(db)
    options = _flatten_category_options(tree)
    brands = await product_service.get_all_brands(db)
    types = (await db.execute(select(ProductType).where(ProductType.is_removed == False))).scalars().all()
    units = (await db.execute(select(ProductUnit).where(ProductUnit.is_removed == False))).scalars().all()
    currencies = (await db.execute(select(Currency).where(Currency.is_removed == False))).scalars().all()
    suppliers = (await db.execute(select(Supplier).where(Supplier.is_removed == False))).scalars().all()
    return templates.TemplateResponse("admin/product_form.html", {
        "request": request, "current_user": current_user,
        "product": None, "categories_tree": options,
        "brands": brands, "product_types": types,
        "product_units": units, "currencies": currencies, "suppliers": suppliers,
    })


@router.post("/products/new", response_class=HTMLResponse)
async def admin_product_create_submit(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    form = await request.form()
    from app.schemas.product import ProductCreate
    from app.services.product_service import create_product
    import jdatetime

    purchase_date_str = form.get("purchase_date") or ""
    purchase_date = None
    if purchase_date_str:
        try:
            parts = purchase_date_str.replace("/", "-").split("-")
            purchase_date = jdatetime.date(int(parts[0]), int(parts[1]), int(parts[2])).togregorian()
        except:
            pass

    data = ProductCreate(
        name=(form.get("name") or "").strip(),
        en_name=(form.get("en_name") or "").strip() or None,
        slug=(form.get("slug") or "").strip() or None,
        en_slug=(form.get("en_slug") or "").strip() or None,
        part_number=(form.get("part_number") or "").strip() or None,
        model=(form.get("model") or "").strip() or None,
        short_description=form.get("short_description") or None,
        introduction=form.get("introduction") or None,
        keywords=(form.get("keywords") or "").strip() or None,
        meta_description=form.get("meta_description") or None,
        status=form.get("status") or "OutOfStock",
        type=form.get("type") or "Product",
        delivery_day=int(form.get("delivery_day") or 0) or None,
        tax_unique_id=(form.get("tax_unique_id") or "").strip() or None,
        taobao_choice_id=(form.get("taobao_choice_id") or "").strip() or None,
        vat_rate=float(form.get("vat_rate") or 0) or None,
        profit_rate=float(form.get("profit_rate") or 0) or None,
        points_from_purchases=int(form.get("points_from_purchases") or 0),
        max_number_of_purchases=int(form.get("max_number_of_purchases") or 0) or None,
        default_variation=(form.get("default_variation") or "").strip() or None,
        purchase_date=purchase_date,
        price=float(form.get("price") or 0) or None,
        stock_quantity=int(form.get("stock_quantity") or 0),
        minimum_purchase=int(form.get("minimum_purchase") or 1),
        category_id=uuid.UUID(form.get("category_id")) if form.get("category_id") else None,
        brand_id=uuid.UUID(form.get("brand_id")) if form.get("brand_id") else None,
        product_type_id=uuid.UUID(form.get("product_type_id")) if form.get("product_type_id") else None,
        product_unit_id=uuid.UUID(form.get("product_unit_id")) if form.get("product_unit_id") else None,
        currency_id=uuid.UUID(form.get("currency_id")) if form.get("currency_id") else None,
        is_new=form.get("is_new") in ("1", "true", "on"),
        is_special=form.get("is_special") in ("1", "true", "on"),
        on_sale=form.get("on_sale") in ("1", "true", "on"),
        suggested=form.get("suggested") in ("1", "true", "on"),
        no_display=form.get("no_display") in ("1", "true", "on"),
        is_bundle=form.get("is_bundle") in ("1", "true", "on"),
        is_calibrated=form.get("is_calibrated") in ("1", "true", "on"),
        restocked=form.get("restocked") in ("1", "true", "on"),
    )
    if not data.slug and data.name:
        from app.utils.common_works import generate_slug
        data.slug = generate_slug(data.name)
    if not data.en_slug and data.en_name:
        docstring = """<implementation of generate_slug>"""
        from app.utils.common_works import generate_slug
        data.en_slug = generate_slug(data.en_name)

    product = await create_product(db, data, current_user.id)
    log_text = form.get("log") or f"ایجاد محصول: {product.name}"
    db.add(Log(
        record_id=product.id, table_name="products",
        description=log_text, created_by_user_id=current_user.id,
        type="Create",
    ))
    await db.commit()
    return RedirectResponse(url="/administration/products", status_code=303)


@router.get("/products/{product_id}/edit", response_class=HTMLResponse)
async def admin_product_edit(
    request: Request,
    product_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    import uuid
    import jdatetime
    pid = uuid.UUID(product_id)
    product = await product_service.get_product_by_id(db, pid)
    if not product:
        raise HTTPException(status_code=404, detail="Product not found")
    tree = await product_service.get_admin_category_tree(db)
    options = _flatten_category_options(tree)
    brands = await product_service.get_all_brands(db)
    types = (await db.execute(select(ProductType).where(ProductType.is_removed == False))).scalars().all()
    units = (await db.execute(select(ProductUnit).where(ProductUnit.is_removed == False))).scalars().all()
    currencies = (await db.execute(select(Currency).where(Currency.is_removed == False))).scalars().all()
    suppliers = (await db.execute(select(Supplier).where(Supplier.is_removed == False))).scalars().all()
    purchase_date_str = ""
    if product.purchase_date:
        try:
            dt = product.purchase_date
            if dt.tzinfo is not None:
                dt = dt.replace(tzinfo=None)
            purchase_date_str = jdatetime.datetime.fromgregorian(datetime=dt).strftime("%Y/%m/%d")
        except Exception:
            purchase_date_str = product.purchase_date.strftime("%Y/%m/%d")
    return templates.TemplateResponse("admin/product_form.html", {
        "request": request, "current_user": current_user,
        "product": product, "categories_tree": options,
        "brands": brands, "product_types": types,
        "product_units": units, "currencies": currencies, "suppliers": suppliers,
        "purchase_date_str": purchase_date_str,
    })


@router.post("/products/{product_id}/edit", response_class=HTMLResponse)
async def admin_product_edit_submit(
    request: Request,
    product_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    import uuid
    import jdatetime
    from app.schemas.product import ProductUpdate
    from app.services.product_service import update_product

    pid = uuid.UUID(product_id)
    product = await product_service.get_product_by_id(db, pid)
    if not product:
        raise HTTPException(status_code=404, detail="Product not found")

    form = await request.form()

    purchase_date_str = form.get("purchase_date") or ""
    purchase_date = product.purchase_date
    if purchase_date_str:
        try:
            parts = purchase_date_str.replace("/", "-").split("-")
            purchase_date = jdatetime.date(int(parts[0]), int(parts[1]), int(parts[2])).togregorian()
        except Exception:
            purchase_date = product.purchase_date

    data = ProductUpdate(
        name=(form.get("name") or "").strip(),
        en_name=(form.get("en_name") or "").strip() or None,
        slug=(form.get("slug") or "").strip() or None,
        en_slug=(form.get("en_slug") or "").strip() or None,
        part_number=(form.get("part_number") or "").strip() or None,
        model=(form.get("model") or "").strip() or None,
        short_description=form.get("short_description") or None,
        introduction=form.get("introduction") or None,
        keywords=(form.get("keywords") or "").strip() or None,
        meta_description=form.get("meta_description") or None,
        status=form.get("status") or "OutOfStock",
        type=form.get("type") or "Product",
        delivery_day=int(form.get("delivery_day") or 0) or None,
        tax_unique_id=(form.get("tax_unique_id") or "").strip() or None,
        taobao_choice_id=(form.get("taobao_choice_id") or "").strip() or None,
        vat_rate=float(form.get("vat_rate") or 0) or None,
        profit_rate=float(form.get("profit_rate") or 0) or None,
        points_from_purchases=int(form.get("points_from_purchases") or 0),
        max_number_of_purchases=int(form.get("max_number_of_purchases") or 0) or None,
        default_variation=(form.get("default_variation") or "").strip() or None,
        purchase_date=purchase_date,
        price=float(form.get("price") or 0) or None,
        stock_quantity=int(form.get("stock_quantity") or 0),
        minimum_purchase=int(form.get("minimum_purchase") or 1),
        category_id=uuid.UUID(form.get("category_id")) if form.get("category_id") else None,
        brand_id=uuid.UUID(form.get("brand_id")) if form.get("brand_id") else None,
        product_type_id=uuid.UUID(form.get("product_type_id")) if form.get("product_type_id") else None,
        product_unit_id=uuid.UUID(form.get("product_unit_id")) if form.get("product_unit_id") else None,
        currency_id=uuid.UUID(form.get("currency_id")) if form.get("currency_id") else None,
        is_new=form.get("is_new") in ("1", "true", "on"),
        is_special=form.get("is_special") in ("1", "true", "on"),
        on_sale=form.get("on_sale") in ("1", "true", "on"),
        suggested=form.get("suggested") in ("1", "true", "on"),
        no_display=form.get("no_display") in ("1", "true", "on"),
        is_bundle=form.get("is_bundle") in ("1", "true", "on"),
        is_calibrated=form.get("is_calibrated") in ("1", "true", "on"),
        restocked=form.get("restocked") in ("1", "true", "on"),
    )
    if not data.slug and data.name:
        from app.utils.common_works import generate_slug
        data.slug = generate_slug(data.name)
    if not data.en_slug and data.en_name:
        from app.utils.common_works import generate_slug
        data.en_slug = generate_slug(data.en_name)

    product = await update_product(db, product, data)
    log_text = form.get("log") or f"ویرایش محصول: {product.name}"
    db.add(Log(
        record_id=product.id, table_name="products",
        description=log_text, created_by_user_id=current_user.id,
        type="Update",
    ))
    await db.commit()
    return RedirectResponse(url="/administration/products", status_code=303)


def _normalize_media_url(url):
    if not url:
        return url
    if url.startswith(("http://", "https://", "//", "/static/", "/media/")):
        return url
    normalized = url.replace("\\", "/").lstrip("/")
    if normalized.lower().startswith("media/"):
        normalized = normalized[len("media/"):]
    return "/media/" + normalized


templates.env.filters["media_url"] = _normalize_media_url


@router.get("/products/{product_id}/details", response_class=HTMLResponse)
async def admin_product_details(
    request: Request,
    product_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    import uuid
    pid = uuid.UUID(product_id)
    product = await product_service.get_product_full_details(db, pid)
    if not product:
        raise HTTPException(status_code=404, detail="Product not found")

    tech_feature_ids = {
        value.technical_feature_id
        for table_product in product.technical_table_products
        for value in table_product.technical_feature_values
    }

    general_features = []
    if product.category_id:
        ctf_stmt = (
            select(CategoryTechnicalFeature)
            .options(selectinload(CategoryTechnicalFeature.technical_feature))
            .where(
                CategoryTechnicalFeature.category_id == product.category_id,
                CategoryTechnicalFeature.is_removed == False,
            )
            .order_by(CategoryTechnicalFeature.insert_date)
        )
        ctf_rows = (await db.execute(ctf_stmt)).scalars().all()
        for ctf in ctf_rows:
            if ctf.technical_feature_id in tech_feature_ids:
                continue
            general_features.append({
                "id": str(ctf.technical_feature_id),
                "text": ctf.technical_feature.fa_name or ctf.technical_feature.name or "",
            })

    special_features = []
    sp_stmt = (
        select(TechnicalFeature)
        .where(TechnicalFeature.is_removed == False)
        .order_by(TechnicalFeature.fa_name)
    )
    sp_rows = (await db.execute(sp_stmt)).scalars().all()
    for feat in sp_rows:
        if feat.id in tech_feature_ids:
            continue
        special_features.append({
            "id": str(feat.id),
            "text": feat.fa_name or feat.name or "",
        })

    tecnical_tables = []
    for table_product in product.technical_table_products:
        if table_product.technical_table:
            tecnical_tables.append({
                "id": str(table_product.id),
                "title": table_product.technical_table.title or "",
            })

    suppliers = (await db.execute(select(Supplier).where(Supplier.is_removed == False).order_by(Supplier.intermediary_name))).scalars().all()

    technical_products = []
    for table_product in product.technical_table_products:
        headers = []
        if table_product.technical_table and table_product.technical_table.header:
            headers = [h for h in table_product.technical_table.header.split(";") if h.strip()]
        column_count = table_product.technical_table.columns if table_product.technical_table else len(headers)
        rows = []
        for value in table_product.technical_feature_values:
            fmt_headers = []
            if value.technical_feature and value.technical_feature.display_format:
                fmt_headers = value.technical_feature.display_format.split(";")
            if column_count and len(fmt_headers) < column_count:
                fmt_headers.extend([""] * (column_count - len(fmt_headers)))
            elif column_count and len(fmt_headers) > column_count:
                fmt_headers = fmt_headers[:column_count]
            rows.append({
                "value": value,
                "feature": value.technical_feature,
                "value_headers": fmt_headers,
                "cells": [_format_tfv(value, f) for f in fmt_headers],
            })
        technical_products.append({
            "table_product": table_product,
            "technical_table": table_product.technical_table,
            "headers": headers,
            "rows": rows,
        })

    def _normalize_img(img):
        from copy import copy
        c = copy(img)
        c.medium_image_url = _normalize_media_url(img.medium_image_url)
        c.small_image_url = _normalize_media_url(img.small_image_url)
        c.large_image_url = _normalize_media_url(img.large_image_url)
        return c

    return templates.TemplateResponse("admin/product_details.html", {
        "request": request, "current_user": current_user,
        "product": product,
        "general_features": general_features,
        "special_features": special_features,
        "tecnical_tables": tecnical_tables,
        "technical_products": technical_products,
        "suppliers": suppliers,
        "insert_date_fa": _to_fa_datetime(product.insert_date),
        "update_date_fa": _to_fa_datetime(product.update_date),
        "purchase_date_fa": _to_fa_date(product.purchase_date),
        "product_images": [_normalize_img(i) for i in sorted(product.product_images, key=lambda x: x.picture_order or 0)],
    })


# ── Product sub-resources (details page) ──

@router.post("/products/{product_id}/images/new", response_class=HTMLResponse)
async def admin_product_image_create(
    request: Request,
    product_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    import uuid
    pid = uuid.UUID(product_id)
    product = await product_service.get_product_by_id(db, pid)
    if not product:
        raise HTTPException(status_code=404, detail="Product not found")
    form = await request.form()
    image = ProductImage(
        id=uuid.uuid4(),
        product_id=pid,
        title=form.get("title") or None,
        medium_image_url=form.get("medium_image_url") or None,
        small_image_url=form.get("medium_image_url") or None,
        large_image_url=form.get("medium_image_url") or None,
        display_photo=form.get("display_photo") in ("1", "true", "on"),
        picture_order=int(form.get("picture_order") or 0),
        created_by_user_id=current_user.id,
        insert_date=datetime.now(timezone.utc),
        update_date=datetime.now(timezone.utc),
    )
    db.add(image)
    # Also update the product-level thumbnails so product cards everywhere show the image
    url = image.medium_image_url
    if url and (product.medium_image_url is None or "placehold" in str(product.medium_image_url).lower()):
        product.medium_image_url = url
        product.large_image_url = image.large_image_url or url
        product.feature_image_url = url
    db.add(Log(
        record_id=image.id, table_name="product_images",
        description=f"افزودن عکس به محصول: {product.name}",
        created_by_user_id=current_user.id, type="Create",
    ))
    await db.commit()
    return RedirectResponse(url=f"/administration/products/{product_id}/details", status_code=303)


@router.post("/products/{product_id}/images/{image_id}/delete", response_class=HTMLResponse)
async def admin_product_image_delete(
    request: Request,
    product_id: str,
    image_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    import uuid
    image = (await db.execute(
        select(ProductImage).where(ProductImage.id == uuid.UUID(image_id), ProductImage.is_removed == False)
    )).scalar_one_or_none()
    if not image:
        raise HTTPException(status_code=404, detail="Image not found")
    image.is_removed = True
    image.update_date = datetime.now(timezone.utc)
    await db.commit()
    return RedirectResponse(url=f"/administration/products/{product_id}/details", status_code=303)


@router.post("/products/{product_id}/datasheets/new", response_class=HTMLResponse)
async def admin_product_datasheet_create(
    request: Request,
    product_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    import uuid
    pid = uuid.UUID(product_id)
    product = await product_service.get_product_by_id(db, pid)
    if not product:
        raise HTTPException(status_code=404, detail="Product not found")
    form = await request.form()
    datasheet = MenuDatasheet(
        id=uuid.uuid4(),
        product_id=pid,
        type=form.get("type") or None,
        file_url=form.get("file_url") or None,
        complete_file_url=form.get("file_url") or None,
        created_by_user_id=current_user.id,
        insert_date=datetime.now(timezone.utc),
        update_date=datetime.now(timezone.utc),
    )
    db.add(datasheet)
    db.add(Log(
        record_id=datasheet.id, table_name="menu_datasheets",
        description=f"افزودن برگه اطلاعات به محصول: {product.name}",
        created_by_user_id=current_user.id, type="Create",
    ))
    await db.commit()
    return RedirectResponse(url=f"/administration/products/{product_id}/details", status_code=303)


@router.post("/products/{product_id}/datasheets/{datasheet_id}/delete", response_class=HTMLResponse)
async def admin_product_datasheet_delete(
    request: Request,
    product_id: str,
    datasheet_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    import uuid
    datasheet = (await db.execute(
        select(MenuDatasheet).where(MenuDatasheet.id == uuid.UUID(datasheet_id), MenuDatasheet.is_removed == False)
    )).scalar_one_or_none()
    if not datasheet:
        raise HTTPException(status_code=404, detail="Datasheet not found")
    datasheet.is_removed = True
    datasheet.update_date = datetime.now(timezone.utc)
    await db.commit()
    return RedirectResponse(url=f"/administration/products/{product_id}/details", status_code=303)


@router.post("/products/{product_id}/supplier-products/new", response_class=HTMLResponse)
async def admin_product_supplier_create(
    request: Request,
    product_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    import uuid
    pid = uuid.UUID(product_id)
    product = await product_service.get_product_by_id(db, pid)
    if not product:
        raise HTTPException(status_code=404, detail="Product not found")
    form = await request.form()
    supplier_id = form.get("supplier_id")
    link = form.get("link")
    if not supplier_id:
        raise HTTPException(status_code=400, detail="Supplier is required")
    supplier_product = SupplierProduct(
        id=uuid.uuid4(),
        product_id=pid,
        supplier_id=uuid.UUID(supplier_id) if supplier_id else None,
        link=link or None,
        created_by_user_id=current_user.id,
        insert_date=datetime.now(timezone.utc),
        update_date=datetime.now(timezone.utc),
    )
    db.add(supplier_product)
    db.add(Log(
        record_id=supplier_product.id, table_name="supplier_products",
        description=f"افزودن تامین‌کننده به محصول: {product.name}",
        created_by_user_id=current_user.id, type="Create",
    ))
    await db.commit()
    return RedirectResponse(url=f"/administration/products/{product_id}/details", status_code=303)


@router.post("/products/{product_id}/supplier-products/{sp_id}/delete", response_class=HTMLResponse)
async def admin_product_supplier_delete(
    request: Request,
    product_id: str,
    sp_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    import uuid
    sp = (await db.execute(
        select(SupplierProduct).where(SupplierProduct.id == uuid.UUID(sp_id), SupplierProduct.is_removed == False)
    )).scalar_one_or_none()
    if not sp:
        raise HTTPException(status_code=404, detail="Supplier product not found")
    sp.is_removed = True
    sp.update_date = datetime.now(timezone.utc)
    await db.commit()
    return RedirectResponse(url=f"/administration/products/{product_id}/details", status_code=303)


@router.post("/products/{product_id}/technical-values/set", response_class=HTMLResponse)
async def admin_product_technical_value_set(
    request: Request,
    product_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    import uuid
    pid = uuid.UUID(product_id)
    product = await product_service.get_product_by_id(db, pid)
    if not product:
        raise HTTPException(status_code=404, detail="Product not found")
    form = await request.form()
    feature_id = form.get("technical_feature_id")
    table_product_id = form.get("table_product_id")
    if not feature_id or not table_product_id:
        raise HTTPException(status_code=400, detail="Feature and table are required")

    table_product = (await db.execute(
        select(TechnicalTableProduct).where(
            TechnicalTableProduct.id == uuid.UUID(table_product_id),
            TechnicalTableProduct.product_id == pid,
            TechnicalTableProduct.is_removed == False,
        )
    )).scalar_one_or_none()
    if not table_product:
        raise HTTPException(status_code=404, detail="Technical table product not found")

    existing = (await db.execute(
        select(TechnicalFeatureValue).where(
            TechnicalFeatureValue.technical_table_product_id == table_product.id,
            TechnicalFeatureValue.technical_feature_id == uuid.UUID(feature_id),
            TechnicalFeatureValue.is_removed == False,
        )
    )).scalar_one_or_none()

    def _num(key):
        val = form.get(key)
        if val is None or val == "":
            return None
        try:
            return float(val)
        except Exception:
            return None

    enum_id = form.get("technical_feature_enum_id")
    enum1_id = form.get("technical_feature_enum1_id")

    payload = dict(
        technical_feature_id=uuid.UUID(feature_id),
        technical_table_product_id=table_product.id,
        min_value=_num("min_value"),
        max_value=_num("max_value"),
        min_unit=form.get("min_unit") or None,
        max_unit=form.get("max_unit") or None,
        d_value=_num("d_value"),
        unit=form.get("unit") or None,
        s_value=form.get("s_value") or None,
        b_value=True if form.get("b_value") == "True" else False if form.get("b_value") == "False" else None,
        x_value=_num("x_value"),
        x_unit=form.get("x_unit") or None,
        y_value=_num("y_value"),
        y_unit=form.get("y_unit") or None,
        z_value=_num("z_value"),
        z_unit=form.get("z_unit") or None,
        technical_feature_enum_id=uuid.UUID(enum_id) if enum_id else None,
        technical_feature_enum1_id=uuid.UUID(enum1_id) if enum1_id else None,
        general_feature=form.get("general_feature") in ("1", "true", "on"),
        created_by_user_id=current_user.id,
    )
    if existing:
        for key, value in payload.items():
            setattr(existing, key, value)
        existing.update_date = datetime.now(timezone.utc)
        db.add(Log(
            record_id=existing.id, table_name="technical_feature_values",
            description=f"به‌روزرسانی ویژگی فنی محصول: {product.name}",
            created_by_user_id=current_user.id, type="Update",
        ))
    else:
        new_value = TechnicalFeatureValue(id=uuid.uuid4(), **payload)
        db.add(new_value)
        db.add(Log(
            record_id=new_value.id, table_name="technical_feature_values",
            description=f"افزودن ویژگی فنی به محصول: {product.name}",
            created_by_user_id=current_user.id, type="Create",
        ))
    await db.commit()
    return RedirectResponse(url=f"/administration/products/{product_id}/details", status_code=303)


@router.post("/products/{product_id}/technical-values/{value_id}/delete", response_class=HTMLResponse)
async def admin_product_technical_value_delete(
    request: Request,
    product_id: str,
    value_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    import uuid
    value = (await db.execute(
        select(TechnicalFeatureValue).where(TechnicalFeatureValue.id == uuid.UUID(value_id), TechnicalFeatureValue.is_removed == False)
    )).scalar_one_or_none()
    if not value:
        raise HTTPException(status_code=404, detail="Feature value not found")
    value.is_removed = True
    value.update_date = datetime.now(timezone.utc)
    await db.commit()
    return RedirectResponse(url=f"/administration/products/{product_id}/details", status_code=303)


@router.get("/products/{product_id}/delete", response_class=HTMLResponse)
async def admin_product_delete(
    request: Request,
    product_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    pid = uuid.UUID(product_id)
    product = await product_service.get_product_by_id(db, pid)
    if product is None:
        raise HTTPException(status_code=404, detail="Product not found")
    return templates.TemplateResponse("admin/product_delete.html", {
        "request": request, "current_user": current_user, "product": product,
    })


@router.post("/products/{product_id}/delete", response_class=HTMLResponse)
async def admin_product_delete_confirm(
    request: Request,
    product_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    pid = uuid.UUID(product_id)
    product = await product_service.get_product_by_id(db, pid)
    if product is None:
        raise HTTPException(status_code=404, detail="Product not found")
    name = product.name
    category = product.category
    await product_service.delete_product(db, product)
    if category is not None:
        category.product_count = max((category.product_count or 0) - 1, 0)
    db.add(Log(
        record_id=pid,
        table_name="products",
        description=f"حذف محصول: {name}",
        created_by_user_id=current_user.id,
        type="Delete",
    ))
    await db.commit()
    return RedirectResponse(url="/administration/products", status_code=303)


def _format_tfv(value: TechnicalFeatureValue, display_format: str) -> str:
    """Render a technical feature value cell — mirrors the .NET
    TechnicalFeatureValue.Value(displayFormat) string.Format call."""
    if not display_format:
        return ""
    def _enum_name(enum):
        return enum.persian_name if enum else None
    args = (
        value.d_value if value.d_value is not None else "-",
        value.unit or "",
        value.s_value or "",
        _enum_name(value.technical_feature_enum) or value.e_value or "",
        str(value.b_value) if value.b_value is not None else "",
        value.min_value if value.min_value is not None else "-",
        value.min_unit or "",
        value.max_value if value.max_value is not None else "-",
        value.max_unit or "",
        value.x_value if value.x_value is not None else "-",
        value.x_unit or "",
        value.y_value if value.y_value is not None else "-",
        value.y_unit or "",
        value.z_value if value.z_value is not None else "-",
        value.z_unit or "",
        _enum_name(value.technical_feature_enum1) or value.e_value1 or "",
    )
    try:
        return display_format.format(*args)
    except Exception:
        return display_format


def _to_fa_datetime(value):
    """Convert an aware/naive datetime to a Persian (jalali) datetime string."""
    import jdatetime
    if not value:
        return "—"
    try:
        dt = value
        if dt.tzinfo is not None:
            dt = dt.replace(tzinfo=None)
        return jdatetime.datetime.fromgregorian(datetime=dt).strftime("%Y/%m/%d - %H:%M")
    except Exception:
        return "—"


def _to_fa_date(value):
    """Convert an aware/naive date to a Persian (jalali) date string."""
    import jdatetime
    if not value:
        return "—"
    try:
        dt = value
        if dt.tzinfo is not None:
            dt = dt.replace(tzinfo=None)
        return jdatetime.datetime.fromgregorian(datetime=dt).strftime("%Y/%m/%d")
    except Exception:
        return "—"


# ── Categories ──

def _flatten_category_options(categories: list[Category], depth: int = 0) -> list[dict]:
    """Flatten the category list into (id, title, depth) options for the parent picker.

    Builds the hierarchy in memory from parent_category_id to avoid async lazy-loading.
    """
    by_parent: dict[str, list[Category]] = {}
    id_to_cat: dict[str, Category] = {}
    for cat in categories:
        id_to_cat[str(cat.id)] = cat
        key = str(cat.parent_category_id) if cat.parent_category_id is not None else ""
        by_parent.setdefault(key, []).append(cat)

    options: list[dict] = []

    def _walk(parent_key: str, depth: int) -> None:
        for cat in by_parent.get(parent_key, []):
            options.append({"id": str(cat.id), "title": cat.title, "depth": depth})
            _walk(str(cat.id), depth + 1)

    _walk("", 0)
    return options


async def _get_descendant_ids(db: AsyncSession, category_id: uuid.UUID) -> set[str]:
    """Collect ids of the category and all its descendants (to exclude from the parent picker)."""
    ids = {str(category_id)}
    stack = [category_id]
    while stack:
        current = stack.pop()
        stmt = select(Category.id).where(
            Category.parent_category_id == current, Category.is_removed == False
        )
        child_ids = [row[0] for row in (await db.execute(stmt)).all()]
        for child_id in child_ids:
            sid = str(child_id)
            if sid not in ids:
                ids.add(sid)
                stack.append(child_id)
    return ids


@router.get("/categories", response_class=HTMLResponse)
async def admin_categories(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    cats = await product_service.get_admin_categories(db)
    return templates.TemplateResponse("admin/categories.html", {
        "request": request, "current_user": current_user, "categories": cats,
    })


@router.get("/categories/new", response_class=HTMLResponse)
async def admin_category_create(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    parent: str | None = request.query_params.get("parent")
    tree = await product_service.get_admin_category_tree(db)
    options = _flatten_category_options(tree)
    return templates.TemplateResponse("admin/category_form.html", {
        "request": request, "current_user": current_user,
        "categories": options, "category": None, "parent_id": parent,
        "exclude_ids": set(),
    })


@router.post("/categories/new", response_class=HTMLResponse)
async def admin_category_create_submit(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    form = await request.form()
    parent_raw = form.get("parent_category_id")
    data = {
        "title": (form.get("title") or "").strip(),
        "en_title": (form.get("en_title") or "").strip(),
        "slug": (form.get("slug") or "").strip(),
        "keywords": (form.get("keywords") or "").strip(),
        "description": form.get("description") or "",
        "meta_description": form.get("meta_description") or "",
        "place": (form.get("place") or "").strip(),
        "priority": int(form.get("priority") or 0),
        "no_display": form.get("no_display") in ("1", "true", "on"),
        "is_disable": form.get("is_disable") in ("1", "true", "on"),
        "parent_category_id": uuid.UUID(parent_raw) if parent_raw else None,
    }
    if data["en_title"]:
        data["en_title"] = generate_slug(data["en_title"])
    if not data["slug"] and data["title"]:
        data["slug"] = generate_slug(data["title"])

    category = await product_service.create_category(db, CategoryCreate(**data), current_user.id)
    log_text = form.get("log") or ""
    db.add(Log(
        record_id=category.id,
        table_name="categories",
        description=log_text,
        created_by_user_id=current_user.id,
        type="Create",
    ))
    return RedirectResponse(url="/administration/categories", status_code=303)


@router.get("/categories/{category_id}/edit", response_class=HTMLResponse)
async def admin_category_edit(
    request: Request,
    category_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    cid = uuid.UUID(category_id)
    category = await product_service.get_category_by_id(db, cid)
    if category is None:
        raise HTTPException(status_code=404, detail="Category not found")
    tree = await product_service.get_admin_category_tree(db)
    options = _flatten_category_options(tree)
    exclude_ids = await _get_descendant_ids(db, cid)
    return templates.TemplateResponse("admin/category_form.html", {
        "request": request, "current_user": current_user,
        "category": category, "categories": options, "parent_id": None,
        "exclude_ids": exclude_ids,
    })


@router.post("/categories/{category_id}/edit", response_class=HTMLResponse)
async def admin_category_edit_submit(
    request: Request,
    category_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    cid = uuid.UUID(category_id)
    category = await product_service.get_category_by_id(db, cid)
    if category is None:
        raise HTTPException(status_code=404, detail="Category not found")

    form = await request.form()
    parent_raw = form.get("parent_category_id")
    data = {
        "title": (form.get("title") or "").strip(),
        "en_title": (form.get("en_title") or "").strip(),
        "slug": (form.get("slug") or "").strip(),
        "keywords": (form.get("keywords") or "").strip(),
        "description": form.get("description") or "",
        "meta_description": form.get("meta_description") or "",
        "place": (form.get("place") or "").strip(),
        "priority": int(form.get("priority") or 0),
        "no_display": form.get("no_display") in ("1", "true", "on"),
        "is_disable": form.get("is_disable") in ("1", "true", "on"),
        "parent_category_id": uuid.UUID(parent_raw) if parent_raw else None,
    }
    if data["en_title"]:
        data["en_title"] = generate_slug(data["en_title"])
    if not data["slug"] and data["title"]:
        data["slug"] = generate_slug(data["title"])

    await product_service.update_category(db, category, CategoryUpdate(**data))

    if category.no_display:
        await _hide_category_children(db, category)

    log_text = form.get("log") or ""
    if log_text:
        db.add(Log(
            record_id=category.id,
            table_name="categories",
            description=log_text,
            created_by_user_id=current_user.id,
            type="Update",
        ))
    return RedirectResponse(url="/administration/categories", status_code=303)


@router.get("/categories/{category_id}/details", response_class=HTMLResponse)
async def admin_category_details(
    request: Request,
    category_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    cid = uuid.UUID(category_id)
    stmt = (
        select(Category)
        .options(
            selectinload(Category.parent),
            selectinload(Category.medias),
            selectinload(Category.products).selectinload(Product.category),
        )
        .where(Category.id == cid, Category.is_removed == False)
    )
    result = await db.execute(stmt)
    category = result.unique().scalar_one_or_none()
    if category is None:
        raise HTTPException(status_code=404, detail="Category not found")
    child_stmt = (
        select(Category)
        .options(selectinload(Category.parent), selectinload(Category.medias))
        .where(Category.parent_category_id == cid, Category.is_removed == False)
        .order_by(Category.title)
    )
    child_result = await db.execute(child_stmt)
    children = child_result.unique().scalars().all()
    # Get products for the category (including child categories, with pagination support)
    products, total_products = await product_service.get_products_by_category(db, cid, 1, 100)
    media_stmt = (
        select(Media)
        .where(Media.category_id == cid, Media.is_removed == False)
        .order_by(Media.picture_order, Media.insert_date)
    )
    medias = (await db.execute(media_stmt)).scalars().all()
    return templates.TemplateResponse("admin/category_detail.html", {
        "request": request, "current_user": current_user,
        "category": category, "children": children,
        "products": products, "total_products": total_products,
        "medias": medias,
    })


@router.get("/categories/{category_id}/delete", response_class=HTMLResponse)
async def admin_category_delete(
    request: Request,
    category_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    cid = uuid.UUID(category_id)
    category = await product_service.get_category_by_id(db, cid)
    if category is None:
        raise HTTPException(status_code=404, detail="Category not found")
    return templates.TemplateResponse("admin/category_delete.html", {
        "request": request, "current_user": current_user, "category": category,
    })


@router.post("/categories/{category_id}/delete", response_class=HTMLResponse)
async def admin_category_delete_confirm(
    request: Request,
    category_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    cid = uuid.UUID(category_id)
    category = await product_service.get_category_by_id(db, cid)
    if category is None:
        raise HTTPException(status_code=404, detail="Category not found")
    title = category.title
    await product_service.delete_category(db, category)
    db.add(Log(
        record_id=cid,
        table_name="categories",
        description=f"حذف دسته بندی: {title}",
        created_by_user_id=current_user.id,
        type="Delete",
    ))
    return RedirectResponse(url="/administration/categories", status_code=303)


async def _hide_category_children(db: AsyncSession, category: Category) -> None:
    """Mirrors .NET HideChildrenAsync — sets no_display True on all descendants."""
    stack = [category]
    while stack:
        current = stack.pop()
        child_stmt = (
            select(Category)
            .options(selectinload(Category.children))
            .where(Category.parent_category_id == current.id, Category.is_removed == False)
        )
        child_result = await db.execute(child_stmt)
        children = child_result.unique().scalars().all()
        for child in children:
            child.no_display = True
            stack.append(child)


# ── Category Media (ManageMedias) ──

async def _get_media(db: AsyncSession, media_id: str) -> Optional[Media]:
    try:
        mid = uuid.UUID(media_id)
    except (ValueError, AttributeError):
        return None
    return (await db.execute(
        select(Media).where(Media.id == mid, Media.is_removed == False)
    )).scalar_one_or_none()


async def _set_category_poster(db: AsyncSession, category: Category, media: Media) -> None:
    """Mirrors .NET behavior: only one PosterImage per category; the URL is
    promoted to Categories.PosterImageURL (used by the shop subcategory boxes)."""
    others = (await db.execute(
        select(Media).where(Media.category_id == category.id, Media.is_removed == False)
    )).scalars().all()
    for m in others:
        if m.poster_image and m.id != media.id:
            m.poster_image = False
    media.poster_image = True
    category.poster_image_url = media.url


@router.get("/categories/{category_id}/add-picture", response_class=HTMLResponse)
async def admin_category_add_picture(
    request: Request,
    category_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    cid = uuid.UUID(category_id)
    category = await product_service.get_category_by_id(db, cid)
    if category is None:
        raise HTTPException(status_code=404, detail="Category not found")
    return templates.TemplateResponse("admin/media_form.html", {
        "request": request, "current_user": current_user,
        "category": category, "media": None,
    })


@router.post("/categories/{category_id}/add-picture", response_class=HTMLResponse)
async def admin_category_add_picture_submit(
    request: Request,
    category_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    cid = uuid.UUID(category_id)
    category = await product_service.get_category_by_id(db, cid)
    if category is None:
        raise HTTPException(status_code=404, detail="Category not found")
    form = await request.form()
    poster_flag = form.get("poster_image") in ("1", "true", "on")
    media = Media(
        id=uuid.uuid4(),
        category_id=cid,
        url=(form.get("url") or "").strip(),
        title=(form.get("title") or "").strip(),
        description=form.get("description") or None,
        display_photo=form.get("display_photo") in ("1", "true", "on"),
        poster_image=poster_flag,
        is_video=False,
        type="Image",
        picture_order=int(form.get("picture_order") or 0),
        created_by_user_id=current_user.id,
        insert_date=datetime.now(timezone.utc),
        update_date=datetime.now(timezone.utc),
    )
    if poster_flag:
        await _set_category_poster(db, category, media)
    db.add(media)
    db.add(Log(
        record_id=media.id, table_name="medias",
        description=f"افزودن عکس به دسته بندی: {category.title}",
        created_by_user_id=current_user.id, type="Create",
    ))
    await db.commit()
    return RedirectResponse(url=f"/administration/categories/{category_id}/details", status_code=303)


@router.get("/medias/{media_id}/edit", response_class=HTMLResponse)
async def admin_media_edit(
    request: Request,
    media_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    media = await _get_media(db, media_id)
    if media is None:
        raise HTTPException(status_code=404, detail="Media not found")
    category = await product_service.get_category_by_id(db, media.category_id) if media.category_id else None
    return templates.TemplateResponse("admin/media_form.html", {
        "request": request, "current_user": current_user,
        "category": category, "media": media,
    })


@router.post("/medias/{media_id}/edit", response_class=HTMLResponse)
async def admin_media_edit_submit(
    request: Request,
    media_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    media = await _get_media(db, media_id)
    if media is None:
        raise HTTPException(status_code=404, detail="Media not found")
    old_url = media.url
    form = await request.form()
    media.url = (form.get("url") or "").strip()
    media.title = (form.get("title") or "").strip()
    media.description = form.get("description") or None
    media.picture_order = int(form.get("picture_order") or 0)
    media.display_photo = form.get("display_photo") in ("1", "true", "on")
    media.update_date = datetime.now(timezone.utc)

    category = await product_service.get_category_by_id(db, media.category_id) if media.category_id else None
    if category is not None:
        if form.get("poster_image") in ("1", "true", "on"):
            await _set_category_poster(db, category, media)
        else:
            media.poster_image = False
            if category.poster_image_url in (old_url, media.url):
                category.poster_image_url = None
    db.add(Log(
        record_id=media.id, table_name="medias",
        description=f"ویرایش عکس دسته بندی: {media.title or category.title if category else 'دسته بندی'}",
        created_by_user_id=current_user.id, type="Update",
    ))
    await db.commit()
    if category is not None:
        return RedirectResponse(url=f"/administration/categories/{category.id}/details", status_code=303)
    return RedirectResponse(url="/administration/categories", status_code=303)


@router.get("/medias/{media_id}/delete", response_class=HTMLResponse)
async def admin_media_delete(
    request: Request,
    media_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    media = await _get_media(db, media_id)
    if media is None:
        raise HTTPException(status_code=404, detail="Media not found")
    category = await product_service.get_category_by_id(db, media.category_id) if media.category_id else None
    return templates.TemplateResponse("admin/media_delete.html", {
        "request": request, "current_user": current_user,
        "category": category, "media": media,
    })


@router.post("/medias/{media_id}/delete", response_class=HTMLResponse)
async def admin_media_delete_submit(
    request: Request,
    media_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    media = await _get_media(db, media_id)
    if media is None:
        raise HTTPException(status_code=404, detail="Media not found")
    category = await product_service.get_category_by_id(db, media.category_id) if media.category_id else None
    if category is not None and category.poster_image_url == media.url:
        category.poster_image_url = None
    db.add(Log(
        record_id=media.id, table_name="medias",
        description=f"حذف عکس دسته بندی: {media.title or (category.title if category else '')}",
        created_by_user_id=current_user.id, type="Delete",
    ))
    media.is_removed = True
    media.update_date = datetime.now(timezone.utc)
    await db.commit()
    if category is not None:
        return RedirectResponse(url=f"/administration/categories/{category.id}/details", status_code=303)
    return RedirectResponse(url="/administration/categories", status_code=303)


# ── Brands ──

@router.get("/brands", response_class=HTMLResponse)
async def admin_brands(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    stmt = (
        select(Brand)
        .options(selectinload(Brand.products))
        .where(Brand.is_removed == False)
        .order_by(Brand.insert_date.desc())
    )
    brands = (await db.execute(stmt)).unique().scalars().all()
    return templates.TemplateResponse("admin/brands.html", {
        "request": request, "current_user": current_user, "items": brands,
    })


@router.get("/brands/new", response_class=HTMLResponse)
async def admin_brand_create(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    return templates.TemplateResponse("admin/brand_form.html", {
        "request": request, "current_user": current_user, "brand": None,
    })


@router.post("/brands/new", response_class=HTMLResponse)
async def admin_brand_create_submit(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    form = await request.form()
    name = (form.get("name") or "").strip()

    brand = Brand(name=name, created_by_user_id=current_user.id)
    db.add(brand)
    await db.flush()

    log_text = form.get("log") or ""
    if log_text:
        db.add(Log(
            record_id=brand.id,
            table_name="brands",
            description=log_text,
            created_by_user_id=current_user.id,
            type="Create",
        ))
    return RedirectResponse(url="/administration/brands", status_code=303)


@router.get("/brands/{brand_id}/edit", response_class=HTMLResponse)
async def admin_brand_edit(
    request: Request,
    brand_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    brand = await db.get(Brand, uuid.UUID(brand_id))
    if brand is None or brand.is_removed:
        raise HTTPException(status_code=404, detail="Brand not found")
    return templates.TemplateResponse("admin/brand_form.html", {
        "request": request, "current_user": current_user, "brand": brand,
    })


@router.post("/brands/{brand_id}/edit", response_class=HTMLResponse)
async def admin_brand_edit_submit(
    request: Request,
    brand_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    brand = await db.get(Brand, uuid.UUID(brand_id))
    if brand is None or brand.is_removed:
        raise HTTPException(status_code=404, detail="Brand not found")

    form = await request.form()
    brand.name = (form.get("name") or "").strip()
    await db.flush()

    log_text = form.get("log") or ""
    if log_text:
        db.add(Log(
            record_id=brand.id,
            table_name="brands",
            description=log_text,
            created_by_user_id=current_user.id,
            type="Update",
        ))
    return RedirectResponse(url="/administration/brands", status_code=303)


@router.get("/brands/{brand_id}", response_class=HTMLResponse)
async def admin_brand_detail(
    request: Request,
    brand_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    brand = await db.get(Brand, uuid.UUID(brand_id))
    if brand is None or brand.is_removed:
        raise HTTPException(status_code=404, detail="Brand not found")
    return templates.TemplateResponse("admin/brand_detail.html", {
        "request": request, "current_user": current_user, "brand": brand,
    })


@router.post("/brands/{brand_id}/delete", response_class=HTMLResponse)
async def admin_brand_delete(
    request: Request,
    brand_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    brand = await db.get(Brand, uuid.UUID(brand_id))
    if brand is None or brand.is_removed:
        raise HTTPException(status_code=404, detail="Brand not found")
    name = brand.name
    brand.is_removed = True
    await db.flush()
    db.add(Log(
        record_id=brand.id,
        table_name="brands",
        description=f"حذف برند: {name}",
        created_by_user_id=current_user.id,
        type="Delete",
    ))
    return RedirectResponse(url="/administration/brands", status_code=303)


# ── Orders ──

@router.get("/orders", response_class=HTMLResponse)
async def admin_orders(
    request: Request,
    page: int = Query(1),
    status_filter: str = Query(None),
    part_number: str = Query(None),
    current_user: User = Depends(require_any_role("Admin", "Orders Manager")),
    db: AsyncSession = Depends(get_db),
):
    orders, total = await order_service.get_admin_orders(db, page, 20, status_filter, part_number)
    # compute has_invoice per order
    result = []
    for o in orders:
        d = order_service.build_admin_order_response(o)
        d["has_invoice"] = (await db.execute(
            select(func.count(Invoice.id)).where(Invoice.order_id == o.id, Invoice.is_removed == False)
        )).scalar() > 0
        result.append(d)
    return templates.TemplateResponse("admin/orders.html", {
        "request": request, "current_user": current_user,
        "orders": result,
        "total": total, "page": page, "total_pages": (total + 19) // 20,
        "status_filter": status_filter, "part_number": part_number or "",
        "order_status_names": order_service.ORDER_STATUS_NAMES,
    })


@router.get("/orders/create", response_class=HTMLResponse)
async def admin_order_create_form(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Orders Manager")),
    db: AsyncSession = Depends(get_db),
):
    users = (await db.execute(
        select(User).where(User.is_removed == False).order_by(User.first_name)
    )).scalars().all()
    post_types = await order_service.get_post_types(db)
    pay_methods = await order_service.get_pay_methods(db)
    return templates.TemplateResponse("admin/order_form.html", {
        "request": request, "current_user": current_user,
        "order": None, "users": users, "post_types": post_types, "pay_methods": pay_methods,
        "order_status_names": order_service.ORDER_STATUS_NAMES,
        "default_status": "Ordering",
    })


@router.post("/orders/create", response_class=HTMLResponse)
async def admin_order_create_submit(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Orders Manager")),
    db: AsyncSession = Depends(get_db),
):
    form = await request.form()
    user_id = form.get("user_id", "")
    try:
        uid = uuid.UUID(user_id)
    except ValueError:
        raise HTTPException(status_code=400, detail="Invalid user")
    user = await db.get(User, uid)
    if not user:
        raise HTTPException(status_code=404, detail="User not found")

    tracking_number = (form.get("tracking_number") or "").strip()
    post_type_id = form.get("post_type_id") or None
    weight = (form.get("weight") or "").strip() or None
    postage_date_str = (form.get("postage_date_str") or "").strip()
    date_str = (form.get("date_str") or "").strip()
    notes = (form.get("notes") or "").strip() or None
    status = form.get("order_status") or "Ordering"

    postage_date = from_farsi_date(postage_date_str) if postage_date_str else None
    date = from_farsi_date(date_str) if date_str else datetime.now(timezone.utc)

    post_type = None
    if post_type_id:
        try:
            post_type = await db.get(PostType, uuid.UUID(post_type_id))
        except ValueError:
            post_type = None

    ref_code = await order_service._generate_reference_code(db)

    order = Order(
        id=uuid.uuid4(),
        reference_code=ref_code,
        user_id=uid,
        email=user.email,
        tracking_number=tracking_number,
        order_status=status,
        count=0,
        notes=notes,
        weight=weight,
        postage_date=postage_date,
        date=date or datetime.now(timezone.utc),
        post_type_id=post_type.id if post_type else None,
        insert_date=datetime.now(timezone.utc),
        update_date=datetime.now(timezone.utc),
    )
    if post_type:
        order.postage_fee = float(post_type.price or 0)
        order.post_vat_rate = float(post_type.post_vat_rate or 0)
        order.post_vat = float(post_type.post_vat or 0)
        order.vat = order.post_vat
        order.payable = order.vat + order.postage_fee

    db.add(order)
    await db.flush()

    status_record = OrderStatusRecord(
        id=uuid.uuid4(),
        order_id=order.id,
        status=status,
        comment="سفارش ثبت شد",
        created_by_user_id=current_user.id,
        insert_date=datetime.now(timezone.utc),
        update_date=datetime.now(timezone.utc),
    )
    db.add(status_record)

    db.add(Log(record_id=order.id, table_name="orders",
               description=f"ایجاد سفارش: {ref_code}",
               created_by_user_id=current_user.id, type="Create"))
    await db.commit()
    return RedirectResponse(url="/administration/orders", status_code=303)


@router.get("/orders/{order_id}", response_class=HTMLResponse)
async def admin_order_detail(
    request: Request,
    order_id: str,
    page_products: int = Query(1, alias="page"),
    current_user: User = Depends(require_any_role("Admin", "Orders Manager")),
    db: AsyncSession = Depends(get_db),
):
    try:
        oid = uuid.UUID(order_id)
    except ValueError:
        raise HTTPException(status_code=404, detail="Invalid order ID")
    order = await order_service.get_admin_order_detail(db, oid)
    if not order:
        raise HTTPException(status_code=404, detail="Order not found")

    # Find the linked invoice (for factor link)
    invoice = (await db.execute(
        select(Invoice).where(Invoice.order_id == oid, Invoice.is_removed == False)
    )).scalar_one_or_none()

    # Resolve created-by users for status records
    status_records = order_service.build_admin_order_response(order)["order_status_records"]
    created_by_map = {}
    ids = {sr["created_by_user_id"] for sr in status_records if sr["created_by_user_id"]}
    if ids:
        users = (await db.execute(select(User).where(User.id.in_(ids)))).scalars().all()
        created_by_map = {u.id: u.full_name for u in users}

    return templates.TemplateResponse("admin/order_detail.html", {
        "request": request, "current_user": current_user,
        "order": order_service.build_admin_order_response(order),
        "invoice": invoice,
        "order_status_names": order_service.ORDER_STATUS_NAMES,
        "created_by_map": created_by_map,
    })


@router.get("/orders/{order_id}/edit", response_class=HTMLResponse)
async def admin_order_edit_form(
    request: Request,
    order_id: str,
    current_user: User = Depends(require_any_role("Admin", "Orders Manager")),
    db: AsyncSession = Depends(get_db),
):
    try:
        oid = uuid.UUID(order_id)
    except ValueError:
        raise HTTPException(status_code=404, detail="Invalid order ID")
    order = await order_service.get_admin_order_detail(db, oid)
    if not order:
        raise HTTPException(status_code=404, detail="Order not found")
    users = (await db.execute(
        select(User).where(User.is_removed == False).order_by(User.first_name)
    )).scalars().all()
    post_types = await order_service.get_post_types(db)
    pay_methods = await order_service.get_pay_methods(db)
    return templates.TemplateResponse("admin/order_form.html", {
        "request": request, "current_user": current_user,
        "order": order_service.build_admin_order_response(order),
        "users": users, "post_types": post_types, "pay_methods": pay_methods,
        "order_status_names": order_service.ORDER_STATUS_NAMES,
    })


@router.post("/orders/{order_id}/edit", response_class=HTMLResponse)
async def admin_order_edit_submit(
    request: Request,
    order_id: str,
    current_user: User = Depends(require_any_role("Admin", "Orders Manager")),
    db: AsyncSession = Depends(get_db),
):
    try:
        oid = uuid.UUID(order_id)
    except ValueError:
        raise HTTPException(status_code=404, detail="Invalid order ID")
    order = await order_service.get_admin_order_detail(db, oid)
    if not order:
        raise HTTPException(status_code=404, detail="Order not found")
    form = await request.form()

    user_id = form.get("user_id", "")
    try:
        uid = uuid.UUID(user_id)
        user = await db.get(User, uid)
        if not user:
            raise HTTPException(status_code=404, detail="User not found")
        order.user_id = uid
    except ValueError:
        pass

    order.tracking_number = (form.get("tracking_number") or "").strip() or None
    order.notes = (form.get("notes") or "").strip() or None
    order.weight = (form.get("weight") or "").strip() or None

    post_type_id = form.get("post_type_id") or None
    if post_type_id:
        try:
            post_type = await db.get(PostType, uuid.UUID(post_type_id))
            if post_type:
                order.post_type_id = post_type.id
        except ValueError:
            pass

    pay_method_id = form.get("pay_method_id") or None
    if pay_method_id:
        try:
            pay_method = await db.get(PayMethod, uuid.UUID(pay_method_id))
            if pay_method:
                order.pay_method_id = pay_method.id
        except ValueError:
            pass

    postage_date_str = (form.get("postage_date_str") or "").strip()
    if postage_date_str:
        order.postage_date = from_farsi_date(postage_date_str)
    date_str = (form.get("date_str") or "").strip()
    if date_str:
        order.date = from_farsi_date(date_str)

    new_status = form.get("order_status")
    if new_status and new_status != order.order_status:
        order.order_status = new_status

    order.update_date = datetime.now(timezone.utc)
    db.add(Log(record_id=order.id, table_name="orders",
               description=f"ویرایش سفارش: {order.reference_code}",
               created_by_user_id=current_user.id, type="Update"))
    await db.commit()
    return RedirectResponse(url=f"/administration/orders/{order.id}", status_code=303)


@router.post("/orders/{order_id}/status", response_class=HTMLResponse)
async def admin_order_update_status(
    request: Request,
    order_id: str,
    current_user: User = Depends(require_any_role("Admin", "Orders Manager")),
    db: AsyncSession = Depends(get_db),
):
    try:
        oid = uuid.UUID(order_id)
    except ValueError:
        raise HTTPException(status_code=404, detail="Invalid order ID")
    order = await db.get(Order, oid)
    if not order:
        raise HTTPException(status_code=404, detail="Order not found")
    form = await request.form()
    status = form.get("status") or order.order_status
    comment = (form.get("comment") or "").strip() or None
    tracking_number = (form.get("tracking_number") or "").strip() or None

    if status == "Sending" and not tracking_number:
        tracking_number = order.tracking_number

    order.order_status = status
    if tracking_number:
        order.tracking_number = tracking_number
    order.update_date = datetime.now(timezone.utc)

    status_record = OrderStatusRecord(
        id=uuid.uuid4(),
        order_id=order.id,
        status=status,
        comment=comment,
        tracking_number=tracking_number,
        created_by_user_id=current_user.id,
        insert_date=datetime.now(timezone.utc),
        update_date=datetime.now(timezone.utc),
    )
    db.add(status_record)
    await db.commit()
    return RedirectResponse(url=f"/administration/orders/{order.id}", status_code=303)


@router.post("/orders/{order_id}/delete", response_class=HTMLResponse)
async def admin_order_delete(
    request: Request,
    order_id: str,
    current_user: User = Depends(require_any_role("Admin", "Orders Manager")),
    db: AsyncSession = Depends(get_db),
):
    try:
        oid = uuid.UUID(order_id)
    except ValueError:
        raise HTTPException(status_code=404, detail="Invalid order ID")
    order = await db.get(Order, oid)
    if not order:
        raise HTTPException(status_code=404, detail="Order not found")
    ref = order.reference_code
    order.is_removed = True
    order.update_date = datetime.now(timezone.utc)
    db.add(Log(record_id=order.id, table_name="orders",
               description=f"حذف سفارش: {ref}",
               created_by_user_id=current_user.id, type="Delete"))
    await db.commit()
    return RedirectResponse(url="/administration/orders", status_code=303)


@router.get("/orders/{order_id}/factor", response_class=HTMLResponse)
async def admin_order_factor(
    request: Request,
    order_id: str,
    current_user: User = Depends(require_any_role("Admin", "Orders Manager")),
    db: AsyncSession = Depends(get_db),
):
    """Printable invoice (فاکتور) for an order — mirrors .NET Factor action."""
    try:
        oid = uuid.UUID(order_id)
    except ValueError:
        raise HTTPException(status_code=404, detail="Invalid order ID")
    order = await order_service.get_admin_order_detail(db, oid)
    if not order:
        raise HTTPException(status_code=404, detail="Order not found")
    invoice = (await db.execute(
        select(Invoice).options(selectinload(Invoice.invoice_products))
        .where(Invoice.order_id == oid, Invoice.is_removed == False)
    )).scalar_one_or_none()
    data = order_service.build_admin_order_response(order)
    data["invoice"] = invoice
    return templates.TemplateResponse("admin/order_factor.html", {
        "request": request, "current_user": current_user,
        "order": data,
    })


@router.get("/orders/{order_id}/labels", response_class=HTMLResponse)
async def admin_order_labels(
    request: Request,
    order_id: str,
    current_user: User = Depends(require_any_role("Admin", "Orders Manager")),
    db: AsyncSession = Depends(get_db),
):
    """Shipping labels (برچسب‌ها) — mirrors .NET Labels action using the order's invoice."""
    try:
        oid = uuid.UUID(order_id)
    except ValueError:
        raise HTTPException(status_code=404, detail="Invalid order ID")
    order = await order_service.get_admin_order_detail(db, oid)
    if not order:
        raise HTTPException(status_code=404, detail="Order not found")
    invoice = (await db.execute(
        select(Invoice).options(selectinload(Invoice.invoice_products))
        .where(Invoice.order_id == oid, Invoice.is_removed == False)
    )).scalar_one_or_none()

    labels = []
    if invoice:
        for ip in invoice.invoice_products:
            if ip.type and ip.type != "Product":
                continue
            labels.append({
                "part_number": ip.part_number or "-",
                "name": ip.name or "",
                "model": ip.model or "",
                "count": ip.count,
                "product_unit": ip.product_unit or "",
                "variety": (ip.variety_value or "").strip(),
            })
    else:
        data = order_service.build_admin_order_response(order)
        for op in data["order_products"]:
            labels.append({
                "part_number": op["part_number"] or "-",
                "name": op["product_name"] or "",
                "model": "",
                "count": op["count"],
                "product_unit": "",
                "variety": (op.get("variety_values") or "").strip(),
            })

    return templates.TemplateResponse("admin/order_labels.html", {
        "request": request, "current_user": current_user,
        "reference_code": order.reference_code or order.id[:8],
        "labels": labels,
    })


@router.get("/orders/{order_id}/post-info", response_class=HTMLResponse)
async def admin_order_post_info(
    request: Request,
    order_id: str,
    current_user: User = Depends(require_any_role("Admin", "Orders Manager")),
    db: AsyncSession = Depends(get_db),
):
    """Post info page (نوع ارسال) — mirrors .NET PostInfo action."""
    try:
        oid = uuid.UUID(order_id)
    except ValueError:
        raise HTTPException(status_code=404, detail="Invalid order ID")
    order = await order_service.get_admin_order_detail(db, oid)
    if not order:
        raise HTTPException(status_code=404, detail="Order not found")
    site_setting = (await db.execute(
        select(SiteSetting).where(SiteSetting.is_removed == False).limit(1)
    )).scalar_one_or_none()
    data = order_service.build_admin_order_response(order)
    return templates.TemplateResponse("admin/order_post_info.html", {
        "request": request, "current_user": current_user,
        "order": data, "site_setting": site_setting,
    })


# ── Order Products ──

@router.get("/order-products/create/{order_id}", response_class=HTMLResponse)
async def admin_order_product_create_form(
    request: Request,
    order_id: str,
    current_user: User = Depends(require_any_role("Admin", "Orders Manager")),
    db: AsyncSession = Depends(get_db),
):
    try:
        oid = uuid.UUID(order_id)
    except ValueError:
        raise HTTPException(status_code=404, detail="Invalid order ID")
    order = await order_service.get_admin_order_detail(db, oid)
    if not order:
        raise HTTPException(status_code=404, detail="Order not found")
    products = (await db.execute(
        select(Product).options(selectinload(Product.varieties))
        .where(Product.is_removed == False).order_by(Product.name)
    )).scalars().all()
    return templates.TemplateResponse("admin/order_product_form.html", {
        "request": request, "current_user": current_user,
        "op": None, "order": order_service.build_admin_order_response(order),
        "products": order_service.product_picker_data(products),
    })


@router.post("/order-products/create", response_class=HTMLResponse)
async def admin_order_product_create_submit(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Orders Manager")),
    db: AsyncSession = Depends(get_db),
):
    form = await request.form()
    try:
        order_id = uuid.UUID(form.get("order_id"))
        product_id = uuid.UUID(form.get("product_id"))
    except (ValueError, TypeError):
        raise HTTPException(status_code=400, detail="Invalid order/product")
    variety_id_str = form.get("variety_id") or ""
    variety_id = None
    if variety_id_str:
        try:
            variety_id = uuid.UUID(variety_id_str)
        except ValueError:
            variety_id = None
    try:
        count = int(form.get("count") or 1)
    except ValueError:
        count = 1

    order = await db.get(Order, order_id)
    if not order:
        raise HTTPException(status_code=404, detail="Order not found")
    product = (await db.execute(
        select(Product).options(selectinload(Product.varieties)).where(Product.id == product_id)
    )).scalar_one_or_none()
    if not product:
        raise HTTPException(status_code=404, detail="Product not found")

    variety = None
    if variety_id:
        variety = await db.get(Variety, variety_id)
    if not variety and product.varieties:
        variety = product.varieties[0]
    if not variety:
        raise HTTPException(status_code=400, detail="Product has no variety")

    # prevent duplicate product/variety
    existing = (await db.execute(
        select(OrderProduct).where(
            OrderProduct.order_id == order_id,
            OrderProduct.product_id == product_id,
            OrderProduct.variety_id == variety.id,
            OrderProduct.is_removed == False,
        )
    )).scalar_one_or_none()
    if existing:
        raise HTTPException(status_code=400, detail="این کالا قبلاً به سفارش اضافه شده است")

    unit_price = float(variety.price or 0)
    discount = float(variety.discount_amount or 0)
    price_after_discount = float(variety.price_after_discount or unit_price)
    total_price = unit_price * count
    total_price_after_discount = (price_after_discount * count) if discount != 0 else (unit_price * count)

    op = OrderProduct(
        id=uuid.uuid4(),
        order_id=order_id,
        product_id=product_id,
        variety_id=variety.id,
        count=count,
        unit_price=unit_price,
        discount=discount,
        price_after_discount=price_after_discount,
        total_price=total_price,
        total_price_after_discount=total_price_after_discount,
        vat_rate=float(product.vat_rate or 0),
        created_by_user_id=current_user.id,
        insert_date=datetime.now(timezone.utc),
        update_date=datetime.now(timezone.utc),
    )
    db.add(op)
    await db.flush()

    # refresh order totals (mirror .NET Refresh + SetPayable)
    order_products = (await db.execute(
        select(OrderProduct).where(OrderProduct.order_id == order_id, OrderProduct.is_removed == False)
    )).scalars().all()
    order.total_price = sum(float(p.unit_price or 0) * p.count for p in order_products)
    order.total_price_after_discount = sum(float(p.total_price_after_discount or 0) for p in order_products)
    order.count = sum(p.count for p in order_products)
    order.vat = sum(float(p.price_after_discount or 0) * p.count * float(p.vat_rate or 0) / 100 for p in order_products)
    order.payable = order.total_price_after_discount + order.vat + float(order.postage_fee or 0) + float(order.packaging_cost or 0)
    order.update_date = datetime.now(timezone.utc)

    db.add(Log(record_id=op.id, table_name="order_products",
               description=f"افزودن کالا {product.name} به سفارش {order.reference_code}",
               created_by_user_id=current_user.id, type="Create"))
    await db.commit()
    return RedirectResponse(url=f"/administration/orders/{order_id}", status_code=303)


@router.get("/order-products/{op_id}", response_class=HTMLResponse)
async def admin_order_product_detail(
    request: Request,
    op_id: str,
    current_user: User = Depends(require_any_role("Admin", "Orders Manager")),
    db: AsyncSession = Depends(get_db),
):
    try:
        oid = uuid.UUID(op_id)
    except ValueError:
        raise HTTPException(status_code=404, detail="Invalid ID")
    op = (await db.execute(
        select(OrderProduct).options(selectinload(OrderProduct.product), selectinload(OrderProduct.variety))
        .where(OrderProduct.id == oid, OrderProduct.is_removed == False)
    )).unique().scalar_one_or_none()
    if not op:
        raise HTTPException(status_code=404, detail="Order product not found")
    return templates.TemplateResponse("admin/order_product_detail.html", {
        "request": request, "current_user": current_user,
        "op": order_service.order_product_detail(op),
    })


@router.get("/order-products/{op_id}/edit", response_class=HTMLResponse)
async def admin_order_product_edit_form(
    request: Request,
    op_id: str,
    current_user: User = Depends(require_any_role("Admin", "Orders Manager")),
    db: AsyncSession = Depends(get_db),
):
    try:
        oid = uuid.UUID(op_id)
    except ValueError:
        raise HTTPException(status_code=404, detail="Invalid ID")
    op = (await db.execute(
        select(OrderProduct).options(selectinload(OrderProduct.product), selectinload(OrderProduct.variety))
        .where(OrderProduct.id == oid, OrderProduct.is_removed == False)
    )).unique().scalar_one_or_none()
    if not op:
        raise HTTPException(status_code=404, detail="Order product not found")
    products = (await db.execute(
        select(Product).options(selectinload(Product.varieties))
        .where(Product.is_removed == False).order_by(Product.name)
    )).scalars().all()
    order = await db.get(Order, op.order_id)
    return templates.TemplateResponse("admin/order_product_form.html", {
        "request": request, "current_user": current_user,
        "op": order_service.order_product_detail(op),
        "order": order_service.build_admin_order_response(order) if order else None,
        "products": order_service.product_picker_data(products),
    })


@router.post("/order-products/{op_id}/edit", response_class=HTMLResponse)
async def admin_order_product_edit_submit(
    request: Request,
    op_id: str,
    current_user: User = Depends(require_any_role("Admin", "Orders Manager")),
    db: AsyncSession = Depends(get_db),
):
    try:
        oid = uuid.UUID(op_id)
    except ValueError:
        raise HTTPException(status_code=404, detail="Invalid ID")
    op = (await db.execute(
        select(OrderProduct).where(OrderProduct.id == oid, OrderProduct.is_removed == False)
    )).scalar_one_or_none()
    if not op:
        raise HTTPException(status_code=404, detail="Order product not found")
    form = await request.form()
    try:
        product_id = uuid.UUID(form.get("product_id"))
        count = int(form.get("count") or op.count)
    except (ValueError, TypeError):
        product_id = op.product_id
        count = op.count
    variety_id_str = form.get("variety_id") or ""
    variety_id = None
    if variety_id_str:
        try:
            variety_id = uuid.UUID(variety_id_str)
        except ValueError:
            variety_id = None

    product = (await db.execute(
        select(Product).options(selectinload(Product.varieties)).where(Product.id == product_id)
    )).scalar_one_or_none()
    if not product:
        raise HTTPException(status_code=404, detail="Product not found")
    variety = None
    if variety_id:
        variety = await db.get(Variety, variety_id)
    if not variety and product.varieties:
        variety = product.varieties[0]

    unit_price = float(variety.price or 0) if variety else float(op.unit_price or 0)
    discount = float(variety.discount_amount or 0) if variety else float(op.discount or 0)
    price_after_discount = float(variety.price_after_discount or unit_price) if variety else float(op.price_after_discount or unit_price)

    op.product_id = product_id
    op.variety_id = variety.id if variety else None
    op.count = count
    op.unit_price = unit_price
    op.discount = discount
    op.price_after_discount = price_after_discount
    op.total_price = unit_price * count
    op.total_price_after_discount = (price_after_discount * count) if discount != 0 else (unit_price * count)
    op.vat_rate = float(product.vat_rate or 0)
    op.update_date = datetime.now(timezone.utc)

    order = await db.get(Order, op.order_id)
    if order:
        order_products = (await db.execute(
            select(OrderProduct).where(OrderProduct.order_id == order.id, OrderProduct.is_removed == False)
        )).scalars().all()
        order.total_price = sum(float(p.unit_price or 0) * p.count for p in order_products)
        order.total_price_after_discount = sum(float(p.total_price_after_discount or 0) for p in order_products)
        order.count = sum(p.count for p in order_products)
        order.vat = sum(float(p.price_after_discount or 0) * p.count * float(p.vat_rate or 0) / 100 for p in order_products)
        order.payable = order.total_price_after_discount + order.vat + float(order.postage_fee or 0) + float(order.packaging_cost or 0)
        order.update_date = datetime.now(timezone.utc)

    db.add(Log(record_id=op.id, table_name="order_products",
               description=f"ویرایش کالا در سفارش",
               created_by_user_id=current_user.id, type="Update"))
    await db.commit()
    return RedirectResponse(url=f"/administration/orders/{op.order_id}", status_code=303)


@router.post("/order-products/{op_id}/delete", response_class=HTMLResponse)
async def admin_order_product_delete(
    request: Request,
    op_id: str,
    current_user: User = Depends(require_any_role("Admin", "Orders Manager")),
    db: AsyncSession = Depends(get_db),
):
    try:
        oid = uuid.UUID(op_id)
    except ValueError:
        raise HTTPException(status_code=404, detail="Invalid ID")
    op = await db.get(OrderProduct, oid)
    if not op:
        raise HTTPException(status_code=404, detail="Order product not found")
    order_id = op.order_id
    op.is_removed = True
    op.update_date = datetime.now(timezone.utc)

    order = await db.get(Order, order_id)
    if order:
        order_products = (await db.execute(
            select(OrderProduct).where(OrderProduct.order_id == order_id, OrderProduct.is_removed == False)
        )).scalars().all()
        order.total_price = sum(float(p.unit_price or 0) * p.count for p in order_products)
        order.total_price_after_discount = sum(float(p.total_price_after_discount or 0) for p in order_products)
        order.count = sum(p.count for p in order_products)
        order.vat = sum(float(p.price_after_discount or 0) * p.count * float(p.vat_rate or 0) / 100 for p in order_products)
        order.payable = order.total_price_after_discount + order.vat + float(order.postage_fee or 0) + float(order.packaging_cost or 0)
        order.update_date = datetime.now(timezone.utc)

    db.add(Log(record_id=op.id, table_name="order_products",
               description="حذف کالا از سفارش",
               created_by_user_id=current_user.id, type="Delete"))
    await db.commit()
    return RedirectResponse(url=f"/administration/orders/{order_id}", status_code=303)


# ── Payment Requests ──

@router.get("/payment-requests/create/{order_id}", response_class=HTMLResponse)
async def admin_payment_request_create_form(
    request: Request,
    order_id: str,
    current_user: User = Depends(require_any_role("Admin", "Financial Manager")),
    db: AsyncSession = Depends(get_db),
):
    try:
        oid = uuid.UUID(order_id)
    except ValueError:
        raise HTTPException(status_code=404, detail="Invalid order ID")
    order = await order_service.get_admin_order_detail(db, oid)
    if not order:
        raise HTTPException(status_code=404, detail="Order not found")
    return templates.TemplateResponse("admin/payment_request_form.html", {
        "request": request, "current_user": current_user,
        "pr": None, "order": order_service.build_admin_order_response(order),
        "payment_status_names": order_service.PAYMENT_STATUS_NAMES,
        "default_status": "Pending",
    })


@router.post("/payment-requests/create", response_class=HTMLResponse)
async def admin_payment_request_create_submit(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Financial Manager")),
    db: AsyncSession = Depends(get_db),
):
    form = await request.form()
    try:
        order_id = uuid.UUID(form.get("order_id"))
    except (ValueError, TypeError):
        raise HTTPException(status_code=400, detail="Invalid order")
    order = await db.get(Order, order_id)
    if not order:
        raise HTTPException(status_code=404, detail="Order not found")
    try:
        amount = float(form.get("amount") or 0)
    except ValueError:
        amount = 0
    status = (form.get("status") or "Pending").strip()
    pr = PaymentRequest(
        id=uuid.uuid4(),
        order_id=order_id,
        user_id=order.user_id,
        amount=amount,
        is_pay=False,
        status=status,
        created_by_user_id=current_user.id,
        insert_date=datetime.now(timezone.utc),
        update_date=datetime.now(timezone.utc),
    )
    db.add(pr)
    db.add(Log(record_id=pr.id, table_name="payment_requests",
               description=f"ایجاد پرداخت برای سفارش {order.reference_code}: {amount}",
               created_by_user_id=current_user.id, type="Create"))
    await db.commit()
    return RedirectResponse(url=f"/administration/orders/{order_id}", status_code=303)


@router.get("/payment-requests/{pr_id}", response_class=HTMLResponse)
async def admin_payment_request_detail(
    request: Request,
    pr_id: str,
    current_user: User = Depends(require_any_role("Admin", "Financial Manager")),
    db: AsyncSession = Depends(get_db),
):
    try:
        pid = uuid.UUID(pr_id)
    except ValueError:
        raise HTTPException(status_code=404, detail="Invalid ID")
    pr = await db.get(PaymentRequest, pid)
    if not pr:
        raise HTTPException(status_code=404, detail="Payment request not found")
    return templates.TemplateResponse("admin/payment_request_detail.html", {
        "request": request, "current_user": current_user,
        "pr": {
            "id": pr.id,
            "amount": float(pr.amount or 0),
            "pay_date": pr.pay_date,
            "pay_date_str": to_farsi_full(pr.pay_date) if pr.pay_date else "",
            "approval": pr.approval,
            "approval_str": to_farsi_full(pr.approval) if pr.approval else "",
            "status": pr.status,
            "status_name": order_service.PAYMENT_STATUS_NAMES.get(pr.status, pr.status or "-"),
            "is_pay": pr.is_pay,
            "ref_id": pr.ref_id,
            "authority": pr.authority,
            "card_pan": pr.card_pan,
            "order_id": pr.order_id,
            "user_id": pr.user_id,
        },
    })


@router.get("/payment-requests/{pr_id}/edit", response_class=HTMLResponse)
async def admin_payment_request_edit_form(
    request: Request,
    pr_id: str,
    current_user: User = Depends(require_any_role("Admin", "Financial Manager")),
    db: AsyncSession = Depends(get_db),
):
    try:
        pid = uuid.UUID(pr_id)
    except ValueError:
        raise HTTPException(status_code=404, detail="Invalid ID")
    pr = await db.get(PaymentRequest, pid)
    if not pr:
        raise HTTPException(status_code=404, detail="Payment request not found")
    return templates.TemplateResponse("admin/payment_request_form.html", {
        "request": request, "current_user": current_user,
        "pr": {
            "id": pr.id,
            "amount": float(pr.amount or 0),
            "status": pr.status,
            "status_name": order_service.PAYMENT_STATUS_NAMES.get(pr.status, pr.status or "-"),
            "is_pay": pr.is_pay,
            "order_id": pr.order_id,
        },
        "payment_status_names": order_service.PAYMENT_STATUS_NAMES,
    })


@router.post("/payment-requests/{pr_id}/edit", response_class=HTMLResponse)
async def admin_payment_request_edit_submit(
    request: Request,
    pr_id: str,
    current_user: User = Depends(require_any_role("Admin", "Financial Manager")),
    db: AsyncSession = Depends(get_db),
):
    try:
        pid = uuid.UUID(pr_id)
    except ValueError:
        raise HTTPException(status_code=404, detail="Invalid ID")
    pr = await db.get(PaymentRequest, pid)
    if not pr:
        raise HTTPException(status_code=404, detail="Payment request not found")
    form = await request.form()
    status = form.get("status")
    if status:
        pr.status = status
    try:
        amount = float(form.get("amount"))
        if amount >= 0:
            pr.amount = amount
    except (ValueError, TypeError):
        pass
    pr.update_date = datetime.now(timezone.utc)
    db.add(Log(record_id=pr.id, table_name="payment_requests",
               description=f"ویرایش پرداخت: {pr.ref_id}",
               created_by_user_id=current_user.id, type="Update"))
    await db.commit()
    return RedirectResponse(url=f"/administration/payment-requests/{pr.id}", status_code=303)


@router.post("/payment-requests/{pr_id}/delete", response_class=HTMLResponse)
async def admin_payment_request_delete(
    request: Request,
    pr_id: str,
    current_user: User = Depends(require_any_role("Admin", "Financial Manager")),
    db: AsyncSession = Depends(get_db),
):
    try:
        pid = uuid.UUID(pr_id)
    except ValueError:
        raise HTTPException(status_code=404, detail="Invalid ID")
    pr = await db.get(PaymentRequest, pid)
    if not pr:
        raise HTTPException(status_code=404, detail="Payment request not found")
    order_id = pr.order_id
    pr.is_removed = True
    pr.update_date = datetime.now(timezone.utc)
    db.add(Log(record_id=pr.id, table_name="payment_requests",
               description="حذف پرداخت",
               created_by_user_id=current_user.id, type="Delete"))
    await db.commit()
    return RedirectResponse(url=f"/administration/orders/{order_id}", status_code=303)


# ── Receipts (order scoped) ──

@router.get("/receipts/create/{order_id}", response_class=HTMLResponse)
async def admin_receipt_create_form(
    request: Request,
    order_id: str,
    current_user: User = Depends(require_any_role("Admin", "Financial Manager")),
    db: AsyncSession = Depends(get_db),
):
    try:
        oid = uuid.UUID(order_id)
    except ValueError:
        raise HTTPException(status_code=404, detail="Invalid order ID")
    order = await order_service.get_admin_order_detail(db, oid)
    if not order:
        raise HTTPException(status_code=404, detail="Order not found")
    users = (await db.execute(
        select(User).where(User.is_removed == False).order_by(User.first_name)
    )).scalars().all()
    return templates.TemplateResponse("admin/receipt_form.html", {
        "request": request, "current_user": current_user,
        "order": order_service.build_admin_order_response(order),
        "users": users,
    })


@router.post("/receipts/create", response_class=HTMLResponse)
async def admin_receipt_create_submit(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Financial Manager")),
    db: AsyncSession = Depends(get_db),
):
    form = await request.form()
    try:
        order_id = uuid.UUID(form.get("order_id"))
    except (ValueError, TypeError):
        raise HTTPException(status_code=400, detail="Invalid order")
    order = await db.get(Order, order_id)
    if not order:
        raise HTTPException(status_code=404, detail="Order not found")
    user_id_str = form.get("user_id") or ""
    user_id = None
    if user_id_str:
        try:
            user_id = uuid.UUID(user_id_str)
        except ValueError:
            user_id = None
    try:
        price = float(form.get("price") or 0)
    except ValueError:
        price = 0
    description = (form.get("description") or "").strip() or None
    destination_bank = (form.get("destination_bank") or "").strip() or None
    reference_code = (form.get("reference_code") or "").strip() or None

    receipt = Receipt(
        id=uuid.uuid4(),
        order_id=order_id,
        user_id=user_id,
        price=price,
        description=description,
        destination_bank=destination_bank,
        reference_code=reference_code or str(order.reference_code),
        status="AwaitingConfirmation",
        created_by_user_id=current_user.id,
        insert_date=datetime.now(timezone.utc),
        update_date=datetime.now(timezone.utc),
    )
    db.add(receipt)
    db.add(Log(record_id=receipt.id, table_name="receipts",
               description=f"ایجاد فیش پرداخت برای سفارش {order.reference_code}: {price}",
               created_by_user_id=current_user.id, type="Create"))
    await db.commit()
    return RedirectResponse(url=f"/administration/orders/{order_id}", status_code=303)


# ── Order Status Records ──

@router.get("/order-status-records/{record_id}", response_class=HTMLResponse)
async def admin_order_status_record_detail(
    request: Request,
    record_id: str,
    current_user: User = Depends(require_any_role("Admin", "Orders Manager")),
    db: AsyncSession = Depends(get_db),
):
    try:
        rid = uuid.UUID(record_id)
    except ValueError:
        raise HTTPException(status_code=404, detail="Invalid ID")
    rec = (await db.execute(
        select(OrderStatusRecord).where(OrderStatusRecord.id == rid, OrderStatusRecord.is_removed == False)
    )).scalar_one_or_none()
    if not rec:
        raise HTTPException(status_code=404, detail="Record not found")
    created_by = None
    if rec.created_by_user_id:
        u = await db.get(User, rec.created_by_user_id)
        created_by = u.full_name if u else None
    return templates.TemplateResponse("admin/order_status_record_detail.html", {
        "request": request, "current_user": current_user,
        "rec": {
            "id": rec.id,
            "order_id": rec.order_id,
            "status": rec.status,
            "status_name": order_service.ORDER_STATUS_NAMES.get(rec.status, rec.status or "-"),
            "comment": rec.comment,
            "tracking_number": rec.tracking_number,
            "insert_date": rec.insert_date,
            "insert_date_str": to_farsi_full(rec.insert_date) if rec.insert_date else "",
            "created_by": created_by,
        },
    })


@router.get("/order-status-records/{record_id}/edit", response_class=HTMLResponse)
async def admin_order_status_record_edit_form(
    request: Request,
    record_id: str,
    current_user: User = Depends(require_any_role("Admin", "Orders Manager")),
    db: AsyncSession = Depends(get_db),
):
    try:
        rid = uuid.UUID(record_id)
    except ValueError:
        raise HTTPException(status_code=404, detail="Invalid ID")
    rec = (await db.execute(
        select(OrderStatusRecord).where(OrderStatusRecord.id == rid, OrderStatusRecord.is_removed == False)
    )).scalar_one_or_none()
    if not rec:
        raise HTTPException(status_code=404, detail="Record not found")
    return templates.TemplateResponse("admin/order_status_record_form.html", {
        "request": request, "current_user": current_user,
        "rec": {
            "id": rec.id,
            "order_id": rec.order_id,
            "status": rec.status,
            "comment": rec.comment,
        },
        "order_status_names": order_service.ORDER_STATUS_NAMES,
    })


@router.post("/order-status-records/{record_id}/edit", response_class=HTMLResponse)
async def admin_order_status_record_edit_submit(
    request: Request,
    record_id: str,
    current_user: User = Depends(require_any_role("Admin", "Orders Manager")),
    db: AsyncSession = Depends(get_db),
):
    try:
        rid = uuid.UUID(record_id)
    except ValueError:
        raise HTTPException(status_code=404, detail="Invalid ID")
    rec = (await db.execute(
        select(OrderStatusRecord).where(OrderStatusRecord.id == rid, OrderStatusRecord.is_removed == False)
    )).scalar_one_or_none()
    if not rec:
        raise HTTPException(status_code=404, detail="Record not found")
    form = await request.form()
    rec.comment = (form.get("comment") or "").strip() or None
    new_status = form.get("status")
    if new_status:
        rec.status = new_status
    rec.update_date = datetime.now(timezone.utc)
    db.add(Log(record_id=rec.id, table_name="order_status_records",
               description="ویرایش سابقه وضعیت سفارش",
               created_by_user_id=current_user.id, type="Update"))
    await db.commit()
    return RedirectResponse(url=f"/administration/order-status-records/{rec.id}", status_code=303)


# ── Users ──

@router.get("/users", response_class=HTMLResponse)
async def admin_users(
    request: Request,
    page: int = Query(1),
    search: str = Query(""),
    current_user: User = Depends(get_admin_user),
    db: AsyncSession = Depends(get_db),
):
    users, total = await identity_service.get_users_paginated(db, page, 20, search)
    roles = (await db.execute(select(Role).where(Role.is_removed == False))).scalars().all()
    return templates.TemplateResponse("admin/users.html", {
        "request": request, "current_user": current_user,
        "users": users, "total": total, "page": page, "total_pages": (total + 19) // 20,
        "search": search, "roles": roles,
    })


@router.get("/users/new", response_class=HTMLResponse)
async def admin_user_create(
    request: Request,
    current_user: User = Depends(get_admin_user),
    db: AsyncSession = Depends(get_db),
):
    return templates.TemplateResponse("admin/user_form.html", {
        "request": request, "current_user": current_user, "user": None,
    })


@router.post("/users/new", response_class=HTMLResponse)
async def admin_user_create_submit(
    request: Request,
    current_user: User = Depends(get_admin_user),
    db: AsyncSession = Depends(get_db),
):
    form = await request.form()
    try:
        user = await identity_service.create_user(db, dict(form), current_user.id)
        return RedirectResponse(url="/administration/users", status_code=303)
    except ValueError as e:
        return templates.TemplateResponse("admin/user_form.html", {
            "request": request, "current_user": current_user, "user": None, "error": str(e),
        })


@router.get("/users/{user_id}/edit", response_class=HTMLResponse)
async def admin_user_edit(
    request: Request,
    user_id: str,
    current_user: User = Depends(get_admin_user),
    db: AsyncSession = Depends(get_db),
):
    uid = uuid.UUID(user_id)
    user = await identity_service.get_user_by_id(db, uid)
    if not user:
        raise HTTPException(status_code=404, detail="User not found")
    return templates.TemplateResponse("admin/user_form.html", {
        "request": request, "current_user": current_user, "user": user,
    })


@router.post("/users/{user_id}/edit", response_class=HTMLResponse)
async def admin_user_edit_submit(
    request: Request,
    user_id: str,
    current_user: User = Depends(get_admin_user),
    db: AsyncSession = Depends(get_db),
):
    uid = uuid.UUID(user_id)
    user = await identity_service.get_user_by_id(db, uid)
    if not user:
        raise HTTPException(status_code=404, detail="User not found")
    form = await request.form()
    try:
        await identity_service.update_user(db, user, dict(form), current_user.id)
        return RedirectResponse(url="/administration/users", status_code=303)
    except ValueError as e:
        return templates.TemplateResponse("admin/user_form.html", {
            "request": request, "current_user": current_user, "user": user, "error": str(e),
        })


@router.get("/users/{user_id}", response_class=HTMLResponse)
async def admin_user_detail(
    request: Request,
    user_id: str,
    page: int = Query(1),
    current_user: User = Depends(get_admin_user),
    db: AsyncSession = Depends(get_db),
):
    from app.models.common import Address, BankInfo
    from app.models.order import OrderModel as Order
    uid = uuid.UUID(user_id)
    user = await identity_service.get_user_by_id(db, uid)
    if not user:
        raise HTTPException(status_code=404, detail="User not found")

    # Load addresses
    addr_page = address_page = page
    addr_stmt = (
        select(Address)
        .options(selectinload(Address.province_city))
        .where(Address.user_id == uid, Address.is_removed == False)
        .order_by(Address.insert_date.desc())
    )
    addr_count = (await db.execute(select(func.count(Address.id)).where(Address.user_id == uid, Address.is_removed == False))).scalar() or 0
    addresses = (await db.execute(addr_stmt.offset(0).limit(100))).unique().scalars().all()

    # Load bank infos
    bank_stmt = (
        select(BankInfo)
        .where(BankInfo.user_id == uid, BankInfo.is_removed == False)
        .order_by(BankInfo.insert_date.desc())
    )
    bank_count = (await db.execute(select(func.count(BankInfo.id)).where(BankInfo.user_id == uid, BankInfo.is_removed == False))).scalar() or 0
    bank_infos = (await db.execute(bank_stmt.offset(0).limit(100))).scalars().all()

    # Load orders
    order_stmt = (
        select(Order)
        .where(Order.user_id == uid, Order.is_removed == False)
        .order_by(Order.insert_date.desc())
    )
    order_count = (await db.execute(select(func.count(Order.id)).where(Order.user_id == uid, Order.is_removed == False))).scalar() or 0
    orders = (await db.execute(order_stmt.offset(0).limit(100))).scalars().all()

    roles = (await db.execute(select(Role).where(Role.is_removed == False))).scalars().all()
    return templates.TemplateResponse("admin/user_detail.html", {
        "request": request, "current_user": current_user,
        "user": user, "roles": roles,
        "addresses": addresses, "addr_count": addr_count, "addr_page": 1, "addr_total_pages": (addr_count + 99) // 100,
        "bank_infos": bank_infos, "bank_count": bank_count, "bank_page": 1, "bank_total_pages": (bank_count + 99) // 100,
        "orders": orders, "order_count": order_count, "order_page": 1, "order_total_pages": (order_count + 99) // 100,
    })


@router.post("/users/{user_id}/delete", response_class=HTMLResponse)
async def admin_user_delete(
    request: Request,
    user_id: str,
    current_user: User = Depends(get_admin_user),
    db: AsyncSession = Depends(get_db),
):
    uid = uuid.UUID(user_id)
    user = await identity_service.get_user_by_id(db, uid)
    if not user:
        raise HTTPException(status_code=404, detail="User not found")
    await identity_service.soft_delete_user(db, user, current_user.id)
    return RedirectResponse(url="/administration/users", status_code=303)


@router.post("/users/{user_id}/roles", response_class=HTMLResponse)
async def admin_user_assign_role(
    request: Request,
    user_id: str,
    current_user: User = Depends(get_admin_user),
    db: AsyncSession = Depends(get_db),
):
    form = await request.form()
    role_name = form.get("role_name", "")
    try:
        await identity_service.assign_role_to_user(db, uuid.UUID(user_id), role_name, current_user.id)
    except ValueError as e:
        return RedirectResponse(url=f"/administration/users/{user_id}", status_code=303)
    return RedirectResponse(url=f"/administration/users/{user_id}", status_code=303)


# ── Roles ──

@router.get("/roles", response_class=HTMLResponse)
async def admin_roles(
    request: Request,
    current_user: User = Depends(get_admin_user),
    db: AsyncSession = Depends(get_db),
):
    items = await identity_service.get_roles_with_counts(db)
    return templates.TemplateResponse("admin/roles.html", {
        "request": request, "current_user": current_user, "roles": items,
    })


@router.get("/roles/create", response_class=HTMLResponse)
async def admin_role_create(
    request: Request,
    current_user: User = Depends(get_admin_user),
    db: AsyncSession = Depends(get_db),
):
    return templates.TemplateResponse("admin/role_form.html", {
        "request": request, "current_user": current_user, "role": None,
    })


@router.post("/roles/create", response_class=HTMLResponse)
async def admin_role_create_submit(
    request: Request,
    current_user: User = Depends(get_admin_user),
    db: AsyncSession = Depends(get_db),
):
    form = await request.form()
    try:
        role = await identity_service.create_role(db, dict(form), current_user.id)
        return RedirectResponse(url="/administration/roles", status_code=303)
    except ValueError as e:
        return templates.TemplateResponse("admin/role_form.html", {
            "request": request, "current_user": current_user, "role": None, "error": str(e),
        })


@router.get("/roles/{role_id}/edit", response_class=HTMLResponse)
async def admin_role_edit(
    request: Request,
    role_id: str,
    current_user: User = Depends(get_admin_user),
    db: AsyncSession = Depends(get_db),
):
    role = await identity_service.get_role_by_id(db, uuid.UUID(role_id))
    if not role or role.is_removed:
        raise HTTPException(status_code=404, detail="Role not found")
    return templates.TemplateResponse("admin/role_form.html", {
        "request": request, "current_user": current_user, "role": role,
    })


@router.post("/roles/{role_id}/edit", response_class=HTMLResponse)
async def admin_role_edit_submit(
    request: Request,
    role_id: str,
    current_user: User = Depends(get_admin_user),
    db: AsyncSession = Depends(get_db),
):
    role = await identity_service.get_role_by_id(db, uuid.UUID(role_id))
    if not role or role.is_removed:
        raise HTTPException(status_code=404, detail="Role not found")
    form = await request.form()
    try:
        await identity_service.update_role(db, role, dict(form), current_user.id)
        return RedirectResponse(url="/administration/roles", status_code=303)
    except ValueError as e:
        return templates.TemplateResponse("admin/role_form.html", {
            "request": request, "current_user": current_user, "role": role, "error": str(e),
        })


@router.get("/roles/{role_id}", response_class=HTMLResponse)
async def admin_role_detail(
    request: Request,
    role_id: str,
    current_user: User = Depends(get_admin_user),
    db: AsyncSession = Depends(get_db),
):
    role = await identity_service.get_role_by_id(db, uuid.UUID(role_id))
    if not role or role.is_removed:
        raise HTTPException(status_code=404, detail="Role not found")
    return templates.TemplateResponse("admin/role_detail.html", {
        "request": request, "current_user": current_user, "role": role,
    })


@router.post("/roles/{role_id}/delete", response_class=HTMLResponse)
async def admin_role_delete(
    request: Request,
    role_id: str,
    current_user: User = Depends(get_admin_user),
    db: AsyncSession = Depends(get_db),
):
    role = await identity_service.get_role_by_id(db, uuid.UUID(role_id))
    if not role or role.is_removed:
        raise HTTPException(status_code=404, detail="Role not found")
    await identity_service.soft_delete_role(db, role, current_user.id)
    return RedirectResponse(url="/administration/roles", status_code=303)


# ── Comments ──

@router.get("/comments", response_class=HTMLResponse)
async def admin_comments(
    request: Request,
    pending_only: bool = Query(False),
    current_user: User = Depends(get_admin_user),
    db: AsyncSession = Depends(get_db),
):
    stmt = select(Comment).where(Comment.is_removed == False).order_by(Comment.insert_date.desc()).limit(50)
    if pending_only:
        stmt = stmt.where(Comment.is_confirmed == False)
    comments = (await db.execute(stmt)).scalars().all()
    return templates.TemplateResponse("admin/comments.html", {
        "request": request, "current_user": current_user, "comments": comments,
    })


# ── Settings (SiteSettings) ──

def _price_str(value) -> str:
    if value is None:
        return ""
    try:
        return f"{float(value):,.0f}"
    except (TypeError, ValueError):
        return str(value)


def _parse_form_bool(v) -> bool:
    return v in ("on", "true", "True", "1")


def _parse_form_uuid(v):
    v = (v or "").strip()
    if not v:
        return None
    try:
        return uuid.UUID(v)
    except ValueError:
        return None


def _parse_form_float(v, default=0.0) -> float:
    v = (v or "").strip().replace(",", "")
    if not v:
        return default
    try:
        return float(v)
    except ValueError:
        return default


def _apply_site_settings(s: SiteSetting, form) -> str | None:
    """Apply form values onto a SiteSetting. Returns an error message, or None."""
    s.bank_name = form.get("BankName") or None
    s.account_number = form.get("AccountNumber") or None
    s.card_number = form.get("CardNumber") or None
    s.sheba_number = form.get("ShebaNumber") or None
    s.account_owner = form.get("AccountOwner") or None
    s.how_to_buy = form.get("HowToBuy") or None
    s.free_delivery = form.get("FreeDelivery") or None
    s.contact_us = form.get("ContactUs") or None
    s.technical_support = form.get("TechnicalSupport") or None
    s.email = form.get("Email") or None
    s.telephone = form.get("Telephone") or None
    s.address = form.get("Address") or None
    s.copy_right = form.get("CopyRight") or None
    s.disable_captcha = _parse_form_bool(form.get("DisableCaptcha"))
    s.free_packaging = _parse_form_bool(form.get("FreePackaging"))
    s.free_postage = _parse_form_bool(form.get("FreePostage"))
    s.top_category_id = _parse_form_uuid(form.get("TopCategoryId"))
    s.middle_category_id = _parse_form_uuid(form.get("MiddleCategoryId"))
    s.bottom_category_id = _parse_form_uuid(form.get("BottomCategoryId"))
    s.top_poster_category_id = _parse_form_uuid(form.get("TopPosterCategoryId"))
    s.mid_left_poster_category_id = _parse_form_uuid(form.get("MidLeftPosterCategoryId"))
    s.mid_right_poster_category_id = _parse_form_uuid(form.get("MidRightPosterCategoryId"))
    s.middle_poster_category_id = _parse_form_uuid(form.get("MiddlePosterCategoryId"))
    s.bottom_poster_category_id = _parse_form_uuid(form.get("BottomPosterCategoryId"))
    s.top_poster_image_url = (form.get("TopPosterImageUrl") or "").strip() or None
    s.middle_poster_image_url = (form.get("MiddlePosterImageUrl") or "").strip() or None
    s.mid_left_poster_image_url = (form.get("MidLeftPosterImageUrl") or "").strip() or None
    s.mid_right_poster_image_url = (form.get("MidRightPosterImageUrl") or "").strip() or None
    s.bottom_poster_image_url = (form.get("BottomPosterImageUrl") or "").strip() or None
    s.sidebar_support_category_id = _parse_form_uuid(form.get("SideBarSupportCategoryId"))
    s.sidebar_support_image_url = (form.get("SideBarSupportImageUrl") or "").strip() or None
    s.technical_table_id = _parse_form_uuid(form.get("TechnicalTableId"))
    free_limit_str = (form.get("FreePostageLimitStr") or "").strip().replace(",", "")
    if free_limit_str and not free_limit_str.replace(".", "", 1).isdigit():
        return "خطا در عملیات"
    s.free_postage_limit = float(free_limit_str) if free_limit_str else 0
    s.payment_status_per_hour = _parse_form_float(form.get("PaymentStatusPerHour"))
    s.postal_code = (form.get("PostalCode") or "").strip() or None
    return None


async def _get_site_setting(db) -> SiteSetting | None:
    return (await db.execute(
        select(SiteSetting).where(SiteSetting.is_removed == False).limit(1)
    )).scalar_one_or_none()


async def _site_setting_related(db, setting) -> dict | None:
    if setting is None:
        return None
    cat_ids = {setting.top_category_id, setting.middle_category_id, setting.bottom_category_id,
               setting.top_poster_category_id, setting.mid_left_poster_category_id,
               setting.mid_right_poster_category_id, setting.middle_poster_category_id,
               setting.bottom_poster_category_id, setting.sidebar_support_category_id}
    cat_ids = {c for c in cat_ids if c}
    names = {}
    if cat_ids:
        rows = (await db.execute(select(Category).where(Category.id.in_(cat_ids)))).scalars().all()
        names = {str(c.id): c.title or "" for c in rows}
    table_title = ""
    if setting.technical_table_id:
        tt = (await db.execute(select(TechnicalTable).where(TechnicalTable.id == setting.technical_table_id))).scalars().first()
        table_title = tt.title if tt else ""
    return {
        "top_category": names.get(str(setting.top_category_id), ""),
        "middle_category": names.get(str(setting.middle_category_id), ""),
        "bottom_category": names.get(str(setting.bottom_category_id), ""),
        "top_poster_category": names.get(str(setting.top_poster_category_id), ""),
        "mid_left_poster_category": names.get(str(setting.mid_left_poster_category_id), ""),
        "mid_right_poster_category": names.get(str(setting.mid_right_poster_category_id), ""),
        "middle_poster_category": names.get(str(setting.middle_poster_category_id), ""),
        "bottom_poster_category": names.get(str(setting.bottom_poster_category_id), ""),
        "sidebar_support_category": names.get(str(setting.sidebar_support_category_id), ""),
        "technical_table": table_title,
        "free_postage_limit_str": _price_str(setting.free_postage_limit),
    }


async def _settings_form_data(db):
    cats = (await db.execute(
        select(Category).where(Category.is_removed == False).order_by(Category.title)
    )).scalars().all()
    tables = (await db.execute(
        select(TechnicalTable).where(TechnicalTable.is_removed == False).order_by(TechnicalTable.title)
    )).scalars().all()
    return cats, tables


def _settings_form_values(s) -> dict:
    def uuid_str(v):
        return str(v) if v else ""
    empty = s is None
    if empty:
        s = SiteSetting()
    return {
        "BankName": s.bank_name or "",
        "AccountNumber": s.account_number or "",
        "CardNumber": s.card_number or "",
        "ShebaNumber": s.sheba_number or "",
        "AccountOwner": s.account_owner or "",
        "HowToBuy": s.how_to_buy or "",
        "TechnicalSupport": s.technical_support or "",
        "FreeDelivery": s.free_delivery or "",
        "ContactUs": s.contact_us or "",
        "Email": s.email or "",
        "Telephone": s.telephone or "",
        "Address": s.address or "",
        "CopyRight": s.copy_right or "",
        "PostalCode": s.postal_code or "",
        "DisableCaptcha": "on" if s.disable_captcha else "",
        "FreePackaging": "on" if s.free_packaging else "",
        "FreePostage": "on" if s.free_postage else "",
        "FreePostageLimitStr": _price_str(s.free_postage_limit),
        "PaymentStatusPerHour": str(s.payment_status_per_hour) if s.payment_status_per_hour is not None else "",
        "TopCategoryId": uuid_str(s.top_category_id),
        "MiddleCategoryId": uuid_str(s.middle_category_id),
        "BottomCategoryId": uuid_str(s.bottom_category_id),
        "TopPosterCategoryId": uuid_str(s.top_poster_category_id),
        "MidLeftPosterCategoryId": uuid_str(s.mid_left_poster_category_id),
        "MidRightPosterCategoryId": uuid_str(s.mid_right_poster_category_id),
        "MiddlePosterCategoryId": uuid_str(s.middle_poster_category_id),
        "BottomPosterCategoryId": uuid_str(s.bottom_poster_category_id),
        "TopPosterImageUrl": s.top_poster_image_url or "",
        "MiddlePosterImageUrl": s.middle_poster_image_url or "",
        "MidLeftPosterImageUrl": s.mid_left_poster_image_url or "",
        "MidRightPosterImageUrl": s.mid_right_poster_image_url or "",
        "BottomPosterImageUrl": s.bottom_poster_image_url or "",
        "SideBarSupportCategoryId": uuid_str(s.sidebar_support_category_id),
        "SideBarSupportImageUrl": s.sidebar_support_image_url or "",
        "TechnicalTableId": uuid_str(s.technical_table_id),
    }


# Fields on SiteSetting marked [Logged] in .NET (used to build create/update log descriptions)
_SITE_LOGGED_FIELDS = [
    ("top_category_id", "TopCategoryId"),
    ("middle_category_id", "MiddleCategoryId"),
    ("bottom_category_id", "BottomCategoryId"),
    ("top_poster_category_id", "TopPosterCategoryId"),
    ("mid_left_poster_category_id", "MidLeftPosterCategoryId"),
    ("mid_right_poster_category_id", "MidRightPosterCategoryId"),
    ("middle_poster_category_id", "MiddlePosterCategoryId"),
    ("bottom_poster_category_id", "BottomPosterCategoryId"),
    ("top_poster_image_url", "TopPosterImageUrl"),
    ("middle_poster_image_url", "MiddlePosterImageUrl"),
    ("mid_left_poster_image_url", "MidLeftPosterImageUrl"),
    ("mid_right_poster_image_url", "MidRightPosterImageUrl"),
    ("bottom_poster_image_url", "BottomPosterImageUrl"),
    ("sidebar_support_category_id", "SideBarSupportCategoryId"),
    ("sidebar_support_image_url", "SideBarSupportImageUrl"),
    ("bank_name", "BankName"),
    ("account_number", "AccountNumber"),
    ("card_number", "CardNumber"),
    ("sheba_number", "ShebaNumber"),
    ("account_owner", "AccountOwner"),
    ("how_to_buy", "HowToBuy"),
    ("free_delivery", "FreeDelivery"),
    ("contact_us", "ContactUs"),
    ("technical_support", "TechnicalSupport"),
    ("email", "Email"),
    ("telephone", "Telephone"),
    ("address", "Address"),
    ("copy_right", "CopyRight"),
    ("disable_captcha", "DisableCaptcha"),
    ("technical_table_id", "TechnicalTableId"),
    ("free_postage_limit", "FreePostageLimit"),
    ("postal_code", "PostalCode"),
]


def _log_val(v) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "True" if v else "False"
    return str(v)


def _site_logged_snapshot(s) -> dict:
    return {name: getattr(s, attr) for attr, name in _SITE_LOGGED_FIELDS}


def _site_create_desc(s) -> str:
    return "\n".join(f"{name}: {_log_val(getattr(s, attr))}" for attr, name in _SITE_LOGGED_FIELDS)


def _site_update_desc(before: dict, after: dict) -> str:
    lines = []
    for attr, name in _SITE_LOGGED_FIELDS:
        oldv = _log_val(before.get(name))
        newv = _log_val(after.get(name))
        if oldv != newv:
            lines.append(f"{name} : {oldv} --> {newv}")
    return "\n".join(lines)


@router.get("/settings", response_class=HTMLResponse)
async def admin_settings_details(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    setting = await _get_site_setting(db)
    related = await _site_setting_related(db, setting)
    return templates.TemplateResponse("admin/site_settings_details.html", {
        "request": request, "current_user": current_user,
        "settings": setting, "related": related,
    })


@router.get("/settings/create", response_class=HTMLResponse)
async def admin_settings_create_page(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    cats, tables = await _settings_form_data(db)
    return templates.TemplateResponse("admin/site_settings_form.html", {
        "request": request, "current_user": current_user,
        "settings": None, "cats": cats, "tables": tables,
        "action_url": "/administration/settings/create", "form_title": "ایجاد", "errors": [],
        "form": _settings_form_values(None),
    })


@router.post("/settings/create", response_class=HTMLResponse)
async def admin_settings_create(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    form = await request.form()
    setting = SiteSetting()
    err = _apply_site_settings(setting, form)
    if err:
        cats, tables = await _settings_form_data(db)
        return templates.TemplateResponse("admin/site_settings_form.html", {
            "request": request, "current_user": current_user,
            "settings": None, "cats": cats, "tables": tables,
            "action_url": "/administration/settings/create", "form_title": "ایجاد", "errors": [err],
            "form": form,
        })
    db.add(setting)
    db.add(Log(record_id=setting.id, table_name="site_settings",
               description=_site_create_desc(setting),
               created_by_user_id=current_user.id, type="Create"))
    await db.commit()
    return RedirectResponse(url="/administration/settings", status_code=303)


@router.get("/settings/{settings_id}/edit", response_class=HTMLResponse)
async def admin_settings_edit_page(
    settings_id: str,
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    try:
        sid = uuid.UUID(settings_id)
    except ValueError:
        raise HTTPException(status_code=404, detail="Invalid settings ID")
    setting = (await db.execute(
        select(SiteSetting).where(SiteSetting.id == sid, SiteSetting.is_removed == False)
    )).scalars().first()
    if not setting:
        raise HTTPException(status_code=404, detail="Site settings not found")
    cats, tables = await _settings_form_data(db)
    return templates.TemplateResponse("admin/site_settings_form.html", {
        "request": request, "current_user": current_user,
        "settings": setting, "cats": cats, "tables": tables,
        "action_url": f"/administration/settings/{settings_id}/edit", "form_title": "ویرایش", "errors": [],
        "form": _settings_form_values(setting),
    })


@router.post("/settings/{settings_id}/edit", response_class=HTMLResponse)
async def admin_settings_edit(
    settings_id: str,
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    try:
        sid = uuid.UUID(settings_id)
    except ValueError:
        raise HTTPException(status_code=404, detail="Invalid settings ID")
    setting = (await db.execute(
        select(SiteSetting).where(SiteSetting.id == sid, SiteSetting.is_removed == False)
    )).scalars().first()
    if not setting:
        raise HTTPException(status_code=404, detail="Site settings not found")
    form = await request.form()
    before = _site_logged_snapshot(setting)
    err = _apply_site_settings(setting, form)
    if err:
        cats, tables = await _settings_form_data(db)
        return templates.TemplateResponse("admin/site_settings_form.html", {
            "request": request, "current_user": current_user,
            "settings": setting, "cats": cats, "tables": tables,
            "action_url": f"/administration/settings/{settings_id}/edit", "form_title": "ویرایش",
            "errors": [err], "form": form,
        })
    update_desc = _site_update_desc(before, _site_logged_snapshot(setting))
    setting.update_date = datetime.now(timezone.utc)
    db.add(setting)
    if update_desc:
        db.add(Log(record_id=setting.id, table_name="site_settings",
                   description=update_desc,
                   created_by_user_id=current_user.id, type="Update"))
    await db.commit()
    return RedirectResponse(url="/administration/settings", status_code=303)


@router.post("/settings/{settings_id}/delete", response_class=HTMLResponse)
async def admin_settings_delete(
    settings_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    try:
        sid = uuid.UUID(settings_id)
    except ValueError:
        raise HTTPException(status_code=404, detail="Invalid settings ID")
    setting = (await db.execute(
        select(SiteSetting).where(SiteSetting.id == sid, SiteSetting.is_removed == False)
    )).scalars().first()
    if not setting:
        raise HTTPException(status_code=404, detail="Site settings not found")
    setting.is_removed = True
    setting.update_date = datetime.now(timezone.utc)
    db.add(Log(record_id=setting.id, table_name="site_settings",
               description="حذف تنظیمات سایت",
               created_by_user_id=current_user.id, type="Delete"))
    await db.commit()
    return RedirectResponse(url="/administration/settings", status_code=303)


@router.get("/settings/restore", response_class=HTMLResponse)
async def admin_settings_restore(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    setting = await _get_site_setting(db)
    if setting is None:
        db.add(SiteSetting(free_postage_limit=0, payment_status_per_hour=0))
    else:
        setting.update_date = datetime.now(timezone.utc)
    db.add(setting)
    db.add(Log(record_id=setting.id, table_name="site_settings",
               description=_site_create_desc(setting),
               created_by_user_id=current_user.id, type="Recovery"))
    await db.commit()
    return RedirectResponse(url="/administration/settings", status_code=303)


@router.get("/settings/update-products-by-variety", response_class=HTMLResponse)
async def admin_settings_update_products_by_variety(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    """Mirrors .NET SiteSettingsController.UpdateProductByVariety."""
    products = (await db.execute(
        select(Product).options(selectinload(Product.varieties)).where(Product.is_removed == False)
    )).scalars().all()
    for p in products:
        varieties = [v for v in (p.varieties or []) if not getattr(v, "is_removed", False)]
        p.number_of_variations = len(varieties)
        if not varieties:
            p.price = 0
            p.discount_amount = 0
            p.discount_percentage = 0
            p.stock_quantity = 0
            p.max_price = None
            db.add(p)
            continue
        min_v = min(varieties, key=lambda v: v.price_after_discount if v.price_after_discount is not None else float("inf"))
        p.price = min_v.price
        p.discount_amount = min_v.discount_amount
        p.discount_percentage = None
        p.stock_quantity = min_v.stock_quantity
        if len(varieties) > 1:
            max_v = max(varieties, key=lambda v: v.price_after_discount if v.price_after_discount is not None else float("-inf"))
            if (max_v.price_after_discount or 0) != (min_v.price_after_discount or 0):
                p.max_price = max_v.price_after_discount
            p.stock_quantity = sum(v.stock_quantity or 0 for v in varieties)
        db.add(p)
    await db.commit()
    return RedirectResponse(url="/administration/settings", status_code=303)


# ── Invoices ──

@router.get("/invoices", response_class=HTMLResponse)
async def admin_invoices(
    request: Request,
    page: int = Query(1),
    type_filter: str = Query(None),
    page_size: int = Query(20, ge=1, le=100),
    current_user: User = Depends(require_any_role("Admin", "Financial Manager")),
    db: AsyncSession = Depends(get_db),
):
    invoices, total = await invoice_service.get_all_invoices(db, page, page_size, type_filter)
    return templates.TemplateResponse("admin/invoices.html", {
        "request": request, "current_user": current_user,
        "invoices": [invoice_service.build_invoice_response(inv) for inv in invoices],
        "total": total, "page": page, "page_size": page_size,
        "total_pages": (total + page_size - 1) // page_size,
        "type_filter": type_filter,
    })


@router.get("/invoices/new", response_class=HTMLResponse)
async def admin_invoice_create_form(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Financial Manager")),
    db: AsyncSession = Depends(get_db),
):
    from app.models.product import Product as ProductModel
    users = (await db.execute(select(User).where(User.is_removed == False).order_by(User.first_name).limit(200))).scalars().all()
    products = (await db.execute(select(ProductModel).where(ProductModel.is_removed == False, ProductModel.no_display == False).limit(200))).scalars().all()
    return templates.TemplateResponse("admin/invoice_form.html", {
        "request": request, "current_user": current_user,
        "invoice": None, "users": users, "products": products,
    })


@router.post("/invoices/new", response_class=HTMLResponse)
async def admin_invoice_create_submit(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Financial Manager")),
    db: AsyncSession = Depends(get_db),
):
    from app.schemas.invoice import InvoiceCreate, InvoiceProductBase
    form = await request.form()
    try:
        invoice_products = []
        product_ids = form.getlist("product_ids[]")
        counts = form.getlist("counts[]")
        prices = form.getlist("prices[]")
        for i in range(len(product_ids)):
            if not product_ids[i].strip():
                continue
            ip = InvoiceProductBase(
                product_id=uuid.UUID(product_ids[i]),
                count=int(counts[i]) if i < len(counts) and counts[i] else 1,
                unit_price=float(prices[i]) if i < len(prices) and prices[i] else 0,
            )
            invoice_products.append(ip)
        if not invoice_products:
            raise ValueError("حداقل یک محصول الزامی است")
        create_data = InvoiceCreate(
            type=form.get("type", "Sale"),
            status=form.get("status", "Bought"),
            date=datetime.now(timezone.utc),
            description=form.get("description", ""),
            invoice_products=invoice_products,
        )
        invoice = await invoice_service.create_invoice(db, create_data, current_user.id)
        db.add(Log(record_id=invoice.id, table_name="invoices",
                   description=f"ایجاد فاکتور: {invoice.reference_code}",
                   created_by_user_id=current_user.id, type="Create"))
        await db.commit()
        return RedirectResponse(url=f"/administration/invoices/{invoice.id}", status_code=303)
    except (ValueError, TypeError) as e:
        users = (await db.execute(select(User).where(User.is_removed == False).order_by(User.first_name).limit(200))).scalars().all()
        products = (await db.execute(select(Product).where(Product.is_removed == False, Product.no_display == False).limit(200))).scalars().all()
        return templates.TemplateResponse("admin/invoice_form.html", {
            "request": request, "current_user": current_user,
            "invoice": None, "users": users, "products": products, "error": str(e),
        })


@router.get("/invoices/create-batch", response_class=HTMLResponse)
async def admin_invoice_create_batch_form(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Financial Manager")),
):
    return templates.TemplateResponse("admin/invoice_batch_form.html", {
        "request": request, "current_user": current_user, "batch_type": "invoice",
    })


@router.post("/invoices/create-batch", response_class=HTMLResponse)
async def admin_invoice_create_batch_submit(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Financial Manager")),
    db: AsyncSession = Depends(get_db),
):
    form = await request.form()
    file = request.files.get("file")
    if not file:
        return templates.TemplateResponse("admin/invoice_batch_form.html", {
            "request": request, "current_user": current_user,
            "batch_type": "invoice", "error": "فایل اکسل انتخاب شود",
        })
    db.add(Log(record_id=None, table_name="invoices",
               description=f"ایجاد دسته‌ای فاکتور توسط {current_user.user_name}",
               created_by_user_id=current_user.id, type="Batch"))
    await db.commit()
    return RedirectResponse(url="/administration/invoices", status_code=303)


@router.get("/invoices/create-batch-invoice-product", response_class=HTMLResponse)
async def admin_invoice_create_batch_product_form(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Financial Manager")),
):
    return templates.TemplateResponse("admin/invoice_batch_form.html", {
        "request": request, "current_user": current_user, "batch_type": "product",
    })


@router.post("/invoices/create-batch-invoice-product", response_class=HTMLResponse)
async def admin_invoice_create_batch_product_submit(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Financial Manager")),
    db: AsyncSession = Depends(get_db),
):
    form = await request.form()
    file = request.files.get("file")
    if not file:
        return templates.TemplateResponse("admin/invoice_batch_form.html", {
            "request": request, "current_user": current_user,
            "batch_type": "product", "error": "فایل اکسل انتخاب شود",
        })
    db.add(Log(record_id=None, table_name="invoice_products",
               description=f"ایجاد دسته‌ای اقلام فاکتور توسط {current_user.user_name}",
               created_by_user_id=current_user.id, type="Batch"))
    await db.commit()
    return RedirectResponse(url="/administration/invoices", status_code=303)


@router.get("/invoices/{invoice_id}", response_class=HTMLResponse)
async def admin_invoice_detail(
    request: Request,
    invoice_id: str,
    current_user: User = Depends(require_any_role("Admin", "Financial Manager")),
    db: AsyncSession = Depends(get_db),
):
    import uuid
    iid = uuid.UUID(invoice_id)
    invoice = await invoice_service.get_invoice_by_id(db, iid)
    if not invoice:
        raise HTTPException(status_code=404, detail="Invoice not found")
    return templates.TemplateResponse("admin/invoice_detail.html", {
        "request": request, "current_user": current_user,
        "invoice": invoice_service.build_invoice_response(invoice),
    })


@router.get("/invoices/{invoice_id}/edit", response_class=HTMLResponse)
async def admin_invoice_edit_form(
    request: Request,
    invoice_id: str,
    current_user: User = Depends(require_any_role("Admin", "Financial Manager")),
    db: AsyncSession = Depends(get_db),
):
    iid = uuid.UUID(invoice_id)
    invoice = await invoice_service.get_invoice_by_id(db, iid)
    if not invoice:
        raise HTTPException(status_code=404, detail="Invoice not found")
    return templates.TemplateResponse("admin/invoice_form.html", {
        "request": request, "current_user": current_user,
        "invoice": invoice_service.build_invoice_response(invoice),
        "users": [], "products": [],
    })


@router.post("/invoices/{invoice_id}/edit", response_class=HTMLResponse)
async def admin_invoice_edit_submit(
    request: Request,
    invoice_id: str,
    current_user: User = Depends(require_any_role("Admin", "Financial Manager")),
    db: AsyncSession = Depends(get_db),
):
    iid = uuid.UUID(invoice_id)
    invoice = await invoice_service.get_invoice_by_id(db, iid)
    if not invoice:
        raise HTTPException(status_code=404, detail="Invoice not found")
    form = await request.form()
    data = {
        "type": form.get("type", invoice.type),
        "status": form.get("status", invoice.status),
        "description": form.get("description", invoice.description),
    }
    await invoice_service.update_invoice(db, invoice, data)
    db.add(Log(record_id=invoice.id, table_name="invoices",
               description=f"ویرایش فاکتور: {invoice.reference_code}",
               created_by_user_id=current_user.id, type="Update"))
    await db.commit()
    return RedirectResponse(url=f"/administration/invoices/{invoice.id}", status_code=303)


@router.post("/invoices/{invoice_id}/delete", response_class=HTMLResponse)
async def admin_invoice_delete(
    request: Request,
    invoice_id: str,
    current_user: User = Depends(require_any_role("Admin", "Financial Manager")),
    db: AsyncSession = Depends(get_db),
):
    iid = uuid.UUID(invoice_id)
    invoice = await invoice_service.get_invoice_by_id(db, iid)
    if not invoice:
        raise HTTPException(status_code=404, detail="Invoice not found")
    code = invoice.reference_code
    await invoice_service.delete_invoice(db, invoice)
    db.add(Log(record_id=invoice.id, table_name="invoices",
               description=f"حذف فاکتور: {code}",
               created_by_user_id=current_user.id, type="Delete"))
    await db.commit()
    return RedirectResponse(url="/administration/invoices", status_code=303)


@router.post("/invoices/{invoice_id}/convert-to-return", response_class=HTMLResponse)
async def admin_invoice_convert_to_return(
    request: Request,
    invoice_id: str,
    current_user: User = Depends(require_any_role("Admin", "Financial Manager")),
    db: AsyncSession = Depends(get_db),
):
    iid = uuid.UUID(invoice_id)
    invoice = await invoice_service.get_invoice_by_id(db, iid)
    if not invoice:
        raise HTTPException(status_code=404, detail="Invoice not found")
    invoice.type = "ReturnFromSale"
    invoice.update_date = datetime.now(timezone.utc)
    db.add(Log(record_id=invoice.id, table_name="invoices",
               description=f"تبدیل فاکتور {invoice.reference_code} به بازگشتی",
               created_by_user_id=current_user.id, type="Update"))
    await db.commit()
    return RedirectResponse(url=f"/administration/invoices/{invoice.id}", status_code=303)


# ── Purchase Orders ──

@router.get("/purchase-orders", response_class=HTMLResponse)
async def admin_purchase_orders(
    request: Request,
    page: int = Query(1),
    current_user: User = Depends(require_any_role("Admin", "Financial Manager")),
    db: AsyncSession = Depends(get_db),
):
    pos, total = await invoice_service.get_all_purchase_orders(db, page, 20)
    po_list = []
    for po in pos:
        d = invoice_service.build_purchase_order_response(po)
        if po.created_by_user_id:
            user = await db.get(User, po.created_by_user_id)
            d["created_by_user_name"] = user.full_name or user.user_name if user else "---"
        else:
            d["created_by_user_name"] = "---"
        po_list.append(d)
    return templates.TemplateResponse("admin/purchase_orders.html", {
        "request": request, "current_user": current_user,
        "purchase_orders": po_list,
        "total": total, "page": page, "total_pages": (total + 19) // 20,
    })


@router.get("/purchase-orders/new", response_class=HTMLResponse)
async def admin_purchase_order_create(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Financial Manager")),
    db: AsyncSession = Depends(get_db),
):
    products = (await db.execute(select(Product).where(Product.is_removed == False, Product.no_display == False).limit(100))).scalars().all()
    suppliers = (await db.execute(select(Supplier).where(Supplier.is_removed == False))).scalars().all()
    currencies = (await db.execute(select(Currency).where(Currency.is_removed == False))).scalars().all()
    return templates.TemplateResponse("admin/purchase_order_form.html", {
        "request": request, "current_user": current_user,
        "products": products, "suppliers": suppliers, "currencies": currencies,
    })


@router.get("/purchase-orders/create", response_class=HTMLResponse)
async def admin_purchase_order_create_get(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Financial Manager")),
    db: AsyncSession = Depends(get_db),
):
    products = (await db.execute(select(Product).where(Product.is_removed == False, Product.no_display == False).limit(100))).scalars().all()
    suppliers = (await db.execute(select(Supplier).where(Supplier.is_removed == False))).scalars().all()
    currencies = (await db.execute(select(Currency).where(Currency.is_removed == False))).scalars().all()
    return templates.TemplateResponse("admin/purchase_order_form.html", {
        "request": request, "current_user": current_user,
        "products": products, "suppliers": suppliers, "currencies": currencies,
    })


@router.post("/purchase-orders/create", response_class=HTMLResponse)
async def admin_purchase_order_create_submit(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Financial Manager")),
    db: AsyncSession = Depends(get_db),
):
    from app.schemas.invoice import PurchaseOrderCreate
    form = await request.form()
    try:
        data = PurchaseOrderCreate(
            status="Ordered",
            shipping_and_clearance_price=float(form.get("shipping_and_clearance_price", 0)) if form.get("shipping_and_clearance_price") else None,
            details=[],
        )
        po = await invoice_service.create_purchase_order(db, data, current_user.id)
        db.add(Log(record_id=po.id, table_name="purchase_orders",
                   description=f"ایجاد سفارش خرید: {po.reference_code}",
                   created_by_user_id=current_user.id, type="Create"))
        await db.commit()
        return RedirectResponse(url="/administration/purchase-orders", status_code=303)
    except (ValueError, TypeError) as e:
        products = (await db.execute(select(Product).where(Product.is_removed == False, Product.no_display == False).limit(100))).scalars().all()
        suppliers = (await db.execute(select(Supplier).where(Supplier.is_removed == False))).scalars().all()
        currencies = (await db.execute(select(Currency).where(Currency.is_removed == False))).scalars().all()
        return templates.TemplateResponse("admin/purchase_order_form.html", {
            "request": request, "current_user": current_user,
            "products": products, "suppliers": suppliers, "currencies": currencies, "error": str(e),
        })


@router.get("/purchase-orders/{po_id}", response_class=HTMLResponse)
async def admin_purchase_order_detail(
    request: Request,
    po_id: str,
    current_user: User = Depends(require_any_role("Admin", "Financial Manager")),
    db: AsyncSession = Depends(get_db),
):
    try:
        pid = uuid.UUID(po_id)
    except ValueError:
        raise HTTPException(status_code=404, detail="Invalid PO ID")
    po = await invoice_service.get_purchase_order_by_id(db, pid)
    if not po:
        raise HTTPException(status_code=404, detail="Purchase order not found")
    return templates.TemplateResponse("admin/purchase_order_detail.html", {
        "request": request, "current_user": current_user,
        "purchase_order": invoice_service.build_purchase_order_response(po),
    })


@router.get("/purchase-orders/{po_id}/edit", response_class=HTMLResponse)
async def admin_purchase_order_edit_form(
    request: Request,
    po_id: str,
    current_user: User = Depends(require_any_role("Admin", "Financial Manager")),
    db: AsyncSession = Depends(get_db),
):
    try:
        pid = uuid.UUID(po_id)
    except ValueError:
        raise HTTPException(status_code=404, detail="Invalid PO ID")
    po = await invoice_service.get_purchase_order_by_id(db, pid)
    if not po:
        raise HTTPException(status_code=404, detail="Purchase order not found")
    return templates.TemplateResponse("admin/purchase_order_form.html", {
        "request": request, "current_user": current_user,
        "purchase_order": invoice_service.build_purchase_order_response(po),
        "products": [], "suppliers": [], "currencies": [],
    })


@router.post("/purchase-orders/{po_id}/edit", response_class=HTMLResponse)
async def admin_purchase_order_edit_submit(
    request: Request,
    po_id: str,
    current_user: User = Depends(require_any_role("Admin", "Financial Manager")),
    db: AsyncSession = Depends(get_db),
):
    try:
        pid = uuid.UUID(po_id)
    except ValueError:
        raise HTTPException(status_code=404, detail="Invalid PO ID")
    po = await invoice_service.get_purchase_order_by_id(db, pid)
    if not po:
        raise HTTPException(status_code=404, detail="Purchase order not found")
    form = await request.form()
    data = {
        "status": form.get("status", po.status),
        "date": form.get("date", po.date),
        "shipping_and_clearance_price": form.get("shipping_and_clearance_price", po.shipping_and_clearance_price),
    }
    await invoice_service.update_purchase_order(db, po, data)
    db.add(Log(record_id=po.id, table_name="purchase_orders",
               description=f"ویرایش سفارش خرید: {po.reference_code}",
               created_by_user_id=current_user.id, type="Update"))
    await db.commit()
    return RedirectResponse(url="/administration/purchase-orders", status_code=303)


@router.post("/purchase-orders/{po_id}/delete", response_class=HTMLResponse)
async def admin_purchase_order_delete(
    request: Request,
    po_id: str,
    current_user: User = Depends(require_any_role("Admin", "Financial Manager")),
    db: AsyncSession = Depends(get_db),
):
    try:
        pid = uuid.UUID(po_id)
    except ValueError:
        raise HTTPException(status_code=404, detail="Invalid PO ID")
    po = await invoice_service.get_purchase_order_by_id(db, pid)
    if not po:
        raise HTTPException(status_code=404, detail="Purchase order not found")
    ref = po.reference_code
    await invoice_service.delete_purchase_order(db, po)
    db.add(Log(record_id=po.id, table_name="purchase_orders",
               description=f"حذف سفارش خرید: {ref}",
               created_by_user_id=current_user.id, type="Delete"))
    await db.commit()
    return RedirectResponse(url="/administration/purchase-orders", status_code=303)


# ── Suppliers ──

@router.get("/suppliers", response_class=HTMLResponse)
async def admin_suppliers(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Financial Manager")),
    db: AsyncSession = Depends(get_db),
):
    suppliers = await invoice_service.get_all_suppliers(db)
    return templates.TemplateResponse("admin/suppliers.html", {
        "request": request, "current_user": current_user, "suppliers": suppliers,
    })


@router.get("/suppliers/new", response_class=HTMLResponse)
async def admin_supplier_create(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Financial Manager")),
    db: AsyncSession = Depends(get_db),
):
    return templates.TemplateResponse("admin/supplier_form.html", {
        "request": request, "current_user": current_user,
    })


@router.get("/suppliers/create", response_class=HTMLResponse)
async def admin_supplier_create_get(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Financial Manager")),
    db: AsyncSession = Depends(get_db),
):
    return templates.TemplateResponse("admin/supplier_form.html", {
        "request": request, "current_user": current_user,
    })


@router.post("/suppliers/create", response_class=HTMLResponse)
async def admin_supplier_create_submit(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Financial Manager")),
    db: AsyncSession = Depends(get_db),
):
    from app.schemas.invoice import SupplierCreate
    form = await request.form()
    data = SupplierCreate(
        telephone=form.get("telephone", ""),
        address=form.get("address", ""),
        site=form.get("site", ""),
        intermediary_name=form.get("intermediary_name", ""),
    )
    supplier = await invoice_service.create_supplier(db, data, current_user.id)
    db.add(Log(record_id=supplier.id, table_name="suppliers",
               description=f"ایجاد تأمین‌کننده: {supplier.intermediary_name}",
               created_by_user_id=current_user.id, type="Create"))
    await db.commit()
    return RedirectResponse(url="/administration/suppliers", status_code=303)


@router.get("/suppliers/{supplier_id}", response_class=HTMLResponse)
async def admin_supplier_detail(
    request: Request,
    supplier_id: str,
    current_user: User = Depends(require_any_role("Admin", "Financial Manager")),
    db: AsyncSession = Depends(get_db),
):
    try:
        sid = uuid.UUID(supplier_id)
    except ValueError:
        raise HTTPException(status_code=404, detail="Invalid supplier ID")
    supplier = await invoice_service.get_supplier_by_id(db, sid)
    if not supplier:
        raise HTTPException(status_code=404, detail="Supplier not found")
    return templates.TemplateResponse("admin/supplier_detail.html", {
        "request": request, "current_user": current_user, "supplier": supplier,
    })


@router.get("/suppliers/{supplier_id}/edit", response_class=HTMLResponse)
async def admin_supplier_edit_form(
    request: Request,
    supplier_id: str,
    current_user: User = Depends(require_any_role("Admin", "Financial Manager")),
    db: AsyncSession = Depends(get_db),
):
    try:
        sid = uuid.UUID(supplier_id)
    except ValueError:
        raise HTTPException(status_code=404, detail="Invalid supplier ID")
    supplier = await invoice_service.get_supplier_by_id(db, sid)
    if not supplier:
        raise HTTPException(status_code=404, detail="Supplier not found")
    return templates.TemplateResponse("admin/supplier_form.html", {
        "request": request, "current_user": current_user, "supplier": supplier,
    })


@router.post("/suppliers/{supplier_id}/edit", response_class=HTMLResponse)
async def admin_supplier_edit_submit(
    request: Request,
    supplier_id: str,
    current_user: User = Depends(require_any_role("Admin", "Financial Manager")),
    db: AsyncSession = Depends(get_db),
):
    try:
        sid = uuid.UUID(supplier_id)
    except ValueError:
        raise HTTPException(status_code=404, detail="Invalid supplier ID")
    supplier = await invoice_service.get_supplier_by_id(db, sid)
    if not supplier:
        raise HTTPException(status_code=404, detail="Supplier not found")
    form = await request.form()
    data = {
        "telephone": form.get("telephone", ""),
        "address": form.get("address", ""),
        "site": form.get("site", ""),
        "intermediary_name": form.get("intermediary_name", ""),
    }
    await invoice_service.update_supplier(db, supplier, data)
    db.add(Log(record_id=supplier.id, table_name="suppliers",
               description=f"ویرایش تأمین‌کننده: {supplier.intermediary_name}",
               created_by_user_id=current_user.id, type="Update"))
    await db.commit()
    return RedirectResponse(url="/administration/suppliers", status_code=303)


@router.post("/suppliers/{supplier_id}/delete", response_class=HTMLResponse)
async def admin_supplier_delete(
    request: Request,
    supplier_id: str,
    current_user: User = Depends(require_any_role("Admin", "Financial Manager")),
    db: AsyncSession = Depends(get_db),
):
    try:
        sid = uuid.UUID(supplier_id)
    except ValueError:
        raise HTTPException(status_code=404, detail="Invalid supplier ID")
    supplier = await invoice_service.get_supplier_by_id(db, sid)
    if not supplier:
        raise HTTPException(status_code=404, detail="Supplier not found")
    name = supplier.intermediary_name
    await invoice_service.delete_supplier(db, supplier)
    db.add(Log(record_id=supplier.id, table_name="suppliers",
               description=f"حذف تأمین‌کننده: {name}",
               created_by_user_id=current_user.id, type="Delete"))
    await db.commit()
    return RedirectResponse(url="/administration/suppliers", status_code=303)


# ── Receipts ──

@router.get("/receipts", response_class=HTMLResponse)
async def admin_receipts(
    request: Request,
    page: int = Query(1),
    status_filter: str = Query(None),
    current_user: User = Depends(require_any_role("Admin", "Financial Manager")),
    db: AsyncSession = Depends(get_db),
):
    receipts, total = await finance_service.get_all_receipts(db, page, 20, status_filter)
    # Enrich each receipt with user full name/email and the linked order's payable
    rows = []
    for r in receipts:
        d = finance_service.build_receipt_response(r)
        user_full_name = ""
        user_email = ""
        if r.user_id:
            u = await db.get(User, r.user_id)
            if u:
                user_full_name = u.full_name
                user_email = u.email or ""
        payable = None
        if r.order_id:
            o = await db.get(Order, r.order_id)
            if o:
                payable = float(o.payable or 0)
        d["user_full_name"] = user_full_name
        d["user_email"] = user_email
        d["payable"] = payable
        rows.append(d)
    return templates.TemplateResponse("admin/receipt_list.html", {
        "request": request, "current_user": current_user,
        "receipts": rows,
        "total": total, "page": page, "total_pages": (total + 19) // 20,
        "status_filter": status_filter,
    })


def _build_admin_receipt_response(receipt: Receipt) -> dict:
    return {
        "id": str(receipt.id),
        "reference_code": receipt.reference_code,
        "price": float(receipt.price or 0),
        "status": receipt.status,
        "status_name": order_service.RECEIPT_STATUS_NAMES.get(receipt.status, receipt.status or "-"),
        "tab": receipt.tab or "",
        "description": receipt.description,
        "deposit_date": receipt.deposit_date,
        "deposit_date_str": to_farsi_full(receipt.deposit_date) if receipt.deposit_date else "",
        "destination_bank": receipt.destination_bank,
        "paya": receipt.paya,
        "image_url": receipt.image_url,
        "user_id": str(receipt.user_id) if receipt.user_id else None,
        "user_full_name": receipt.user.full_name if receipt.user else "",
        "order_id": str(receipt.order_id) if receipt.order_id else None,
        "insert_date": receipt.insert_date,
        "insert_date_str": to_farsi_full(receipt.insert_date) if receipt.insert_date else "",
    }


@router.get("/receipts/{receipt_id}", response_class=HTMLResponse)
async def admin_receipt_detail(
    request: Request,
    receipt_id: str,
    current_user: User = Depends(require_any_role("Admin", "Financial Manager")),
    db: AsyncSession = Depends(get_db),
):
    try:
        rid = uuid.UUID(receipt_id)
    except ValueError:
        raise HTTPException(status_code=404, detail="Invalid receipt ID")
    receipt = (await db.execute(
        select(Receipt).options(selectinload(Receipt.user)).where(Receipt.id == rid, Receipt.is_removed == False)
    )).scalar_one_or_none()
    if not receipt:
        raise HTTPException(status_code=404, detail="Receipt not found")
    return templates.TemplateResponse("admin/receipt_detail.html", {
        "request": request, "current_user": current_user,
        "receipt": _build_admin_receipt_response(receipt),
    })


@router.post("/receipts/{receipt_id}/confirm", response_class=HTMLResponse)
async def admin_receipt_confirm(
    request: Request,
    receipt_id: str,
    current_user: User = Depends(require_any_role("Admin", "Financial Manager")),
    db: AsyncSession = Depends(get_db),
):
    try:
        rid = uuid.UUID(receipt_id)
    except ValueError:
        raise HTTPException(status_code=404, detail="Invalid receipt ID")
    receipt = (await db.execute(
        select(Receipt).where(Receipt.id == rid, Receipt.is_removed == False)
    )).scalar_one_or_none()
    if not receipt:
        raise HTTPException(status_code=404, detail="Receipt not found")

    # Mirror .NET Accept: set receipt status to Confirmed, then update order status.
    receipt.status = "Confirmed"
    receipt.update_date = datetime.now(timezone.utc)

    if receipt.order_id:
        order = await db.get(Order, receipt.order_id)
        if order and order.is_removed == False:
            if order.order_status in ("AwaitingPayment", "Paid"):
                rsum = (await db.execute(
                    select(func.coalesce(func.sum(Receipt.price), 0))
                    .where(Receipt.order_id == order.id, Receipt.status == "Confirmed", Receipt.is_removed == False)
                )).scalar() or 0
                psum = (await db.execute(
                    select(func.coalesce(func.sum(PaymentRequest.amount), 0))
                    .where(PaymentRequest.order_id == order.id, PaymentRequest.status == "Success", PaymentRequest.is_removed == False)
                )).scalar() or 0
                total = float(rsum) + float(psum)
                payable = float(order.payable or 0)
                if total == payable:
                    order.order_status = "ConfirmedPayment"
                    db.add(OrderStatusRecord(
                        id=uuid.uuid4(), order_id=order.id, status="ConfirmedPayment",
                        comment="Accept payment", created_by_user_id=current_user.id,
                        insert_date=datetime.now(timezone.utc), update_date=datetime.now(timezone.utc),
                    ))
                elif total > payable:
                    order.order_status = "NeedsToBeChecked"
                    db.add(OrderStatusRecord(
                        id=uuid.uuid4(), order_id=order.id, status="NeedsToBeChecked",
                        comment="Need to be checked", created_by_user_id=current_user.id,
                        insert_date=datetime.now(timezone.utc), update_date=datetime.now(timezone.utc),
                    ))
                order.update_date = datetime.now(timezone.utc)

    db.add(Log(record_id=receipt.id, table_name="receipts",
               description=f"تائید فیش پرداخت: {receipt.reference_code or receipt.id}",
               created_by_user_id=current_user.id, type="Update"))
    await db.commit()
    return RedirectResponse(url=f"/administration/receipts/{receipt.id}", status_code=303)


@router.post("/receipts/{receipt_id}/reject", response_class=HTMLResponse)
async def admin_receipt_reject(
    request: Request,
    receipt_id: str,
    current_user: User = Depends(require_any_role("Admin", "Financial Manager")),
    db: AsyncSession = Depends(get_db),
):
    try:
        rid = uuid.UUID(receipt_id)
    except ValueError:
        raise HTTPException(status_code=404, detail="Invalid receipt ID")
    receipt = (await db.execute(
        select(Receipt).where(Receipt.id == rid, Receipt.is_removed == False)
    )).scalar_one_or_none()
    if not receipt:
        raise HTTPException(status_code=404, detail="Receipt not found")
    receipt.status = "Failed"
    receipt.update_date = datetime.now(timezone.utc)
    db.add(Log(record_id=receipt.id, table_name="receipts",
               description=f"رد فیش پرداخت: {receipt.reference_code or receipt.id}",
               created_by_user_id=current_user.id, type="Update"))
    await db.commit()
    return RedirectResponse(url=f"/administration/receipts/{receipt.id}", status_code=303)


@router.post("/receipts/{receipt_id}/delete", response_class=HTMLResponse)
async def admin_receipt_delete(
    request: Request,
    receipt_id: str,
    current_user: User = Depends(require_any_role("Admin", "Financial Manager")),
    db: AsyncSession = Depends(get_db),
):
    try:
        rid = uuid.UUID(receipt_id)
    except ValueError:
        raise HTTPException(status_code=404, detail="Invalid receipt ID")
    receipt = (await db.execute(
        select(Receipt).where(Receipt.id == rid, Receipt.is_removed == False)
    )).scalar_one_or_none()
    if not receipt:
        raise HTTPException(status_code=404, detail="Receipt not found")
    if receipt.status != "AwaitingConfirmation":
        raise HTTPException(status_code=400, detail="فقط فیش‌های در انتظار تائید قابل حذف هستند")
    ref = receipt.reference_code
    receipt.is_removed = True
    receipt.update_date = datetime.now(timezone.utc)
    db.add(Log(record_id=receipt.id, table_name="receipts",
               description=f"حذف فیش پرداخت: {ref or receipt.id}",
               created_by_user_id=current_user.id, type="Delete"))
    await db.commit()
    return RedirectResponse(url="/administration/receipts", status_code=303)


# ── Currency ──

@router.get("/currencies", response_class=HTMLResponse)
async def admin_currencies(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Financial Manager")),
    db: AsyncSession = Depends(get_db),
):
    currencies = await finance_service.get_all_currencies(db)
    return templates.TemplateResponse("admin/currency_list.html", {
        "request": request, "current_user": current_user, "currencies": currencies,
    })


@router.get("/currencies/new", response_class=HTMLResponse)
async def admin_currency_create_form(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Financial Manager")),
):
    return templates.TemplateResponse("admin/currency_form.html", {
        "request": request, "current_user": current_user,
        "currency": None, "update_price": False,
    })


@router.post("/currencies/new", response_class=HTMLResponse)
async def admin_currency_create_submit(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Financial Manager")),
    db: AsyncSession = Depends(get_db),
):
    form = await request.form()
    name = form.get("name", "").strip()
    if not name:
        return templates.TemplateResponse("admin/currency_form.html", {
            "request": request, "current_user": current_user,
            "currency": None, "update_price": False, "error": "نام ارز الزامی است",
        })
    currency = await finance_service.create_currency(db, name, current_user.id)
    try:
        price = float(form.get("price", 0))
        if price > 0:
            from app.services.finance_service import add_currency_price
            await add_currency_price(db, currency.id, price)
    except (ValueError, TypeError):
        pass
    db.add(Log(record_id=currency.id, table_name="currencies", description=f"ایجاد ارز: {name}", created_by_user_id=current_user.id, type="Create"))
    await db.commit()
    return RedirectResponse(url="/administration/currency", status_code=303)


@router.get("/currencies/{currency_id}", response_class=HTMLResponse)
async def admin_currency_detail(
    request: Request,
    currency_id: str,
    page: int = Query(1),
    current_user: User = Depends(require_any_role("Admin", "Financial Manager")),
    db: AsyncSession = Depends(get_db),
):
    try:
        cid = uuid.UUID(currency_id)
    except ValueError:
        raise HTTPException(status_code=404, detail="Invalid currency ID")
    currency = await finance_service.get_currency_by_id(db, cid)
    if not currency:
        raise HTTPException(status_code=404, detail="Currency not found")
    details, total = await finance_service.get_all_currency_details(db, cid, page, 20)
    return templates.TemplateResponse("admin/currency_detail.html", {
        "request": request, "current_user": current_user,
        "currency": currency, "details": details,
        "total": total, "page": page, "total_pages": (total + 19) // 20,
    })


@router.get("/currencies/{currency_id}/edit", response_class=HTMLResponse)
async def admin_currency_edit_form(
    request: Request,
    currency_id: str,
    current_user: User = Depends(require_any_role("Admin", "Financial Manager")),
    db: AsyncSession = Depends(get_db),
):
    try:
        cid = uuid.UUID(currency_id)
    except ValueError:
        raise HTTPException(status_code=404, detail="Invalid currency ID")
    currency = await finance_service.get_currency_by_id(db, cid)
    if not currency:
        raise HTTPException(status_code=404, detail="Currency not found")
    return templates.TemplateResponse("admin/currency_form.html", {
        "request": request, "current_user": current_user,
        "currency": currency, "update_price": False,
    })


@router.post("/currencies/{currency_id}/edit", response_class=HTMLResponse)
async def admin_currency_edit_submit(
    request: Request,
    currency_id: str,
    current_user: User = Depends(require_any_role("Admin", "Financial Manager")),
    db: AsyncSession = Depends(get_db),
):
    try:
        cid = uuid.UUID(currency_id)
    except ValueError:
        raise HTTPException(status_code=404, detail="Invalid currency ID")
    currency = await finance_service.get_currency_by_id(db, cid)
    if not currency:
        raise HTTPException(status_code=404, detail="Currency not found")
    form = await request.form()
    name = form.get("name", "").strip()
    if not name:
        return templates.TemplateResponse("admin/currency_form.html", {
            "request": request, "current_user": current_user,
            "currency": currency, "update_price": False, "error": "نام ارز الزامی است",
        })
    await finance_service.update_currency(db, currency, name)
    db.add(Log(record_id=currency.id, table_name="currencies", description=f"ویرایش ارز: {name}", created_by_user_id=current_user.id, type="Update"))
    await db.commit()
    return RedirectResponse(url="/administration/currency", status_code=303)


@router.post("/currencies/{currency_id}/delete", response_class=HTMLResponse)
async def admin_currency_delete(
    request: Request,
    currency_id: str,
    current_user: User = Depends(require_any_role("Admin", "Financial Manager")),
    db: AsyncSession = Depends(get_db),
):
    try:
        cid = uuid.UUID(currency_id)
    except ValueError:
        raise HTTPException(status_code=404, detail="Invalid currency ID")
    currency = await finance_service.get_currency_by_id(db, cid)
    if not currency:
        raise HTTPException(status_code=404, detail="Currency not found")
    name = currency.name
    await finance_service.delete_currency(db, currency)
    db.add(Log(record_id=currency.id, table_name="currencies", description=f"حذف ارز: {name}", created_by_user_id=current_user.id, type="Delete"))
    await db.commit()
    return RedirectResponse(url="/administration/currency", status_code=303)


@router.get("/currencies/{currency_id}/update-price", response_class=HTMLResponse)
async def admin_currency_update_price_form(
    request: Request,
    currency_id: str,
    page: int = Query(1),
    product_search: str = Query(None),
    part_search: str = Query(None),
    current_user: User = Depends(require_any_role("Admin", "Financial Manager")),
    db: AsyncSession = Depends(get_db),
):
    try:
        cid = uuid.UUID(currency_id)
    except ValueError:
        raise HTTPException(status_code=404, detail="Invalid currency ID")
    currency = await finance_service.get_currency_by_id(db, cid)
    if not currency:
        raise HTTPException(status_code=404, detail="Currency not found")
    detail_stmt = (
        select(CurrencyDetail)
        .where(CurrencyDetail.currency_id == cid, CurrencyDetail.is_removed == False)
        .order_by(CurrencyDetail.insert_date.desc())
        .limit(1)
    )
    detail_result = await db.execute(detail_stmt)
    last_detail = detail_result.scalar_one_or_none()
    currency._last_price = last_detail
    conditions = [Product.is_removed == False, Product.no_display == False]
    if product_search:
        conditions.append(Product.name.ilike(f"%{product_search}%"))
    if part_search:
        conditions.append(Product.part_number.ilike(f"%{part_search}%"))
    count_stmt = select(func.count(Product.id)).where(*conditions)
    total_products = (await db.execute(count_stmt)).scalar() or 0
    stmt = (
        select(Product)
        .where(*conditions)
        .order_by(Product.insert_date.desc())
        .offset((page - 1) * 25)
        .limit(25)
    )
    products = (await db.execute(stmt)).scalars().all()
    return templates.TemplateResponse("admin/currency_update_price.html", {
        "request": request, "current_user": current_user,
        "currency": currency, "products": products,
        "total_products": total_products, "page": page,
        "total_pages": (total_products + 24) // 25,
        "product_search": product_search or "", "part_search": part_search or "",
    })


@router.post("/currencies/{currency_id}/update-price", response_class=HTMLResponse)
async def admin_currency_update_price_submit(
    request: Request,
    currency_id: str,
    current_user: User = Depends(require_any_role("Admin", "Financial Manager")),
    db: AsyncSession = Depends(get_db),
):
    try:
        cid = uuid.UUID(currency_id)
    except ValueError:
        raise HTTPException(status_code=404, detail="Invalid currency ID")
    currency = await finance_service.get_currency_by_id(db, cid)
    if not currency:
        raise HTTPException(status_code=404, detail="Currency not found")
    form = await request.form()
    try:
        price = float(form.get("price", 0))
        profit_rate = float(form.get("profit_rate", 0))
        discount_percent = float(form.get("discount_percent", 0))
    except (ValueError, TypeError):
        return RedirectResponse(url=f"/administration/currencies/{currency_id}/update-price", status_code=303)
    cd = await finance_service.add_currency_price(db, cid, price)
    # Update all products/varieties with this currency
    from app.models.product import Variety as VarietyModel
    conditions = [VarietyModel.is_removed == False, VarietyModel.currency_id == cid]
    stmt = select(VarietyModel).where(*conditions)
    varieties = (await db.execute(stmt)).scalars().all()
    for v in varieties:
        v.currency_price = price
        if profit_rate >= 0:
            v.profit_rate = profit_rate
        v.automatic_price_calculation = True
        v.update_date = datetime.now(timezone.utc)
    db.add(Log(record_id=cd.id, table_name="currency_details",
               description=f"بروزرسانی قیمت دسته‌ای {currency.name}: قیمت={price}, نرخ سود={profit_rate}, تخفیف={discount_percent}",
               created_by_user_id=current_user.id, type="Update"))
    await db.commit()
    return RedirectResponse(url=f"/administration/currencies/{currency_id}/update-price", status_code=303)


@router.get("/currencies/{currency_id}/export-excel")
async def admin_currency_export_excel(
    request: Request,
    currency_id: str,
    current_user: User = Depends(require_any_role("Admin", "Financial Manager")),
    db: AsyncSession = Depends(get_db),
):
    try:
        cid = uuid.UUID(currency_id)
    except ValueError:
        raise HTTPException(status_code=404, detail="Invalid currency ID")
    currency = await finance_service.get_currency_by_id(db, cid)
    if not currency:
        raise HTTPException(status_code=404, detail="Currency not found")
    import openpyxl
    from io import BytesIO
    from datetime import datetime
    import jdatetime
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Currencies"
    headers = ["ردیف", "دسته بندی", "شماره قطعه", "نام", "قیمت", "قیمت ارزی", "موجودی انبار", "درصد تخفیف", "نرخ سود", "تاریخ ایجاد", "محصول", "نام انگلیسی", "تاریخ خرید", "لینک"]
    ws.append(headers)
    from app.models.product import Product as ProductModel, Category
    stmt = (
        select(ProductModel)
        .options(selectinload(ProductModel.category))
        .where(ProductModel.is_removed == False, ProductModel.no_display == False)
        .order_by(ProductModel.insert_date.desc())
    )
    products = (await db.execute(stmt)).unique().scalars().all()
    for idx, p in enumerate(products, 1):
        cat_name = p.category.name if p.category else ""
        link = f"https://eshop.eca.ir/product/{p.slug}" if p.slug else ""
        ws.append([
            idx, cat_name, p.part_number or "", p.name or "",
            float(p.price) if p.price else 0,
            float(p.currency_price) if p.currency_price else 0,
            p.stock_quantity or 0,
            float(p.discount_percentage) if p.discount_percentage else 0,
            float(p.profit_rate) if p.profit_rate else 0,
            p.insert_date.strftime("%Y/%m/%d") if p.insert_date else "",
            "", p.en_name or "",
            p.purchase_date.strftime("%Y/%m/%d") if p.purchase_date else "",
            link,
        ])
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    from fastapi.responses import StreamingResponse
    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=Currencies.xlsx"},
    )


# ── Currency Details ──

@router.get("/currency-details/create/{currency_id}", response_class=HTMLResponse)
async def admin_currency_detail_create_form(
    request: Request,
    currency_id: str,
    current_user: User = Depends(require_any_role("Admin", "Financial Manager")),
    db: AsyncSession = Depends(get_db),
):
    try:
        cid = uuid.UUID(currency_id)
    except ValueError:
        raise HTTPException(status_code=404, detail="Invalid currency ID")
    currency = await finance_service.get_currency_by_id(db, cid)
    if not currency:
        raise HTTPException(status_code=404, detail="Currency not found")
    return templates.TemplateResponse("admin/currency_detail_form.html", {
        "request": request, "current_user": current_user, "currency": currency,
    })


@router.post("/currency-details/create/{currency_id}", response_class=HTMLResponse)
async def admin_currency_detail_create_submit(
    request: Request,
    currency_id: str,
    current_user: User = Depends(require_any_role("Admin", "Financial Manager")),
    db: AsyncSession = Depends(get_db),
):
    try:
        cid = uuid.UUID(currency_id)
    except ValueError:
        raise HTTPException(status_code=404, detail="Invalid currency ID")
    currency = await finance_service.get_currency_by_id(db, cid)
    if not currency:
        raise HTTPException(status_code=404, detail="Currency not found")
    form = await request.form()
    import jdatetime
    try:
        price = float(form.get("price", 0))
        date_str = form.get("date", "")
        date = None
        if date_str:
            try:
                date = jdatetime.datetime.strptime(date_str, "%Y/%m/%d %H:%M:%S").togregorian()
            except ValueError:
                try:
                    date = jdatetime.datetime.strptime(date_str, "%Y/%m/%d").togregorian()
                except ValueError:
                    date = datetime.now(timezone.utc)
        cd = await finance_service.add_currency_price(db, cid, price, date)
        db.add(Log(record_id=cd.id, table_name="currency_details",
                   description=f"ایجاد جزئیات ارز {currency.name}: {price}",
                   created_by_user_id=current_user.id, type="Create"))
        await db.commit()
        return RedirectResponse(url="/administration/currencies", status_code=303)
    except (ValueError, TypeError) as e:
        return templates.TemplateResponse("admin/currency_detail_form.html", {
            "request": request, "current_user": current_user,
            "currency": currency, "error": "مقادیر نامعتبر",
        })


# ── Warehouse ──

@router.get("/warehouse", response_class=HTMLResponse)
async def admin_warehouse(
    request: Request,
    page: int = Query(1),
    current_user: User = Depends(require_any_role("Admin", "Warehouse Keeper")),
    db: AsyncSession = Depends(get_db),
):
    movements, total = await warehouse_service.get_all_movements(db, page, 20)
    low_stock = await warehouse_service.get_low_stock_products(db, 5)
    return templates.TemplateResponse("admin/warehouse_movements.html", {
        "request": request, "current_user": current_user,
        "movements": [warehouse_service.build_movement_response(m) for m in movements],
        "low_stock_products": low_stock,
        "total": total, "page": page, "total_pages": (total + 19) // 20,
    })


# ── Tickets ──

def _normalize_ticket_status(value):
    return {
        "open": "Open", "answered": "Answered", "closed": "Closed",
        "Open": "Open", "Answered": "Answered", "Closed": "Closed",
    }.get(value or "", value or None)


@router.get("/tickets", response_class=HTMLResponse)
async def admin_tickets(
    request: Request,
    page: int = Query(1),
    status_filter: str = Query(None),
    category_filter: str = Query(None),
    priority_filter: str = Query(None),
    current_user: User = Depends(get_admin_user),
    db: AsyncSession = Depends(get_db),
):
    tickets, total = await support_service.get_all_tickets(
        db, page, 20,
        _normalize_ticket_status(status_filter), category_filter, priority_filter,
    )
    from app.services.support_service import (
        TICKET_CATEGORIES, TICKET_PRIORITIES, TICKET_STATUSES,
    )
    return templates.TemplateResponse("admin/tickets.html", {
        "request": request, "current_user": current_user,
        "tickets": [support_service.build_ticket_response(t) for t in tickets],
        "total": total, "page": page, "total_pages": (total + 19) // 20,
        "status_filter": status_filter, "category_filter": category_filter, "priority_filter": priority_filter,
        "categories": TICKET_CATEGORIES, "priorities": TICKET_PRIORITIES, "statuses": TICKET_STATUSES,
    })


@router.get("/tickets/{ticket_id}", response_class=HTMLResponse)
async def admin_ticket_detail(
    request: Request,
    ticket_id: str,
    current_user: User = Depends(get_admin_user),
    db: AsyncSession = Depends(get_db),
):
    import uuid
    tid = uuid.UUID(ticket_id)
    ticket = await support_service.get_ticket_by_id(db, tid)
    if not ticket:
        raise HTTPException(status_code=404, detail="Ticket not found")
    from app.services.support_service import (
        TICKET_CATEGORIES, TICKET_PRIORITIES, TICKET_STATUSES,
    )
    return templates.TemplateResponse("admin/ticket_detail.html", {
        "request": request, "current_user": current_user,
        "ticket": support_service.build_ticket_response(ticket),
        "categories": TICKET_CATEGORIES, "priorities": TICKET_PRIORITIES, "statuses": TICKET_STATUSES,
    })


@router.post("/tickets/{ticket_id}/reply")
async def admin_ticket_reply(
    request: Request,
    ticket_id: str,
    message: str = Form(""),
    file: Optional[UploadFile] = File(None),
    current_user: User = Depends(get_admin_user),
    db: AsyncSession = Depends(get_db),
):
    import uuid
    try:
        tid = uuid.UUID(ticket_id)
    except ValueError:
        return RedirectResponse(url="/administration/tickets", status_code=303)
    ticket = await support_service.get_ticket_by_id(db, tid)
    if not ticket:
        return RedirectResponse(url="/administration/tickets", status_code=303)

    file_path = None
    file_name = None
    if file and file.filename:
        import aiofiles
        upload_dir = "app/static/uploads/tickets"
        os.makedirs(upload_dir, exist_ok=True)
        ext = file.filename.split(".")[-1] if "." in file.filename else "bin"
        fname = f"ticket_admin_{current_user.id.hex[:8]}_{uuid.uuid4().hex[:8]}.{ext}"
        file_path = f"/static/uploads/tickets/{fname}"
        file_name = file.filename
        content = await file.read()
        async with aiofiles.open(os.path.join(upload_dir, fname), "wb") as f:
            await f.write(content)

    if message.strip():
        await support_service.reply_to_ticket(
            db, ticket, message.strip(), current_user.id,
            is_admin=True, file_path=file_path, file_name=file_name,
        )
        await db.commit()
    return RedirectResponse(url=f"/administration/tickets/{ticket.id}", status_code=303)


@router.post("/tickets/{ticket_id}/close")
async def admin_ticket_close(
    request: Request,
    ticket_id: str,
    current_user: User = Depends(get_admin_user),
    db: AsyncSession = Depends(get_db),
):
    import uuid
    try:
        tid = uuid.UUID(ticket_id)
    except ValueError:
        return RedirectResponse(url="/administration/tickets", status_code=303)
    ticket = await support_service.get_ticket_by_id(db, tid)
    if not ticket:
        return RedirectResponse(url="/administration/tickets", status_code=303)
    await support_service.close_ticket(db, ticket)
    await db.commit()
    return RedirectResponse(url=f"/administration/tickets/{ticket.id}", status_code=303)


# ── Product Types ──

@router.get("/product-types", response_class=HTMLResponse)
async def admin_product_types(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    items = (await db.execute(
        select(ProductType).where(ProductType.is_removed == False).order_by(ProductType.insert_date.desc())
    )).scalars().all()
    return templates.TemplateResponse("admin/product_types.html", {
        "request": request, "current_user": current_user, "items": items,
    })


@router.get("/product-types/new", response_class=HTMLResponse)
async def admin_product_type_create(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    return templates.TemplateResponse("admin/product_type_form.html", {
        "request": request, "current_user": current_user, "product_type": None,
    })


@router.post("/product-types/new", response_class=HTMLResponse)
async def admin_product_type_create_submit(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    form = await request.form()
    fa_name = (form.get("fa_name") or "").strip()
    en_name = (form.get("en_name") or "").strip()

    pt = ProductType(fa_name=fa_name or None, en_name=en_name or None, created_by_user_id=current_user.id)
    db.add(pt)
    await db.flush()

    log_text = form.get("log") or ""
    if log_text:
        db.add(Log(
            record_id=pt.id,
            table_name="product_types",
            description=log_text,
            created_by_user_id=current_user.id,
            type="Create",
        ))
    return RedirectResponse(url="/administration/product-types", status_code=303)


@router.get("/product-types/{product_type_id}/edit", response_class=HTMLResponse)
async def admin_product_type_edit(
    request: Request,
    product_type_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    pt = await db.get(ProductType, uuid.UUID(product_type_id))
    if pt is None or pt.is_removed:
        raise HTTPException(status_code=404, detail="Product type not found")
    return templates.TemplateResponse("admin/product_type_form.html", {
        "request": request, "current_user": current_user, "product_type": pt,
    })


@router.post("/product-types/{product_type_id}/edit", response_class=HTMLResponse)
async def admin_product_type_edit_submit(
    request: Request,
    product_type_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    pt = await db.get(ProductType, uuid.UUID(product_type_id))
    if pt is None or pt.is_removed:
        raise HTTPException(status_code=404, detail="Product type not found")

    form = await request.form()
    pt.fa_name = (form.get("fa_name") or "").strip() or None
    pt.en_name = (form.get("en_name") or "").strip() or None
    await db.flush()

    log_text = form.get("log") or ""
    if log_text:
        db.add(Log(
            record_id=pt.id,
            table_name="product_types",
            description=log_text,
            created_by_user_id=current_user.id,
            type="Update",
        ))
    return RedirectResponse(url="/administration/product-types", status_code=303)


@router.get("/product-types/{product_type_id}", response_class=HTMLResponse)
async def admin_product_type_detail(
    request: Request,
    product_type_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    pt = await db.get(ProductType, uuid.UUID(product_type_id))
    if pt is None or pt.is_removed:
        raise HTTPException(status_code=404, detail="Product type not found")
    return templates.TemplateResponse("admin/product_type_detail.html", {
        "request": request, "current_user": current_user, "product_type": pt,
    })


@router.post("/product-types/{product_type_id}/delete", response_class=HTMLResponse)
async def admin_product_type_delete(
    request: Request,
    product_type_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    pt = await db.get(ProductType, uuid.UUID(product_type_id))
    if pt is None or pt.is_removed:
        raise HTTPException(status_code=404, detail="Product type not found")
    name = pt.fa_name or pt.en_name or ""
    pt.is_removed = True
    await db.flush()
    db.add(Log(
        record_id=pt.id,
        table_name="product_types",
        description=f"حذف نوع محصول: {name}",
        created_by_user_id=current_user.id,
        type="Delete",
    ))
    return RedirectResponse(url="/administration/product-types", status_code=303)


# ── Product Units ──

@router.get("/product-units", response_class=HTMLResponse)
async def admin_product_units(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    items = (await db.execute(
        select(ProductUnit).where(ProductUnit.is_removed == False).order_by(ProductUnit.insert_date.desc())
    )).scalars().all()
    return templates.TemplateResponse("admin/product_units.html", {
        "request": request, "current_user": current_user, "items": items,
    })


@router.get("/product-units/new", response_class=HTMLResponse)
async def admin_product_unit_create(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    return templates.TemplateResponse("admin/product_unit_form.html", {
        "request": request, "current_user": current_user, "product_unit": None,
    })


@router.post("/product-units/new", response_class=HTMLResponse)
async def admin_product_unit_create_submit(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    form = await request.form()
    name = (form.get("name") or "").strip()
    abbreviation = (form.get("abbreviation") or "").strip()
    product_unit_tax_id = (form.get("product_unit_tax_id") or "").strip()

    pu = ProductUnit(
        name=name or None,
        abbreviation=abbreviation or None,
        product_unit_tax_id=product_unit_tax_id or None,
        created_by_user_id=current_user.id,
    )
    db.add(pu)
    await db.flush()

    log_text = form.get("log") or ""
    if log_text:
        db.add(Log(
            record_id=pu.id,
            table_name="product_units",
            description=log_text,
            created_by_user_id=current_user.id,
            type="Create",
        ))
    return RedirectResponse(url="/administration/product-units", status_code=303)


@router.get("/product-units/{product_unit_id}/edit", response_class=HTMLResponse)
async def admin_product_unit_edit(
    request: Request,
    product_unit_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    pu = await db.get(ProductUnit, uuid.UUID(product_unit_id))
    if pu is None or pu.is_removed:
        raise HTTPException(status_code=404, detail="Product unit not found")
    return templates.TemplateResponse("admin/product_unit_form.html", {
        "request": request, "current_user": current_user, "product_unit": pu,
    })


@router.post("/product-units/{product_unit_id}/edit", response_class=HTMLResponse)
async def admin_product_unit_edit_submit(
    request: Request,
    product_unit_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    pu = await db.get(ProductUnit, uuid.UUID(product_unit_id))
    if pu is None or pu.is_removed:
        raise HTTPException(status_code=404, detail="Product unit not found")

    form = await request.form()
    pu.name = (form.get("name") or "").strip() or None
    pu.abbreviation = (form.get("abbreviation") or "").strip() or None
    pu.product_unit_tax_id = (form.get("product_unit_tax_id") or "").strip() or None
    await db.flush()

    log_text = form.get("log") or ""
    if log_text:
        db.add(Log(
            record_id=pu.id,
            table_name="product_units",
            description=log_text,
            created_by_user_id=current_user.id,
            type="Update",
        ))
    return RedirectResponse(url="/administration/product-units", status_code=303)


@router.get("/product-units/{product_unit_id}", response_class=HTMLResponse)
async def admin_product_unit_detail(
    request: Request,
    product_unit_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    pu = await db.get(ProductUnit, uuid.UUID(product_unit_id))
    if pu is None or pu.is_removed:
        raise HTTPException(status_code=404, detail="Product unit not found")
    return templates.TemplateResponse("admin/product_unit_detail.html", {
        "request": request, "current_user": current_user, "product_unit": pu,
    })


@router.post("/product-units/{product_unit_id}/delete", response_class=HTMLResponse)
async def admin_product_unit_delete(
    request: Request,
    product_unit_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    pu = await db.get(ProductUnit, uuid.UUID(product_unit_id))
    if pu is None or pu.is_removed:
        raise HTTPException(status_code=404, detail="Product unit not found")
    name = pu.name or pu.abbreviation or ""
    pu.is_removed = True
    await db.flush()
    db.add(Log(
        record_id=pu.id,
        table_name="product_units",
        description=f"حذف واحد اندازه‌گیری: {name}",
        created_by_user_id=current_user.id,
        type="Delete",
    ))
    return RedirectResponse(url="/administration/product-units", status_code=303)


# ── Tags ──

@router.get("/tags", response_class=HTMLResponse)
async def admin_tags(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    stmt = (
        select(Tag)
        .options(selectinload(Tag.product_tags))
        .where(Tag.is_removed == False)
        .order_by(Tag.insert_date.desc())
    )
    items = (await db.execute(stmt)).unique().scalars().all()
    return templates.TemplateResponse("admin/tags.html", {
        "request": request, "current_user": current_user, "items": items,
    })


@router.get("/tags/new", response_class=HTMLResponse)
async def admin_tag_create(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    return templates.TemplateResponse("admin/tag_form.html", {
        "request": request, "current_user": current_user, "tag": None,
    })


@router.post("/tags/new", response_class=HTMLResponse)
async def admin_tag_create_submit(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    form = await request.form()
    name = (form.get("name") or "").strip()

    tag = Tag(name=name, created_by_user_id=current_user.id)
    db.add(tag)
    await db.flush()

    log_text = form.get("log") or ""
    if log_text:
        db.add(Log(
            record_id=tag.id,
            table_name="tags",
            description=log_text,
            created_by_user_id=current_user.id,
            type="Create",
        ))
    return RedirectResponse(url="/administration/tags", status_code=303)


@router.get("/tags/{tag_id}/edit", response_class=HTMLResponse)
async def admin_tag_edit(
    request: Request,
    tag_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    tag = await db.get(Tag, uuid.UUID(tag_id))
    if tag is None or tag.is_removed:
        raise HTTPException(status_code=404, detail="Tag not found")
    return templates.TemplateResponse("admin/tag_form.html", {
        "request": request, "current_user": current_user, "tag": tag,
    })


@router.post("/tags/{tag_id}/edit", response_class=HTMLResponse)
async def admin_tag_edit_submit(
    request: Request,
    tag_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    tag = await db.get(Tag, uuid.UUID(tag_id))
    if tag is None or tag.is_removed:
        raise HTTPException(status_code=404, detail="Tag not found")

    form = await request.form()
    tag.name = (form.get("name") or "").strip()
    await db.flush()

    log_text = form.get("log") or ""
    if log_text:
        db.add(Log(
            record_id=tag.id,
            table_name="tags",
            description=log_text,
            created_by_user_id=current_user.id,
            type="Update",
        ))
    return RedirectResponse(url="/administration/tags", status_code=303)


@router.get("/tags/{tag_id}", response_class=HTMLResponse)
async def admin_tag_detail(
    request: Request,
    tag_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    tag = await db.get(Tag, uuid.UUID(tag_id))
    if tag is None or tag.is_removed:
        raise HTTPException(status_code=404, detail="Tag not found")
    return templates.TemplateResponse("admin/tag_detail.html", {
        "request": request, "current_user": current_user, "tag": tag,
    })


@router.post("/tags/{tag_id}/delete", response_class=HTMLResponse)
async def admin_tag_delete(
    request: Request,
    tag_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    tag = await db.get(Tag, uuid.UUID(tag_id))
    if tag is None or tag.is_removed:
        raise HTTPException(status_code=404, detail="Tag not found")
    name = tag.name
    tag.is_removed = True
    await db.flush()
    db.add(Log(
        record_id=tag.id,
        table_name="tags",
        description=f"حذف برچسب: {name}",
        created_by_user_id=current_user.id,
        type="Delete",
    ))
    return RedirectResponse(url="/administration/tags", status_code=303)


# ── Discounts ──

@router.get("/discounts", response_class=HTMLResponse)
async def admin_discounts(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Orders Manager")),
    db: AsyncSession = Depends(get_db),
):
    items = (await db.execute(
        select(Discount).where(Discount.is_removed == False).order_by(Discount.insert_date.desc()).limit(100)
    )).scalars().all()
    return templates.TemplateResponse("admin/generic_list.html", {
        "request": request, "current_user": current_user,
        "title": "تخفیف‌ها", "items": items,
        "columns": [
            {"key": "code", "label": "کد تخفیف"},
            {"key": "percent", "label": "درصد"},
            {"key": "amount", "label": "مبلغ"},
            {"key": "is_enable", "label": "فعال"},
        ],
        "create_url": "/administration/discounts/new", "create_label": "تخفیف جدید",
        "edit_url": "/administration/discounts",
    })


@router.get("/discounts/new", response_class=HTMLResponse)
async def admin_discount_create(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Orders Manager")),
    db: AsyncSession = Depends(get_db),
):
    return templates.TemplateResponse("admin/generic_form.html", {
        "request": request, "current_user": current_user,
        "title": "تخفیف جدید",
        "fields": [
            {"name": "code", "label": "کد تخفیف", "type": "text", "required": True},
            {"name": "percent", "label": "درصد تخفیف", "type": "number"},
            {"name": "amount", "label": "مبلغ تخفیف", "type": "number"},
            {"name": "is_enable", "label": "فعال", "type": "checkbox"},
        ],
    })


# ── Pay Methods ──

PAY_METHOD_TYPES = {
    "Zarinpal": "زرین‌پال",
    "BankReceipt": "فیش بانکی",
}


def _pay_method_type_name(value) -> str:
    return PAY_METHOD_TYPES.get(value, value or "-")


@router.get("/pay-methods", response_class=HTMLResponse)
async def admin_pay_methods(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Orders Manager")),
    db: AsyncSession = Depends(get_db),
):
    items = (await db.execute(
        select(PayMethod).where(PayMethod.is_removed == False).order_by(PayMethod.insert_date.desc()).limit(200)
    )).scalars().all()
    return templates.TemplateResponse("admin/pay_methods.html", {
        "request": request, "current_user": current_user, "items": items,
    })


@router.get("/pay-methods/create", response_class=HTMLResponse)
async def admin_pay_method_create_form(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Orders Manager")),
):
    return templates.TemplateResponse("admin/pay_method_form.html", {
        "request": request, "current_user": current_user,
        "pay_method": None, "pay_method_types": PAY_METHOD_TYPES,
    })


@router.post("/pay-methods/create", response_class=HTMLResponse)
async def admin_pay_method_create_submit(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Orders Manager")),
    db: AsyncSession = Depends(get_db),
):
    form = await request.form()
    name = (form.get("name") or "").strip()
    if not name:
        raise HTTPException(status_code=400, detail="نام الزامی است")
    type_value = (form.get("type") or "").strip()
    if type_value not in PAY_METHOD_TYPES:
        type_value = type_value or None
    enable = (form.get("enable") or "") in ("on", "true", "1", "True")

    pm = PayMethod(
        id=uuid.uuid4(),
        name=name,
        enable=enable,
        type=type_value,
        description=(form.get("description") or "").strip(),
        created_by_user_id=current_user.id,
        insert_date=datetime.now(timezone.utc),
        update_date=datetime.now(timezone.utc),
    )
    db.add(pm)
    db.add(Log(record_id=pm.id, table_name="pay_methods",
               description=f"ایجاد روش پرداخت: {pm.name}",
               created_by_user_id=current_user.id, type="Create"))
    await db.commit()
    return RedirectResponse(url="/administration/pay-methods", status_code=303)


@router.get("/pay-methods/{pay_method_id}", response_class=HTMLResponse)
async def admin_pay_method_detail(
    request: Request,
    pay_method_id: str,
    current_user: User = Depends(require_any_role("Admin", "Orders Manager")),
    db: AsyncSession = Depends(get_db),
):
    try:
        pid = uuid.UUID(pay_method_id)
    except ValueError:
        raise HTTPException(status_code=404, detail="Invalid pay method ID")
    pm = await db.get(PayMethod, pid)
    if not pm or pm.is_removed:
        raise HTTPException(status_code=404, detail="Pay method not found")
    return templates.TemplateResponse("admin/pay_method_detail.html", {
        "request": request, "current_user": current_user,
        "pay_method": pm, "pay_method_type_name": _pay_method_type_name(pm.type),
    })


@router.get("/pay-methods/{pay_method_id}/edit", response_class=HTMLResponse)
async def admin_pay_method_edit_form(
    request: Request,
    pay_method_id: str,
    current_user: User = Depends(require_any_role("Admin", "Orders Manager")),
    db: AsyncSession = Depends(get_db),
):
    try:
        pid = uuid.UUID(pay_method_id)
    except ValueError:
        raise HTTPException(status_code=404, detail="Invalid pay method ID")
    pm = await db.get(PayMethod, pid)
    if not pm or pm.is_removed:
        raise HTTPException(status_code=404, detail="Pay method not found")
    return templates.TemplateResponse("admin/pay_method_form.html", {
        "request": request, "current_user": current_user,
        "pay_method": pm, "pay_method_types": PAY_METHOD_TYPES,
    })


@router.post("/pay-methods/{pay_method_id}/edit", response_class=HTMLResponse)
async def admin_pay_method_edit_submit(
    request: Request,
    pay_method_id: str,
    current_user: User = Depends(require_any_role("Admin", "Orders Manager")),
    db: AsyncSession = Depends(get_db),
):
    try:
        pid = uuid.UUID(pay_method_id)
    except ValueError:
        raise HTTPException(status_code=404, detail="Invalid pay method ID")
    pm = await db.get(PayMethod, pid)
    if not pm or pm.is_removed:
        raise HTTPException(status_code=404, detail="Pay method not found")
    form = await request.form()
    name = (form.get("name") or "").strip()
    if name:
        pm.name = name
    type_value = (form.get("type") or "").strip()
    if type_value in PAY_METHOD_TYPES:
        pm.type = type_value
    pm.enable = (form.get("enable") or "") in ("on", "true", "1", "True")
    pm.description = (form.get("description") or "").strip()
    pm.update_date = datetime.now(timezone.utc)
    db.add(Log(record_id=pm.id, table_name="pay_methods",
               description=f"ویرایش روش پرداخت: {pm.name}",
               created_by_user_id=current_user.id, type="Update"))
    await db.commit()
    return RedirectResponse(url=f"/administration/pay-methods/{pm.id}", status_code=303)


@router.post("/pay-methods/{pay_method_id}/delete", response_class=HTMLResponse)
async def admin_pay_method_delete(
    request: Request,
    pay_method_id: str,
    current_user: User = Depends(require_any_role("Admin", "Orders Manager")),
    db: AsyncSession = Depends(get_db),
):
    try:
        pid = uuid.UUID(pay_method_id)
    except ValueError:
        raise HTTPException(status_code=404, detail="Invalid pay method ID")
    pm = await db.get(PayMethod, pid)
    if not pm or pm.is_removed:
        raise HTTPException(status_code=404, detail="Pay method not found")
    name = pm.name
    pm.is_removed = True
    pm.update_date = datetime.now(timezone.utc)
    db.add(Log(record_id=pm.id, table_name="pay_methods",
               description=f"حذف روش پرداخت: {name}",
               created_by_user_id=current_user.id, type="Delete"))
    await db.commit()
    return RedirectResponse(url="/administration/pay-methods", status_code=303)


# ── Post Types ──

_POST_TYPE_UPLOAD_DIR = "app/static/uploads/post_types"


async def _save_post_type_image(file) -> str | None:
    """Save an uploaded image and return its public URL (or None)."""
    filename = getattr(file, "filename", "")
    if not filename:
        return None
    content = await file.read()
    if not content:
        return None
    ext = (filename.split(".")[-1] or "jpg").lower()
    if ext not in {"jpg", "jpeg", "png", "gif", "webp"}:
        ext = "jpg"
    os.makedirs(_POST_TYPE_UPLOAD_DIR, exist_ok=True)
    fname = f"pt_{uuid.uuid4().hex[:12]}.{ext}"
    import aiofiles
    async with aiofiles.open(f"{_POST_TYPE_UPLOAD_DIR}/{fname}", "wb") as f:
        await f.write(content)
    return f"/static/uploads/post_types/{fname}"


def _remove_post_type_image(image_url: str | None):
    if not image_url:
        return
    rel = image_url.lstrip("/").replace("static/", "", 1)
    path = os.path.join("app", "static", "uploads", "post_types", os.path.basename(image_url))
    try:
        if os.path.exists(path):
            os.remove(path)
    except Exception:
        pass


@router.get("/post-types", response_class=HTMLResponse)
async def admin_post_types(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Orders Manager")),
    db: AsyncSession = Depends(get_db),
):
    items = (await db.execute(
        select(PostType).where(PostType.is_removed == False).order_by(PostType.insert_date.desc()).limit(200)
    )).scalars().all()
    return templates.TemplateResponse("admin/post_types.html", {
        "request": request, "current_user": current_user, "items": items,
    })


@router.get("/post-types/create", response_class=HTMLResponse)
async def admin_post_type_create_form(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Orders Manager")),
):
    return templates.TemplateResponse("admin/post_type_form.html", {
        "request": request, "current_user": current_user, "post_type": None,
    })


@router.post("/post-types/create", response_class=HTMLResponse)
async def admin_post_type_create_submit(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Orders Manager")),
    db: AsyncSession = Depends(get_db),
):
    form = await request.form()
    name = (form.get("name") or "").strip()
    if not name:
        raise HTTPException(status_code=400, detail="نام الزامی است")

    def _num(val):
        try:
            return float(val) if (val or "").strip() else 0.0
        except ValueError:
            return 0.0

    image_file = form.get("image")
    image_url = await _save_post_type_image(image_file) if image_file else None

    pt = PostType(
        id=uuid.uuid4(),
        name=name,
        site=(form.get("site") or "").strip() or None,
        price=_num(form.get("price")),
        post_vat_rate=_num(form.get("post_vat_rate")),
        description=(form.get("description") or "").strip(),
        image_url=image_url,
        created_by_user_id=current_user.id,
        insert_date=datetime.now(timezone.utc),
        update_date=datetime.now(timezone.utc),
    )
    db.add(pt)
    db.add(Log(record_id=pt.id, table_name="post_types",
               description=f"ایجاد نوع ارسال: {pt.name}",
               created_by_user_id=current_user.id, type="Create"))
    await db.commit()
    return RedirectResponse(url="/administration/post-types", status_code=303)


@router.get("/post-types/{post_type_id}", response_class=HTMLResponse)
async def admin_post_type_detail(
    request: Request,
    post_type_id: str,
    current_user: User = Depends(require_any_role("Admin", "Orders Manager")),
    db: AsyncSession = Depends(get_db),
):
    try:
        pid = uuid.UUID(post_type_id)
    except ValueError:
        raise HTTPException(status_code=404, detail="Invalid post type ID")
    pt = await db.get(PostType, pid)
    if not pt or pt.is_removed:
        raise HTTPException(status_code=404, detail="Post type not found")
    return templates.TemplateResponse("admin/post_type_detail.html", {
        "request": request, "current_user": current_user, "post_type": pt,
    })


@router.get("/post-types/{post_type_id}/edit", response_class=HTMLResponse)
async def admin_post_type_edit_form(
    request: Request,
    post_type_id: str,
    current_user: User = Depends(require_any_role("Admin", "Orders Manager")),
    db: AsyncSession = Depends(get_db),
):
    try:
        pid = uuid.UUID(post_type_id)
    except ValueError:
        raise HTTPException(status_code=404, detail="Invalid post type ID")
    pt = await db.get(PostType, pid)
    if not pt or pt.is_removed:
        raise HTTPException(status_code=404, detail="Post type not found")
    return templates.TemplateResponse("admin/post_type_form.html", {
        "request": request, "current_user": current_user, "post_type": pt,
    })


@router.post("/post-types/{post_type_id}/edit", response_class=HTMLResponse)
async def admin_post_type_edit_submit(
    request: Request,
    post_type_id: str,
    current_user: User = Depends(require_any_role("Admin", "Orders Manager")),
    db: AsyncSession = Depends(get_db),
):
    try:
        pid = uuid.UUID(post_type_id)
    except ValueError:
        raise HTTPException(status_code=404, detail="Invalid post type ID")
    pt = await db.get(PostType, pid)
    if not pt or pt.is_removed:
        raise HTTPException(status_code=404, detail="Post type not found")
    form = await request.form()

    def _num(val):
        try:
            return float(val) if (val or "").strip() else 0.0
        except ValueError:
            return 0.0

    name = (form.get("name") or "").strip()
    if name:
        pt.name = name
    pt.site = (form.get("site") or "").strip() or None
    pt.price = _num(form.get("price"))
    pt.post_vat_rate = _num(form.get("post_vat_rate"))
    pt.description = (form.get("description") or "").strip()

    image_file = form.get("image")
    if image_file and getattr(image_file, "filename", ""):
        new_url = await _save_post_type_image(image_file)
        if new_url:
            _remove_post_type_image(pt.image_url)
            pt.image_url = new_url

    pt.update_date = datetime.now(timezone.utc)
    db.add(Log(record_id=pt.id, table_name="post_types",
               description=f"ویرایش نوع ارسال: {pt.name}",
               created_by_user_id=current_user.id, type="Update"))
    await db.commit()
    return RedirectResponse(url=f"/administration/post-types/{pt.id}", status_code=303)


@router.post("/post-types/{post_type_id}/delete", response_class=HTMLResponse)
async def admin_post_type_delete(
    request: Request,
    post_type_id: str,
    current_user: User = Depends(require_any_role("Admin", "Orders Manager")),
    db: AsyncSession = Depends(get_db),
):
    try:
        pid = uuid.UUID(post_type_id)
    except ValueError:
        raise HTTPException(status_code=404, detail="Invalid post type ID")
    pt = await db.get(PostType, pid)
    if not pt or pt.is_removed:
        raise HTTPException(status_code=404, detail="Post type not found")
    name = pt.name
    pt.is_removed = True
    pt.update_date = datetime.now(timezone.utc)
    db.add(Log(record_id=pt.id, table_name="post_types",
               description=f"حذف نوع ارسال: {name}",
               created_by_user_id=current_user.id, type="Delete"))
    await db.commit()
    return RedirectResponse(url="/administration/post-types", status_code=303)


# ── Media ──

@router.get("/media", response_class=HTMLResponse)
async def admin_media(
    request: Request,
    current_user: User = Depends(get_admin_user),
    db: AsyncSession = Depends(get_db),
):
    items = (await db.execute(
        select(Media).where(Media.is_removed == False).order_by(Media.insert_date.desc()).limit(100)
    )).scalars().all()
    return templates.TemplateResponse("admin/generic_list.html", {
        "request": request, "current_user": current_user,
        "title": "رسانه‌ها", "items": items,
        "columns": [
            {"key": "title", "label": "عنوان"},
            {"key": "type", "label": "نوع"},
        ],
    })


# ── Manufacturers ──

@router.get("/manufacturers", response_class=HTMLResponse)
async def admin_manufacturers(
    request: Request,
    current_user: User = Depends(get_admin_user),
    db: AsyncSession = Depends(get_db),
):
    items = (await db.execute(
        select(Manufacturer).where(Manufacturer.is_removed == False).order_by(Manufacturer.insert_date.desc()).limit(100)
    )).scalars().all()
    return templates.TemplateResponse("admin/generic_list.html", {
        "request": request, "current_user": current_user,
        "title": "تولیدکنندگان", "items": items,
        "columns": [
            {"key": "name", "label": "نام"},
            {"key": "telephone", "label": "تلفن"},
            {"key": "email", "label": "ایمیل"},
        ],
    })


# ── Admin Parameters ──

_admin_param_defaults = {"ConfirmOrderPN": "09930003120", "ConfrimOrderEm": "hamdoos@outlook.com"}


def _admin_param_items(value: str | None) -> list[str]:
    if not value:
        return []
    return [x.strip() for x in value.split(";") if x.strip()]


def _validate_admin_param(confirm_pn: str, confirm_em: str) -> list[str]:
    """Validate ';'-separated phone numbers / emails. Mirrors .NET Tools checks."""
    errors: list[str] = []
    for num in _admin_param_items(confirm_pn):
        if not is_phone_number(num):
            errors.append("شماره تلفن تأیید سفارش معتبر نیست")
            break
    for email in _admin_param_items(confirm_em):
        if not is_email(email):
            errors.append("ایمیل تأیید سفارش معتبر نیست")
            break
    return errors


async def _get_admin_param(db: AsyncSession) -> AdminParameter | None:
    return (await db.execute(
        select(AdminParameter)
        .where(AdminParameter.is_removed == False)
        .order_by(AdminParameter.insert_date.asc())
    )).scalars().first()


async def _get_admin_param_by_id(db, param_id) -> AdminParameter | None:
    try:
        pid = uuid.UUID(param_id)
    except ValueError:
        return None
    return (await db.execute(
        select(AdminParameter).where(AdminParameter.id == pid, AdminParameter.is_removed == False)
    )).scalars().first()


@router.get("/admin-parameters", response_class=HTMLResponse)
async def admin_parameters_details(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    param = await _get_admin_param(db)
    return templates.TemplateResponse("admin/admin_parameter_details.html", {
        "request": request, "current_user": current_user, "param": param,
    })


@router.get("/admin-parameters/create", response_class=HTMLResponse)
async def admin_parameters_create_page(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    return templates.TemplateResponse("admin/admin_parameter_form.html", {
        "request": request, "current_user": current_user,
        "param": None, "action_url": "/administration/admin-parameters/create",
        "form_title": "ایجاد", "errors": [],
    })


@router.post("/admin-parameters/create", response_class=HTMLResponse)
async def admin_parameters_create(
    request: Request,
    ConfirmOrderPN: str = Form(""),
    ConfrimOrderEm: str = Form(""),
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    errors = _validate_admin_param(ConfirmOrderPN, ConfrimOrderEm)
    if errors:
        return templates.TemplateResponse("admin/admin_parameter_form.html", {
            "request": request, "current_user": current_user,
            "param": None, "action_url": "/administration/admin-parameters/create",
            "form_title": "ایجاد", "errors": errors,
            "ConfirmOrderPN": ConfirmOrderPN, "ConfrimOrderEm": ConfrimOrderEm,
        })
    param = AdminParameter(ConfirmOrderPN=ConfirmOrderPN, ConfrimOrderEm=ConfrimOrderEm)
    db.add(param)
    db.add(Log(record_id=param.id, table_name="admin_parameters",
               description=f"ConfirmOrderPN: {ConfirmOrderPN}\nConfrimOrderEm: {ConfrimOrderEm}",
               created_by_user_id=current_user.id, type="Create"))
    await db.commit()
    return RedirectResponse(url="/administration/admin-parameters", status_code=303)


@router.get("/admin-parameters/{param_id}/edit", response_class=HTMLResponse)
async def admin_parameters_edit_page(
    param_id: str,
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    param = await _get_admin_param_by_id(db, param_id)
    if not param:
        raise HTTPException(status_code=404, detail="Admin parameter not found")
    return templates.TemplateResponse("admin/admin_parameter_form.html", {
        "request": request, "current_user": current_user,
        "param": param, "action_url": f"/administration/admin-parameters/{param_id}/edit",
        "form_title": "ویرایش", "errors": [],
    })


@router.post("/admin-parameters/{param_id}/edit", response_class=HTMLResponse)
async def admin_parameters_edit(
    param_id: str,
    request: Request,
    ConfirmOrderPN: str = Form(""),
    ConfrimOrderEm: str = Form(""),
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    param = await _get_admin_param_by_id(db, param_id)
    if not param:
        raise HTTPException(status_code=404, detail="Admin parameter not found")
    errors = _validate_admin_param(ConfirmOrderPN, ConfrimOrderEm)
    if errors:
        return templates.TemplateResponse("admin/admin_parameter_form.html", {
            "request": request, "current_user": current_user,
            "param": param, "action_url": f"/administration/admin-parameters/{param_id}/edit",
            "form_title": "ویرایش", "errors": errors,
            "ConfirmOrderPN": ConfirmOrderPN, "ConfrimOrderEm": ConfrimOrderEm,
        })
    old_pn = param.ConfirmOrderPN or ""
    old_em = param.ConfrimOrderEm or ""
    param.ConfirmOrderPN = ConfirmOrderPN
    param.ConfrimOrderEm = ConfrimOrderEm
    param.update_date = datetime.now(timezone.utc)
    db.add(param)

    update_lines = []
    if old_pn != ConfirmOrderPN:
        update_lines.append(f"ConfirmOrderPN : {old_pn} --> {ConfirmOrderPN}")
    if old_em != ConfrimOrderEm:
        update_lines.append(f"ConfrimOrderEm : {old_em} --> {ConfrimOrderEm}")
    if update_lines:
        db.add(Log(record_id=param.id, table_name="admin_parameters",
                   description="\n".join(update_lines),
                   created_by_user_id=current_user.id, type="Update"))
    await db.commit()
    return RedirectResponse(url="/administration/admin-parameters", status_code=303)


@router.post("/admin-parameters/{param_id}/delete", response_class=HTMLResponse)
async def admin_parameters_delete(
    param_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    param = await _get_admin_param_by_id(db, param_id)
    if not param:
        raise HTTPException(status_code=404, detail="Admin parameter not found")
    param.is_removed = True
    param.update_date = datetime.now(timezone.utc)
    db.add(Log(record_id=param.id, table_name="admin_parameters",
               description=f"حذف پارامتر ادمین: {param.ConfrimOrderEm}",
               created_by_user_id=current_user.id, type="Delete"))
    await db.commit()
    return RedirectResponse(url="/administration/admin-parameters", status_code=303)


@router.get("/admin-parameters/restore", response_class=HTMLResponse)
async def admin_parameters_restore(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    param = await _get_admin_param(db)
    if param is None:
        db.add(AdminParameter(
            ConfirmOrderPN=_admin_param_defaults["ConfirmOrderPN"],
            ConfrimOrderEm=_admin_param_defaults["ConfrimOrderEm"],
        ))
    else:
        param.ConfirmOrderPN = _admin_param_defaults["ConfirmOrderPN"]
        param.ConfrimOrderEm = _admin_param_defaults["ConfrimOrderEm"]
        param.update_date = datetime.now(timezone.utc)
        db.add(param)
    await db.commit()
    return RedirectResponse(url="/administration/admin-parameters", status_code=303)


# ── Logs (Logger) ──

def _log_user_map(users) -> dict:
    return {str(u.id): (u.full_name or u.username or "") for u in users}


@router.get("/logs", response_class=HTMLResponse)
async def admin_logs(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    """Logger index matching .NET — filter section + paginated table."""
    q = request.query_params
    filter_type = (q.get("type") or "").strip()          # Table_t name
    filter_u_id = (q.get("u_id") or q.get("uid") or "").strip()
    filter_log_type = (q.get("log_type") or "").strip()  # LogType_t name

    try:
        page = max(1, int(q.get("page") or 1))
    except (TypeError, ValueError):
        page = 1
    try:
        page_size = int(q.get("page_size") or 25)
        if page_size <= 0:
            page_size = 25
    except (TypeError, ValueError):
        page_size = 25

    filters = [Log.is_removed == False]

    if filter_type and filter_type != "all":
        filters.append(Log.table == resolve_table_int(filter_type))
    if filter_log_type and filter_log_type != "all":
        filters.append(Log.type == resolve_type_int(filter_log_type))
    if filter_u_id and filter_u_id != "all":
        try:
            filters.append(Log.created_by_user_id == uuid.UUID(filter_u_id))
        except (ValueError, AttributeError):
            pass

    total_stmt = select(func.count(Log.id)).select_from(Log).where(*filters)
    total = (await db.execute(total_stmt)).scalar_one()

    stmt = (
        select(Log)
        .where(*filters)
        .order_by(Log.insert_date.desc())
        .offset((page - 1) * page_size)
        .limit(page_size)
    )
    logs = (await db.execute(stmt)).scalars().all()

    user_ids = {l.created_by_user_id for l in logs if l.created_by_user_id}
    users = []
    if user_ids:
        users = (await db.execute(
            select(User).where(User.id.in_(user_ids))
        )).scalars().all()
    user_map = _log_user_map(users)

    # Droplists
    all_users = (await db.execute(
        select(User).where(User.is_removed == False).order_by(User.first_name)
    )).scalars().all()

    table_options = TABLE_OPTIONS                                   # Table_t enum names
    log_type_options = list(LOG_TYPE_OPTIONS.keys())                # LogType_t enum names

    total_pages = (total + page_size - 1) // page_size if total else 1
    if page > total_pages:
        page = total_pages
    start_index = (page - 1) * page_size + 1

    rows = []
    for i, log in enumerate(logs, start=start_index):
        rows.append({
            "index": i,
            "id": str(log.id),
            "insert_date": to_farsi_full(log.insert_date),
            "table_name": log.table_name,
            "type_name": log.type_name,
            "description": log.description or "",
            "user_name": user_map.get(str(log.created_by_user_id), ""),
        })

    def page_url(p: int) -> str:
        params = {"u_id": filter_u_id, "log_type": filter_log_type, "type": filter_type, "page": p, "page_size": page_size}
        keep = "&".join(f"{k}={v}" for k, v in params.items() if v not in ("", None))
        return "/administration/logs?" + keep

    return templates.TemplateResponse("admin/logs.html", {
        "request": request, "current_user": current_user,
        "rows": rows, "total": total, "page": page, "page_size": page_size,
        "total_pages": total_pages,
        "filter_type": filter_type, "filter_u_id": filter_u_id, "filter_log_type": filter_log_type,
        "table_options": table_options,
        "user_options": all_users,
        "log_type_options": log_type_options,
        "page_url": page_url,
        "page_numbers": _page_numbers(page, total_pages),
    })


def _page_numbers(page: int, total_pages: int) -> list:
    if total_pages <= 1:
        return [1]
    start = max(1, min(page - 2, total_pages - 4))
    end = min(total_pages, start + 4)
    return list(range(start, end + 1))


@router.get("/logs/detail/{log_id}", response_class=HTMLResponse)
async def admin_log_detail(
    log_id: str,
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    """Single log detail page matching .NET Logger/Details."""
    try:
        lid = uuid.UUID(log_id)
    except ValueError:
        raise HTTPException(status_code=404, detail="Invalid log ID")
    log = (await db.execute(select(Log).where(Log.id == lid, Log.is_removed == False))).scalars().first()
    if not log:
        raise HTTPException(status_code=404, detail="Log not found")

    user = None
    if log.created_by_user_id:
        user = (await db.execute(select(User).where(User.id == log.created_by_user_id))).scalars().first()

    lines = [ln for ln in (log.description or "").split("\n") if ln.strip()]

    return templates.TemplateResponse("admin/log_detail.html", {
        "request": request, "current_user": current_user,
        "log": log,
        "user_name": (user.full_name or user.username or "") if user else "",
        "insert_date": to_farsi_full(log.insert_date),
        "lines": lines,
    })


@router.post("/logs/{log_id}/delete", response_class=HTMLResponse)
async def admin_log_delete(
    log_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    try:
        lid = uuid.UUID(log_id)
    except ValueError:
        raise HTTPException(status_code=404, detail="Invalid log ID")
    log = (await db.execute(select(Log).where(Log.id == lid, Log.is_removed == False))).scalars().first()
    if not log:
        raise HTTPException(status_code=404, detail="Log not found")
    log.is_removed = True
    log.update_date = datetime.now(timezone.utc)
    db.add(Log(record_id=log.id, table_name="logs",
               description=f"حذف لاگ: {log.type_name} - {log.table_name}",
               created_by_user_id=current_user.id, type="Delete"))
    await db.commit()
    return RedirectResponse(url="/administration/logs", status_code=303)


# ── Logger (سوابق) dedicated page matching .NET ──

@router.get("/logs/{record_id}", response_class=HTMLResponse)
async def admin_logger(
    request: Request,
    record_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    from app.models.common import Log
    from app.models.identity import User

    # Build the base query
    filters = [Log.is_removed == False]

    # Record ID filter (route param)
    cid = uuid.UUID(record_id)
    filters.append(Log.record_id == cid)

    # Optional query filters
    uid = request.query_params.get("uId") or request.query_params.get("user_id") or ""
    log_type = request.query_params.get("log_type") or request.query_params.get("type") or ""
    table_name = "categories"

    try:
        if uid and uid != "all":
            filters.append(Log.created_by_user_id == uuid.UUID(uid))
    except (ValueError, AttributeError):
        pass

    if log_type and log_type != "all":
        filters.append(Log.type == resolve_type_int(log_type))

    # Fetch logs
    stmt = (
        select(Log)
        .where(*filters)
        .order_by(Log.insert_date.desc())
    )
    result = await db.execute(stmt)
    logs = result.unique().scalars().all()

    # Fetch all users for the filter dropdown
    users_result = await db.execute(
        select(User).where(User.is_removed == False).order_by(User.first_name)
    )
    users = users_result.scalars().all()

    # Get the record type from the first log entry
    if logs:
        table_name = logs[0].table_name or "categories"

    # Fetch the related user who created the record
    record_user = None
    for log in logs:
        if log.created_by_user_id:
            record_user = next((u for u in users if u.id == log.created_by_user_id), None)
            if record_user:
                break

    return templates.TemplateResponse("admin/logger.html", {
        "request": request, "current_user": current_user,
        "logs": logs, "users": users,
        "record_id": record_id,
        "table_name": table_name,
        "filter_user_id": uid,
        "filter_log_type": log_type,
        "record_user": record_user,
        "user_map": {str(u.id): u.full_name or u.username or "" for u in users},
    })


# ── Notified Products ──

async def _load_notified_products(db: AsyncSession, np_ids: list | None = None):
    stmt = (
        select(NotifiedProduct)
        .options(
            selectinload(NotifiedProduct.variety).selectinload(Variety.product),
            selectinload(NotifiedProduct.variety).selectinload(Variety.product_varieties).selectinload(ProductVariety.category_option),
            selectinload(NotifiedProduct.created_by_user),
        )
        .where(NotifiedProduct.is_removed == False)
    )
    if np_ids:
        stmt = stmt.where(NotifiedProduct.id.in_(np_ids))
    return (await db.execute(stmt)).scalars().all()


def _notified_type_filter(np: NotifiedProduct, notify_type: str) -> bool:
    """Mirror .NET GetFilteredNotifyProducts Type filter. Returns True when kept."""
    variety = np.variety
    supply_date = variety.stock_supply_date if variety else None
    if not supply_date:
        return False
    if notify_type == "SmsNotification":
        return np.insert_date <= supply_date and (np.sms_response_date is not None and np.sms_response_date < supply_date)
    if notify_type == "EmailNotification":
        return np.insert_date <= supply_date and (np.email_response_date is not None and np.email_response_date < supply_date)
    return True


def _notified_sort_key(np: NotifiedProduct, arrange: str):
    key = {
        "InsertDate": lambda x: x.insert_date,
        "SmsResponseDate": lambda x: x.sms_response_date or x.insert_date,
        "EmailResponseDate": lambda x: x.email_response_date or x.insert_date,
        "UserName": lambda x: (x.created_by_user.full_name or "") if x.created_by_user else "",
        "Product": lambda x: (x.variety.product.name or "") if x.variety and x.variety.product else "",
        "SupplyDate": lambda x: (x.variety.product.stock_supply_date or x.insert_date) if x.variety and x.variety.product else x.insert_date,
    }.get(arrange, lambda x: x.insert_date)
    value = key(np)
    return value.replace("", "\uffff") if isinstance(value, str) else (value or value)


def _notified_variety_values(np: NotifiedProduct) -> str:
    """Mirror _setVarietyValues: 'Name: Value , Name: Value'."""
    if not np.variety:
        return "—"
    pvs = sorted(np.variety.product_varieties, key=lambda pv: pv.category_option.name if pv.category_option else "")
    parts = []
    for pv in pvs:
        if pv.category_option:
            parts.append(f"{pv.category_option.name}: {pv.value or ''}")
    return " , ".join(parts) if parts else "—"


def _notified_row(np: NotifiedProduct, detail_url: bool = True) -> dict:
    variety = np.variety
    product = variety.product if variety else None
    return {
        "id": str(np.id),
        "insert_date": _to_fa_datetime(np.insert_date),
        "sms_response_date": _to_fa_datetime(np.sms_response_date),
        "email_response_date": _to_fa_datetime(np.email_response_date),
        "product_name": product.name if product else "—",
        "supply_date": _to_fa_datetime(product.stock_supply_date) if product else "—",
        "created_by": (np.created_by_user.full_name or np.created_by_user.username) if np.created_by_user else "—",
        "variety_values": _notified_variety_values(np),
    }


@router.get("/notified-products", response_class=HTMLResponse)
async def admin_notified_products(
    request: Request,
    notify_type: str = Query(""),
    arrange: str = Query("InsertDate"),
    desc: bool = Query(True),
    page: int = Query(1),
    page_size: int = Query(25),
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    items = await _load_notified_products(db)

    if notify_type:
        items = [np for np in items if _notified_type_filter(np, notify_type)]

    sorted_items = sorted(items, key=lambda np: _notified_sort_key(np, arrange), reverse=desc)

    total = len(sorted_items)
    total_pages = max(1, (total + page_size - 1) // page_size)
    page = max(1, min(page, total_pages))
    start = (page - 1) * page_size
    page_items = sorted_items[start:start + page_size]

    rows = [_notified_row(np) for np in page_items]

    return templates.TemplateResponse("admin/notified_products.html", {
        "request": request, "current_user": current_user,
        "items": rows,
        "total": total, "page": page, "total_pages": total_pages, "page_size": page_size,
        "notify_type": notify_type, "arrange": arrange, "desc": desc,
    })


@router.post("/notified-products/send-sms", response_class=HTMLResponse)
async def admin_notified_products_send_sms(
    request: Request,
    notify_ids: str = Form(""),
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    try:
        ids = [u.strip() for u in notify_ids.split(",") if u.strip()]
        items = await _load_notified_products(db, ids)
        sent = 0
        for np in items:
            variety = np.variety
            product = variety.product if variety else None
            user = np.created_by_user
            if not (variety and product and user and user.phone_number):
                continue
            from app.services.sms_service import SelectedSmsSender
            sms = SelectedSmsSender()
            await sms.send_notify_product(user.phone_number, user.full_name, product.name, product.part_number or "")
            np.sms_response_date = datetime.now(timezone.utc)
            sent += 1
        await db.commit()
        msg = "sms_sent" if sent else "none"
        return RedirectResponse(url=f"/administration/notified-products?msg={msg}", status_code=303)
    except Exception:
        await db.rollback()
        return RedirectResponse(url="/administration/notified-products?msg=error", status_code=303)


@router.post("/notified-products/send-email", response_class=HTMLResponse)
async def admin_notified_products_send_email(
    request: Request,
    notify_ids: str = Form(""),
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    try:
        ids = [u.strip() for u in notify_ids.split(",") if u.strip()]
        items = await _load_notified_products(db, ids)
        sent = 0
        for np in items:
            variety = np.variety
            product = variety.product if variety else None
            user = np.created_by_user
            if not (variety and product and user and user.email):
                continue
            from app.services.email_service import EmailSender
            email = EmailSender()
            product_url = f"/products/{product.slug or product.id}"
            await email.send_notify_product(user.email, user.full_name, product.name, product_url)
            np.email_response_date = datetime.now(timezone.utc)
            sent += 1
        await db.commit()
        msg = "email_sent" if sent else "none"
        return RedirectResponse(url=f"/administration/notified-products?msg={msg}", status_code=303)
    except Exception:
        await db.rollback()
        return RedirectResponse(url="/administration/notified-products?msg=error", status_code=303)


@router.get("/notified-products/{np_id}/details", response_class=HTMLResponse)
async def admin_notified_products_details(
    np_id: str,
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    try:
        uid = uuid.UUID(np_id)
    except ValueError:
        raise HTTPException(status_code=404, detail="Notified product not found")
    items = await _load_notified_products(db, [uid])
    if not items:
        raise HTTPException(status_code=404, detail="Notified product not found")
    np = items[0]
    detail = _notified_row(np)
    detail.update({"variety_values": _notified_variety_values(np), "record_id": str(np.id)})
    return templates.TemplateResponse("admin/notified_product_details.html", {
        "request": request, "current_user": current_user,
        "item": detail,
        "record_id": str(np.id),
    })


# ── Price History ──

@router.get("/price-history", response_class=HTMLResponse)
async def admin_price_history(
    request: Request,
    current_user: User = Depends(get_admin_user),
    db: AsyncSession = Depends(get_db),
):
    items = (await db.execute(
        select(PriceHistory).where(PriceHistory.is_removed == False).order_by(PriceHistory.insert_date.desc()).limit(200)
    )).scalars().all()
    return templates.TemplateResponse("admin/generic_list.html", {
        "request": request, "current_user": current_user,
        "title": "تاریخچه قیمت", "items": items,
        "columns": [
            {"key": "product_id", "label": "شناسه محصول"},
            {"key": "price", "label": "قیمت"},
            {"key": "insert_date", "label": "تاریخ"},
        ],
    })


# ── Chats (Admin) ──

@router.get("/chats", response_class=HTMLResponse)
async def admin_chats(
    request: Request,
    current_user: User = Depends(get_admin_user),
    db: AsyncSession = Depends(get_db),
):
    items = (await db.execute(
        select(Chat).where(Chat.is_removed == False).order_by(Chat.insert_date.desc()).limit(100)
    )).scalars().all()
    return templates.TemplateResponse("admin/generic_list.html", {
        "request": request, "current_user": current_user,
        "title": "چت‌ها", "items": items,
        "columns": [
            {"key": "title", "label": "عنوان"},
            {"key": "subject", "label": "موضوع"},
        ],
    })


# ── Identity Information ──

@router.get("/identity-informations", response_class=HTMLResponse)
async def admin_identity_informations(
    request: Request,
    page: int = Query(1),
    type_filter: str = Query(""),
    status_filter: str = Query(""),
    user_filter: str = Query(""),
    national_code_filter: str = Query(""),
    current_user: User = Depends(get_admin_user),
    db: AsyncSession = Depends(get_db),
):
    from app.models.enums import IdentityType, IdentityStatus
    items, total = await identity_service.get_identity_infos_with_user(
        db, page, 20, type_filter, status_filter, user_filter, national_code_filter
    )
    users = (await db.execute(select(User).where(User.is_removed == False).order_by(User.first_name))).scalars().all()
    return templates.TemplateResponse("admin/identity_informations.html", {
        "request": request, "current_user": current_user,
        "items": items, "total": total, "page": page, "total_pages": (total + 19) // 20,
        "type_filter": type_filter, "status_filter": status_filter,
        "user_filter": user_filter, "national_code_filter": national_code_filter,
        "users": users, "identity_types": list(IdentityType), "identity_statuses": list(IdentityStatus),
    })


@router.get("/identity-informations/new", response_class=HTMLResponse)
async def admin_identity_information_create(
    request: Request,
    current_user: User = Depends(get_admin_user),
    db: AsyncSession = Depends(get_db),
):
    users = (await db.execute(select(User).where(User.is_removed == False).order_by(User.first_name))).scalars().all()
    return templates.TemplateResponse("admin/identity_information_form.html", {
        "request": request, "current_user": current_user,
        "info": None, "users": users,
    })


@router.post("/identity-informations/new", response_class=HTMLResponse)
async def admin_identity_information_create_submit(
    request: Request,
    current_user: User = Depends(get_admin_user),
    db: AsyncSession = Depends(get_db),
):
    form = await request.form()
    users = (await db.execute(select(User).where(User.is_removed == False).order_by(User.first_name))).scalars().all()
    try:
        info = await identity_service.create_identity_info(db, dict(form), current_user.id)
        return RedirectResponse(url="/administration/identity-informations", status_code=303)
    except ValueError as e:
        return templates.TemplateResponse("admin/identity_information_form.html", {
            "request": request, "current_user": current_user,
            "info": None, "users": users, "error": str(e), "form_data": dict(form),
        })


@router.get("/identity-informations/{info_id}/edit", response_class=HTMLResponse)
async def admin_identity_information_edit(
    request: Request,
    info_id: str,
    current_user: User = Depends(get_admin_user),
    db: AsyncSession = Depends(get_db),
):
    info = await identity_service.get_identity_info_by_id(db, uuid.UUID(info_id))
    if not info:
        raise HTTPException(status_code=404, detail="Identity information not found")
    users = (await db.execute(select(User).where(User.is_removed == False).order_by(User.first_name))).scalars().all()
    return templates.TemplateResponse("admin/identity_information_form.html", {
        "request": request, "current_user": current_user,
        "info": info, "users": users,
    })


@router.post("/identity-informations/{info_id}/edit", response_class=HTMLResponse)
async def admin_identity_information_edit_submit(
    request: Request,
    info_id: str,
    current_user: User = Depends(get_admin_user),
    db: AsyncSession = Depends(get_db),
):
    info = await identity_service.get_identity_info_by_id(db, uuid.UUID(info_id))
    if not info:
        raise HTTPException(status_code=404, detail="Identity information not found")
    form = await request.form()
    users = (await db.execute(select(User).where(User.is_removed == False).order_by(User.first_name))).scalars().all()
    try:
        await identity_service.update_identity_info(db, info, dict(form), current_user.id)
        return RedirectResponse(url="/administration/identity-informations", status_code=303)
    except ValueError as e:
        return templates.TemplateResponse("admin/identity_information_form.html", {
            "request": request, "current_user": current_user,
            "info": info, "users": users, "error": str(e), "form_data": dict(form),
        })


@router.get("/identity-informations/{info_id}", response_class=HTMLResponse)
async def admin_identity_information_detail(
    request: Request,
    info_id: str,
    current_user: User = Depends(get_admin_user),
    db: AsyncSession = Depends(get_db),
):
    info = await identity_service.get_identity_info_by_id(db, uuid.UUID(info_id))
    if not info:
        raise HTTPException(status_code=404, detail="Identity information not found")
    return templates.TemplateResponse("admin/identity_information_detail.html", {
        "request": request, "current_user": current_user, "info": info,
    })


@router.post("/identity-informations/{info_id}/accept", response_class=HTMLResponse)
async def admin_identity_information_accept(
    request: Request,
    info_id: str,
    current_user: User = Depends(get_admin_user),
    db: AsyncSession = Depends(get_db),
):
    info = await identity_service.get_identity_info_by_id(db, uuid.UUID(info_id))
    if not info:
        raise HTTPException(status_code=404, detail="Identity information not found")
    try:
        await identity_service.accept_identity_info(db, info, current_user.id)
    except ValueError as e:
        pass
    return RedirectResponse(url="/administration/identity-informations", status_code=303)


@router.post("/identity-informations/{info_id}/reject", response_class=HTMLResponse)
async def admin_identity_information_reject(
    request: Request,
    info_id: str,
    current_user: User = Depends(get_admin_user),
    db: AsyncSession = Depends(get_db),
):
    info = await identity_service.get_identity_info_by_id(db, uuid.UUID(info_id))
    if not info:
        raise HTTPException(status_code=404, detail="Identity information not found")
    try:
        await identity_service.reject_identity_info(db, info, current_user.id)
    except ValueError as e:
        pass
    return RedirectResponse(url="/administration/identity-informations", status_code=303)


@router.post("/identity-informations/{info_id}/delete", response_class=HTMLResponse)
async def admin_identity_information_delete(
    request: Request,
    info_id: str,
    current_user: User = Depends(get_admin_user),
    db: AsyncSession = Depends(get_db),
):
    info = await identity_service.get_identity_info_by_id(db, uuid.UUID(info_id))
    if not info:
        raise HTTPException(status_code=404, detail="Identity information not found")
    await identity_service.soft_delete_identity_info(db, info, current_user.id)
    return RedirectResponse(url="/administration/identity-informations", status_code=303)


# ── Similar Products ──

@router.get("/similar-products", response_class=HTMLResponse)
async def admin_similar_products(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    stmt = (
        select(SimilarProduct)
        .options(selectinload(SimilarProduct.product), selectinload(SimilarProduct.similar))
        .where(SimilarProduct.is_removed == False)
        .order_by(SimilarProduct.insert_date.desc())
    )
    items = (await db.execute(stmt)).unique().scalars().all()
    return templates.TemplateResponse("admin/similar_products.html", {
        "request": request, "current_user": current_user, "items": items,
    })


@router.get("/similar-products/new", response_class=HTMLResponse)
async def admin_similar_product_create(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    products = (await db.execute(
        select(Product).where(Product.is_removed == False).order_by(Product.name)
    )).scalars().all()
    return templates.TemplateResponse("admin/similar_product_form.html", {
        "request": request, "current_user": current_user,
        "similar_product": None, "products": products,
    })


@router.post("/similar-products/new", response_class=HTMLResponse)
async def admin_similar_product_create_submit(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    form = await request.form()
    product_id_raw = form.get("product_id")
    similar_product_id_raw = form.get("similar_product_id")
    if not product_id_raw or not similar_product_id_raw:
        raise HTTPException(status_code=400, detail="Product and similar product are required")
    if product_id_raw == similar_product_id_raw:
        raise HTTPException(status_code=400, detail="محصول و محصول مشابه نمی‌توانند یکسان باشند")

    existing = (await db.execute(
        select(SimilarProduct).where(
            SimilarProduct.product_id == uuid.UUID(product_id_raw),
            SimilarProduct.similar_product_id == uuid.UUID(similar_product_id_raw),
            SimilarProduct.is_removed == False,
        )
    )).scalar_one_or_none()
    if existing:
        raise HTTPException(status_code=400, detail="این ارتباط از قبل وجود دارد")

    sp = SimilarProduct(
        product_id=uuid.UUID(product_id_raw),
        similar_product_id=uuid.UUID(similar_product_id_raw),
        created_by_user_id=current_user.id,
    )
    db.add(sp)
    await db.flush()

    log_text = form.get("log") or ""
    if log_text:
        db.add(Log(
            record_id=sp.id,
            table_name="similar_products",
            description=log_text,
            created_by_user_id=current_user.id,
            type="Create",
        ))
    return RedirectResponse(url="/administration/similar-products", status_code=303)


@router.get("/similar-products/{similar_product_id}/edit", response_class=HTMLResponse)
async def admin_similar_product_edit(
    request: Request,
    similar_product_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    sp = await db.get(SimilarProduct, uuid.UUID(similar_product_id))
    if sp is None or sp.is_removed:
        raise HTTPException(status_code=404, detail="Similar product not found")
    products = (await db.execute(
        select(Product).where(Product.is_removed == False).order_by(Product.name)
    )).scalars().all()
    return templates.TemplateResponse("admin/similar_product_form.html", {
        "request": request, "current_user": current_user,
        "similar_product": sp, "products": products,
    })


@router.post("/similar-products/{similar_product_id}/edit", response_class=HTMLResponse)
async def admin_similar_product_edit_submit(
    request: Request,
    similar_product_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    sp = await db.get(SimilarProduct, uuid.UUID(similar_product_id))
    if sp is None or sp.is_removed:
        raise HTTPException(status_code=404, detail="Similar product not found")

    form = await request.form()
    product_id_raw = form.get("product_id")
    similar_product_id_raw = form.get("similar_product_id")
    if product_id_raw:
        sp.product_id = uuid.UUID(product_id_raw)
    if similar_product_id_raw:
        sp.similar_product_id = uuid.UUID(similar_product_id_raw)
    await db.flush()

    log_text = form.get("log") or ""
    if log_text:
        db.add(Log(
            record_id=sp.id,
            table_name="similar_products",
            description=log_text,
            created_by_user_id=current_user.id,
            type="Update",
        ))
    return RedirectResponse(url="/administration/similar-products", status_code=303)


@router.get("/similar-products/{similar_product_id}", response_class=HTMLResponse)
async def admin_similar_product_detail(
    request: Request,
    similar_product_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    stmt = (
        select(SimilarProduct)
        .options(selectinload(SimilarProduct.product), selectinload(SimilarProduct.similar))
        .where(SimilarProduct.id == uuid.UUID(similar_product_id), SimilarProduct.is_removed == False)
    )
    sp = (await db.execute(stmt)).unique().scalar_one_or_none()
    if sp is None:
        raise HTTPException(status_code=404, detail="Similar product not found")
    return templates.TemplateResponse("admin/similar_product_detail.html", {
        "request": request, "current_user": current_user, "similar_product": sp,
    })


@router.post("/similar-products/{similar_product_id}/delete", response_class=HTMLResponse)
async def admin_similar_product_delete(
    request: Request,
    similar_product_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    sp = await db.get(SimilarProduct, uuid.UUID(similar_product_id))
    if sp is None or sp.is_removed:
        raise HTTPException(status_code=404, detail="Similar product not found")
    sp.is_removed = True
    await db.flush()
    db.add(Log(
        record_id=sp.id,
        table_name="similar_products",
        description="حذف محصول مشابه",
        created_by_user_id=current_user.id,
        type="Delete",
    ))
    return RedirectResponse(url="/administration/similar-products", status_code=303)


# ── Related Products ──

@router.get("/related-products", response_class=HTMLResponse)
async def admin_related_products(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    stmt = (
        select(RelatedProduct)
        .options(selectinload(RelatedProduct.product), selectinload(RelatedProduct.relate_product))
        .where(RelatedProduct.is_removed == False)
        .order_by(RelatedProduct.insert_date.desc())
    )
    items = (await db.execute(stmt)).unique().scalars().all()
    return templates.TemplateResponse("admin/related_products.html", {
        "request": request, "current_user": current_user, "items": items,
    })


@router.get("/related-products/new", response_class=HTMLResponse)
async def admin_related_product_create(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    products = (await db.execute(
        select(Product).where(Product.is_removed == False).order_by(Product.name)
    )).scalars().all()
    return templates.TemplateResponse("admin/related_product_form.html", {
        "request": request, "current_user": current_user,
        "related_product": None, "products": products,
    })


@router.post("/related-products/new", response_class=HTMLResponse)
async def admin_related_product_create_submit(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    form = await request.form()
    product_id_raw = form.get("product_id")
    relate_product_id_raw = form.get("relate_product_id")
    if not product_id_raw or not relate_product_id_raw:
        raise HTTPException(status_code=400, detail="Product and related product are required")
    if product_id_raw == relate_product_id_raw:
        raise HTTPException(status_code=400, detail="محصول و محصول مرتبط نمی‌توانند یکسان باشند")

    rp = RelatedProduct(
        product_id=uuid.UUID(product_id_raw),
        relate_product_id=uuid.UUID(relate_product_id_raw),
        created_by_user_id=current_user.id,
    )
    db.add(rp)
    await db.flush()

    log_text = form.get("log") or ""
    if log_text:
        db.add(Log(
            record_id=rp.id,
            table_name="related_products",
            description=log_text,
            created_by_user_id=current_user.id,
            type="Create",
        ))
    return RedirectResponse(url="/administration/related-products", status_code=303)


@router.get("/related-products/{related_product_id}/edit", response_class=HTMLResponse)
async def admin_related_product_edit(
    request: Request,
    related_product_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    rp = await db.get(RelatedProduct, uuid.UUID(related_product_id))
    if rp is None or rp.is_removed:
        raise HTTPException(status_code=404, detail="Related product not found")
    products = (await db.execute(
        select(Product).where(Product.is_removed == False).order_by(Product.name)
    )).scalars().all()
    return templates.TemplateResponse("admin/related_product_form.html", {
        "request": request, "current_user": current_user,
        "related_product": rp, "products": products,
    })


@router.post("/related-products/{related_product_id}/edit", response_class=HTMLResponse)
async def admin_related_product_edit_submit(
    request: Request,
    related_product_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    rp = await db.get(RelatedProduct, uuid.UUID(related_product_id))
    if rp is None or rp.is_removed:
        raise HTTPException(status_code=404, detail="Related product not found")

    form = await request.form()
    product_id_raw = form.get("product_id")
    relate_product_id_raw = form.get("relate_product_id")
    if product_id_raw:
        rp.product_id = uuid.UUID(product_id_raw)
    if relate_product_id_raw:
        rp.relate_product_id = uuid.UUID(relate_product_id_raw)
    await db.flush()

    log_text = form.get("log") or ""
    if log_text:
        db.add(Log(
            record_id=rp.id,
            table_name="related_products",
            description=log_text,
            created_by_user_id=current_user.id,
            type="Update",
        ))
    return RedirectResponse(url="/administration/related-products", status_code=303)


@router.get("/related-products/{related_product_id}", response_class=HTMLResponse)
async def admin_related_product_detail(
    request: Request,
    related_product_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    stmt = (
        select(RelatedProduct)
        .options(selectinload(RelatedProduct.product), selectinload(RelatedProduct.relate_product))
        .where(RelatedProduct.id == uuid.UUID(related_product_id), RelatedProduct.is_removed == False)
    )
    rp = (await db.execute(stmt)).unique().scalar_one_or_none()
    if rp is None:
        raise HTTPException(status_code=404, detail="Related product not found")
    return templates.TemplateResponse("admin/related_product_detail.html", {
        "request": request, "current_user": current_user, "related_product": rp,
    })


@router.post("/related-products/{related_product_id}/delete", response_class=HTMLResponse)
async def admin_related_product_delete(
    request: Request,
    related_product_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    rp = await db.get(RelatedProduct, uuid.UUID(related_product_id))
    if rp is None or rp.is_removed:
        raise HTTPException(status_code=404, detail="Related product not found")
    rp.is_removed = True
    await db.flush()
    db.add(Log(
        record_id=rp.id,
        table_name="related_products",
        description="حذف محصول مرتبط",
        created_by_user_id=current_user.id,
        type="Delete",
    ))
    return RedirectResponse(url="/administration/related-products", status_code=303)


# ── Product Tags ──

@router.get("/product-tags", response_class=HTMLResponse)
async def admin_product_tags(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    stmt = (
        select(ProductTag)
        .options(selectinload(ProductTag.product), selectinload(ProductTag.tag))
        .where(ProductTag.is_removed == False)
        .order_by(ProductTag.insert_date.desc())
    )
    items = (await db.execute(stmt)).unique().scalars().all()
    return templates.TemplateResponse("admin/product_tags.html", {
        "request": request, "current_user": current_user, "items": items,
    })


@router.get("/product-tags/new", response_class=HTMLResponse)
async def admin_product_tag_create(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    products = (await db.execute(
        select(Product).where(Product.is_removed == False).order_by(Product.name)
    )).scalars().all()
    tags = (await db.execute(
        select(Tag).where(Tag.is_removed == False).order_by(Tag.name)
    )).scalars().all()
    return templates.TemplateResponse("admin/product_tag_form.html", {
        "request": request, "current_user": current_user,
        "product_tag": None, "products": products, "tags": tags,
    })


@router.post("/product-tags/new", response_class=HTMLResponse)
async def admin_product_tag_create_submit(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    form = await request.form()
    product_id_raw = form.get("product_id")
    tag_id_raw = form.get("tag_id")
    if not product_id_raw or not tag_id_raw:
        raise HTTPException(status_code=400, detail="Product and tag are required")

    existing = (await db.execute(
        select(ProductTag).where(
            ProductTag.product_id == uuid.UUID(product_id_raw),
            ProductTag.tag_id == uuid.UUID(tag_id_raw),
            ProductTag.is_removed == False,
        )
    )).scalar_one_or_none()
    if existing:
        raise HTTPException(status_code=400, detail="این ارتباط از قبل وجود دارد")

    pt = ProductTag(
        product_id=uuid.UUID(product_id_raw),
        tag_id=uuid.UUID(tag_id_raw),
        created_by_user_id=current_user.id,
    )
    db.add(pt)
    await db.flush()

    log_text = form.get("log") or ""
    if log_text:
        db.add(Log(
            record_id=pt.id,
            table_name="product_tags",
            description=log_text,
            created_by_user_id=current_user.id,
            type="Create",
        ))
    return RedirectResponse(url="/administration/product-tags", status_code=303)


@router.get("/product-tags/{product_tag_id}/edit", response_class=HTMLResponse)
async def admin_product_tag_edit(
    request: Request,
    product_tag_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    pt = await db.get(ProductTag, uuid.UUID(product_tag_id))
    if pt is None or pt.is_removed:
        raise HTTPException(status_code=404, detail="Product tag not found")
    products = (await db.execute(
        select(Product).where(Product.is_removed == False).order_by(Product.name)
    )).scalars().all()
    tags = (await db.execute(
        select(Tag).where(Tag.is_removed == False).order_by(Tag.name)
    )).scalars().all()
    return templates.TemplateResponse("admin/product_tag_form.html", {
        "request": request, "current_user": current_user,
        "product_tag": pt, "products": products, "tags": tags,
    })


@router.post("/product-tags/{product_tag_id}/edit", response_class=HTMLResponse)
async def admin_product_tag_edit_submit(
    request: Request,
    product_tag_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    pt = await db.get(ProductTag, uuid.UUID(product_tag_id))
    if pt is None or pt.is_removed:
        raise HTTPException(status_code=404, detail="Product tag not found")

    form = await request.form()
    product_id_raw = form.get("product_id")
    tag_id_raw = form.get("tag_id")
    if product_id_raw:
        pt.product_id = uuid.UUID(product_id_raw)
    if tag_id_raw:
        pt.tag_id = uuid.UUID(tag_id_raw)
    await db.flush()

    log_text = form.get("log") or ""
    if log_text:
        db.add(Log(
            record_id=pt.id,
            table_name="product_tags",
            description=log_text,
            created_by_user_id=current_user.id,
            type="Update",
        ))
    return RedirectResponse(url="/administration/product-tags", status_code=303)


@router.get("/product-tags/{product_tag_id}", response_class=HTMLResponse)
async def admin_product_tag_detail(
    request: Request,
    product_tag_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    stmt = (
        select(ProductTag)
        .options(selectinload(ProductTag.product), selectinload(ProductTag.tag))
        .where(ProductTag.id == uuid.UUID(product_tag_id), ProductTag.is_removed == False)
    )
    pt = (await db.execute(stmt)).unique().scalar_one_or_none()
    if pt is None:
        raise HTTPException(status_code=404, detail="Product tag not found")
    return templates.TemplateResponse("admin/product_tag_detail.html", {
        "request": request, "current_user": current_user, "product_tag": pt,
    })


@router.post("/product-tags/{product_tag_id}/delete", response_class=HTMLResponse)
async def admin_product_tag_delete(
    request: Request,
    product_tag_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    pt = await db.get(ProductTag, uuid.UUID(product_tag_id))
    if pt is None or pt.is_removed:
        raise HTTPException(status_code=404, detail="Product tag not found")
    pt.is_removed = True
    await db.flush()
    db.add(Log(
        record_id=pt.id,
        table_name="product_tags",
        description="حذف برچسب محصول",
        created_by_user_id=current_user.id,
        type="Delete",
    ))
    return RedirectResponse(url="/administration/product-tags", status_code=303)


# ── Technical Feature Enums ──

@router.get("/technical-feature-enums", response_class=HTMLResponse)
async def admin_technical_feature_enums(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    items = (await db.execute(
        select(TechnicalFeatureEnum)
        .where(TechnicalFeatureEnum.is_removed == False)
        .order_by(TechnicalFeatureEnum.persian_name)
    )).scalars().all()
    return templates.TemplateResponse("admin/technical_feature_enums.html", {
        "request": request, "current_user": current_user, "items": items,
    })


@router.get("/technical-feature-enums/new", response_class=HTMLResponse)
async def admin_technical_feature_enum_create(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    return templates.TemplateResponse("admin/technical_feature_enum_form.html", {
        "request": request, "current_user": current_user, "feature_enum": None,
    })


@router.post("/technical-feature-enums/new", response_class=HTMLResponse)
async def admin_technical_feature_enum_create_submit(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    form = await request.form()
    persian_name = (form.get("persian_name") or "").strip()
    name = (form.get("name") or "").strip()

    fe = TechnicalFeatureEnum(
        persian_name=persian_name or None,
        name=name or None,
        created_by_user_id=current_user.id,
    )
    db.add(fe)
    await db.flush()

    log_text = form.get("log") or ""
    if log_text:
        db.add(Log(
            record_id=fe.id,
            table_name="technical_feature_enums",
            description=log_text,
            created_by_user_id=current_user.id,
            type="Create",
        ))
    return RedirectResponse(url="/administration/technical-feature-enums", status_code=303)


@router.get("/technical-feature-enums/{enum_id}/edit", response_class=HTMLResponse)
async def admin_technical_feature_enum_edit(
    request: Request,
    enum_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    fe = await db.get(TechnicalFeatureEnum, uuid.UUID(enum_id))
    if fe is None or fe.is_removed:
        raise HTTPException(status_code=404, detail="Technical feature enum not found")
    return templates.TemplateResponse("admin/technical_feature_enum_form.html", {
        "request": request, "current_user": current_user, "feature_enum": fe,
    })


@router.post("/technical-feature-enums/{enum_id}/edit", response_class=HTMLResponse)
async def admin_technical_feature_enum_edit_submit(
    request: Request,
    enum_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    fe = await db.get(TechnicalFeatureEnum, uuid.UUID(enum_id))
    if fe is None or fe.is_removed:
        raise HTTPException(status_code=404, detail="Technical feature enum not found")

    form = await request.form()
    fe.persian_name = (form.get("persian_name") or "").strip() or None
    fe.name = (form.get("name") or "").strip() or None
    await db.flush()

    log_text = form.get("log") or ""
    if log_text:
        db.add(Log(
            record_id=fe.id,
            table_name="technical_feature_enums",
            description=log_text,
            created_by_user_id=current_user.id,
            type="Update",
        ))
    return RedirectResponse(url="/administration/technical-feature-enums", status_code=303)


@router.get("/technical-feature-enums/{enum_id}", response_class=HTMLResponse)
async def admin_technical_feature_enum_detail(
    request: Request,
    enum_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    fe = await db.get(TechnicalFeatureEnum, uuid.UUID(enum_id))
    if fe is None or fe.is_removed:
        raise HTTPException(status_code=404, detail="Technical feature enum not found")
    return templates.TemplateResponse("admin/technical_feature_enum_detail.html", {
        "request": request, "current_user": current_user, "feature_enum": fe,
    })


@router.post("/technical-feature-enums/{enum_id}/delete", response_class=HTMLResponse)
async def admin_technical_feature_enum_delete(
    request: Request,
    enum_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    fe = await db.get(TechnicalFeatureEnum, uuid.UUID(enum_id))
    if fe is None or fe.is_removed:
        raise HTTPException(status_code=404, detail="Technical feature enum not found")
    name = fe.name or fe.persian_name or ""
    fe.is_removed = True
    await db.flush()
    db.add(Log(
        record_id=fe.id,
        table_name="technical_feature_enums",
        description=f"حذف انوم ویژگی فنی: {name}",
        created_by_user_id=current_user.id,
        type="Delete",
    ))
    return RedirectResponse(url="/administration/technical-feature-enums", status_code=303)


# ── Category Technical Features ──

@router.get("/category-technical-features", response_class=HTMLResponse)
async def admin_category_technical_features(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    items = (await db.execute(
        select(CategoryTechnicalFeature)
        .options(selectinload(CategoryTechnicalFeature.category), selectinload(CategoryTechnicalFeature.technical_feature))
        .where(CategoryTechnicalFeature.is_removed == False)
        .order_by(CategoryTechnicalFeature.insert_date.desc())
    )).unique().scalars().all()
    return templates.TemplateResponse("admin/category_technical_features.html", {
        "request": request, "current_user": current_user, "items": items,
    })


async def _ctf_dropdowns(db: AsyncSession):
    categories = (await db.execute(
        select(Category).where(Category.is_removed == False).order_by(Category.title)
    )).scalars().all()
    technical_features = (await db.execute(
        select(TechnicalFeature).where(TechnicalFeature.is_removed == False).order_by(TechnicalFeature.fa_name)
    )).scalars().all()
    return categories, technical_features


@router.get("/category-technical-features/new", response_class=HTMLResponse)
async def admin_category_technical_feature_create(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    categories, technical_features = await _ctf_dropdowns(db)
    return templates.TemplateResponse("admin/category_technical_feature_form.html", {
        "request": request, "current_user": current_user, "ctf": None,
        "categories": categories, "technical_features": technical_features,
    })


@router.post("/category-technical-features/new", response_class=HTMLResponse)
async def admin_category_technical_feature_create_submit(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    form = await request.form()
    category_id = form.get("category_id")
    technical_feature_id = form.get("technical_feature_id")
    if not category_id or not technical_feature_id:
        raise HTTPException(status_code=400, detail="Category and feature are required")

    exists = (await db.execute(
        select(CategoryTechnicalFeature.id).where(
            CategoryTechnicalFeature.is_removed == False,
            CategoryTechnicalFeature.category_id == uuid.UUID(category_id),
            CategoryTechnicalFeature.technical_feature_id == uuid.UUID(technical_feature_id),
        )
    )).scalar_one_or_none()
    if exists:
        raise HTTPException(status_code=400, detail="این دسته‌بندی و ویژگی فنی قبلاً تعریف شده")

    ctf = CategoryTechnicalFeature(
        category_id=uuid.UUID(category_id),
        technical_feature_id=uuid.UUID(technical_feature_id),
        created_by_user_id=current_user.id,
    )
    db.add(ctf)
    await db.flush()
    db.add(Log(
        record_id=ctf.id,
        table_name="category_technical_features",
        description="ایجاد دسته‌بندی ویژگی فنی",
        created_by_user_id=current_user.id,
        type="Create",
    ))
    return RedirectResponse(url="/administration/category-technical-features", status_code=303)


@router.get("/category-technical-features/{ctf_id}/edit", response_class=HTMLResponse)
async def admin_category_technical_feature_edit(
    request: Request,
    ctf_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    ctf = (await db.execute(
        select(CategoryTechnicalFeature)
        .options(selectinload(CategoryTechnicalFeature.category), selectinload(CategoryTechnicalFeature.technical_feature))
        .where(CategoryTechnicalFeature.id == uuid.UUID(ctf_id), CategoryTechnicalFeature.is_removed == False)
    )).scalar_one_or_none()
    if ctf is None:
        raise HTTPException(status_code=404, detail="Category technical feature not found")
    categories, technical_features = await _ctf_dropdowns(db)
    return templates.TemplateResponse("admin/category_technical_feature_form.html", {
        "request": request, "current_user": current_user, "ctf": ctf,
        "categories": categories, "technical_features": technical_features,
    })


@router.post("/category-technical-features/{ctf_id}/edit", response_class=HTMLResponse)
async def admin_category_technical_feature_edit_submit(
    request: Request,
    ctf_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    ctf = (await db.execute(
        select(CategoryTechnicalFeature)
        .where(CategoryTechnicalFeature.id == uuid.UUID(ctf_id), CategoryTechnicalFeature.is_removed == False)
    )).scalar_one_or_none()
    if ctf is None:
        raise HTTPException(status_code=404, detail="Category technical feature not found")

    form = await request.form()
    category_id = form.get("category_id")
    technical_feature_id = form.get("technical_feature_id")
    if not category_id or not technical_feature_id:
        raise HTTPException(status_code=400, detail="Category and Technical Feature are required")

    ctf.category_id = uuid.UUID(category_id)
    ctf.technical_feature_id = uuid.UUID(technical_feature_id)
    await db.flush()

    db.add(Log(
        record_id=ctf.id,
        table_name="category_technical_features",
        description="ویرایش دسته‌بندی ویژگی فنی",
        created_by_user_id=current_user.id,
        type="Update",
    ))
    return RedirectResponse(url="/administration/category-technical-features", status_code=303)


@router.post("/category-technical-features/{ctf_id}/delete", response_class=HTMLResponse)
async def admin_category_technical_feature_delete(
    request: Request,
    ctf_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    ctf = (await db.execute(
        select(CategoryTechnicalFeature)
        .where(CategoryTechnicalFeature.id == uuid.UUID(ctf_id), CategoryTechnicalFeature.is_removed == False)
    )).scalar_one_or_none()
    if ctf is None:
        raise HTTPException(status_code=404, detail="Category technical feature not found")
    ctf.is_removed = True
    await db.flush()
    db.add(Log(
        record_id=ctf.id,
        table_name="category_technical_features",
        description="حذف دسته‌بندی ویژگی فنی",
        created_by_user_id=current_user.id,
        type="Delete",
    ))
    return RedirectResponse(url="/administration/category-technical-features", status_code=303)


# ── Technical Tables ──

@router.get("/technical-tables", response_class=HTMLResponse)
async def admin_technical_tables(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    items = (await db.execute(
        select(TechnicalTable).where(TechnicalTable.is_removed == False).order_by(TechnicalTable.title)
    )).scalars().all()
    return templates.TemplateResponse("admin/technical_tables.html", {
        "request": request, "current_user": current_user, "items": items,
    })


@router.get("/technical-tables/new", response_class=HTMLResponse)
async def admin_technical_table_create(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    return templates.TemplateResponse("admin/technical_table_form.html", {
        "request": request, "current_user": current_user, "table": None,
    })


@router.post("/technical-tables/new", response_class=HTMLResponse)
async def admin_technical_table_create_submit(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    form = await request.form()
    title = (form.get("title") or "").strip()
    en_title = (form.get("en_title") or "").strip()
    columns = _to_int(form.get("columns"), 0)
    header = (form.get("header") or "").strip()

    tt = TechnicalTable(
        title=title or None, en_title=en_title or None,
        columns=columns, header=header or None, created_by_user_id=current_user.id,
    )
    db.add(tt)
    await db.flush()
    db.add(Log(
        record_id=tt.id, table_name="technical_tables",
        description=f"ایجاد جدول فنی: {title}", created_by_user_id=current_user.id, type="Create",
    ))
    return RedirectResponse(url="/administration/technical-tables", status_code=303)


@router.get("/technical-tables/{table_id}/edit", response_class=HTMLResponse)
async def admin_technical_table_edit(
    request: Request,
    table_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    tt = await db.get(TechnicalTable, uuid.UUID(table_id))
    if tt is None or tt.is_removed:
        raise HTTPException(status_code=404, detail="Technical table not found")
    return templates.TemplateResponse("admin/technical_table_form.html", {
        "request": request, "current_user": current_user, "table": tt,
    })


@router.post("/technical-tables/{table_id}/edit", response_class=HTMLResponse)
async def admin_technical_table_edit_submit(
    request: Request,
    table_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    tt = await db.get(TechnicalTable, uuid.UUID(table_id))
    if tt is None or tt.is_removed:
        raise HTTPException(status_code=404, detail="Technical table not found")

    form = await request.form()
    tt.title = (form.get("title") or "").strip() or None
    tt.en_title = (form.get("en_title") or "").strip() or None
    tt.columns = _to_int(form.get("columns"), 0)
    tt.header = (form.get("header") or "").strip() or None
    await db.flush()

    db.add(Log(
        record_id=tt.id, table_name="technical_tables",
        description=f"ویرایش جدول فنی: {tt.title}", created_by_user_id=current_user.id, type="Update",
    ))
    return RedirectResponse(url="/administration/technical-tables", status_code=303)


@router.get("/technical-tables/{table_id}", response_class=HTMLResponse)
async def admin_technical_table_detail(
    request: Request,
    table_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    tt = await db.get(TechnicalTable, uuid.UUID(table_id))
    if tt is None or tt.is_removed:
        raise HTTPException(status_code=404, detail="Technical table not found")
    return templates.TemplateResponse("admin/technical_table_detail.html", {
        "request": request, "current_user": current_user, "table": tt,
    })


@router.post("/technical-tables/{table_id}/delete", response_class=HTMLResponse)
async def admin_technical_table_delete(
    request: Request,
    table_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    tt = await db.get(TechnicalTable, uuid.UUID(table_id))
    if tt is None or tt.is_removed:
        raise HTTPException(status_code=404, detail="Technical table not found")
    name = tt.title or ""
    tt.is_removed = True
    await db.flush()
    db.add(Log(
        record_id=tt.id, table_name="technical_tables",
        description=f"حذف جدول فنی: {name}", created_by_user_id=current_user.id, type="Delete",
    ))
    return RedirectResponse(url="/administration/technical-tables", status_code=303)


# ── Technical Table Products ──

@router.get("/technical-table-products", response_class=HTMLResponse)
async def admin_technical_table_products(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    items = (await db.execute(
        select(TechnicalTableProduct)
        .options(selectinload(TechnicalTableProduct.product), selectinload(TechnicalTableProduct.technical_table))
        .where(TechnicalTableProduct.is_removed == False)
        .order_by(TechnicalTableProduct.insert_date.desc())
    )).unique().scalars().all()
    return templates.TemplateResponse("admin/technical_table_products.html", {
        "request": request, "current_user": current_user, "items": items,
    })


@router.get("/technical-table-products/new", response_class=HTMLResponse)
async def admin_technical_table_product_create(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    products = (await db.execute(
        select(Product).where(Product.is_removed == False).order_by(Product.name)
    )).scalars().all()
    tables = (await db.execute(
        select(TechnicalTable).where(TechnicalTable.is_removed == False).order_by(TechnicalTable.title)
    )).scalars().all()
    return templates.TemplateResponse("admin/technical_table_product_form.html", {
        "request": request, "current_user": current_user, "ttp": None,
        "products": products, "tables": tables,
    })


@router.post("/technical-table-products/new", response_class=HTMLResponse)
async def admin_technical_table_product_create_submit(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    form = await request.form()
    product_id = form.get("product_id")
    technical_table_id = form.get("technical_table_id")
    if not product_id or not technical_table_id:
        raise HTTPException(status_code=400, detail="Product and technical table are required")

    exists = (await db.execute(
        select(TechnicalTableProduct.id).where(
            TechnicalTableProduct.is_removed == False,
            TechnicalTableProduct.product_id == uuid.UUID(product_id),
            TechnicalTableProduct.technical_table_id == uuid.UUID(technical_table_id),
        )
    )).scalar_one_or_none()
    if exists:
        raise HTTPException(status_code=400, detail="این جدول فنی برای این محصول تعریف شده")

    ttp = TechnicalTableProduct(
        product_id=uuid.UUID(product_id),
        technical_table_id=uuid.UUID(technical_table_id),
        created_by_user_id=current_user.id,
    )
    db.add(ttp)
    await db.flush()
    db.add(Log(
        record_id=ttp.id, table_name="technical_table_products",
        description="ایجاد جدول فنی محصول", created_by_user_id=current_user.id, type="Create",
    ))
    return RedirectResponse(url="/administration/technical-table-products", status_code=303)


@router.get("/technical-table-products/{ttp_id}/edit", response_class=HTMLResponse)
async def admin_technical_table_product_edit(
    request: Request,
    ttp_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    ttp = (await db.execute(
        select(TechnicalTableProduct)
        .options(selectinload(TechnicalTableProduct.product), selectinload(TechnicalTableProduct.technical_table))
        .where(TechnicalTableProduct.id == uuid.UUID(ttp_id), TechnicalTableProduct.is_removed == False)
    )).scalar_one_or_none()
    if ttp is None:
        raise HTTPException(status_code=404, detail="Technical table product not found")
    products = (await db.execute(
        select(Product).where(Product.is_removed == False).order_by(Product.name)
    )).scalars().all()
    tables = (await db.execute(
        select(TechnicalTable).where(TechnicalTable.is_removed == False).order_by(TechnicalTable.title)
    )).scalars().all()
    return templates.TemplateResponse("admin/technical_table_product_form.html", {
        "request": request, "current_user": current_user, "ttp": ttp,
        "products": products, "tables": tables,
    })


@router.post("/technical-table-products/{ttp_id}/edit", response_class=HTMLResponse)
async def admin_technical_table_product_edit_submit(
    request: Request,
    ttp_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    ttp = (await db.execute(
        select(TechnicalTableProduct)
        .where(TechnicalTableProduct.id == uuid.UUID(ttp_id), TechnicalTableProduct.is_removed == False)
    )).scalar_one_or_none()
    if ttp is None:
        raise HTTPException(status_code=404, detail="Technical table product not found")

    form = await request.form()
    product_id = form.get("product_id")
    technical_table_id = form.get("technical_table_id")
    if not product_id or not technical_table_id:
        raise HTTPException(status_code=400, detail="Product and technical table are required")

    ttp.product_id = uuid.UUID(product_id)
    ttp.technical_table_id = uuid.UUID(technical_table_id)
    await db.flush()

    db.add(Log(
        record_id=ttp.id, table_name="technical_table_products",
        description="ویرایش جدول فنی محصول", created_by_user_id=current_user.id, type="Update",
    ))
    return RedirectResponse(url="/administration/technical-table-products", status_code=303)


@router.get("/technical-table-products/{ttp_id}", response_class=HTMLResponse)
async def admin_technical_table_product_detail(
    request: Request,
    ttp_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    ttp = (await db.execute(
        select(TechnicalTableProduct)
        .options(selectinload(TechnicalTableProduct.product), selectinload(TechnicalTableProduct.technical_table))
        .where(TechnicalTableProduct.id == uuid.UUID(ttp_id), TechnicalTableProduct.is_removed == False)
    )).scalar_one_or_none()
    if ttp is None:
        raise HTTPException(status_code=404, detail="Technical table product not found")
    return templates.TemplateResponse("admin/technical_table_product_detail.html", {
        "request": request, "current_user": current_user, "ttp": ttp,
    })


@router.post("/technical-table-products/{ttp_id}/delete", response_class=HTMLResponse)
async def admin_technical_table_product_delete(
    request: Request,
    ttp_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    ttp = (await db.execute(
        select(TechnicalTableProduct)
        .where(TechnicalTableProduct.id == uuid.UUID(ttp_id), TechnicalTableProduct.is_removed == False)
    )).scalar_one_or_none()
    if ttp is None:
        raise HTTPException(status_code=404, detail="Technical table product not found")
    ttp.is_removed = True
    await db.flush()
    db.add(Log(
        record_id=ttp.id, table_name="technical_table_products",
        description="حذف جدول فنی محصول", created_by_user_id=current_user.id, type="Delete",
    ))
    return RedirectResponse(url="/administration/technical-table-products", status_code=303)


# ── Technical Features ──

@router.get("/technical-features", response_class=HTMLResponse)
async def admin_technical_features(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    items = (await db.execute(
        select(TechnicalFeature).where(TechnicalFeature.is_removed == False).order_by(TechnicalFeature.fa_name)
    )).scalars().all()
    return templates.TemplateResponse("admin/technical_features.html", {
        "request": request, "current_user": current_user, "items": items,
    })


@router.get("/technical-features/new", response_class=HTMLResponse)
async def admin_technical_feature_create(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    enums = (await db.execute(
        select(TechnicalFeatureEnum).where(TechnicalFeatureEnum.is_removed == False).order_by(TechnicalFeatureEnum.persian_name)
    )).scalars().all()
    return templates.TemplateResponse("admin/technical_feature_form.html", {
        "request": request, "current_user": current_user, "feature": None, "enums": enums,
    })


@router.post("/technical-features/new", response_class=HTMLResponse)
async def admin_technical_feature_create_submit(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    form = await request.form()
    name = (form.get("name") or "").strip()
    fa_name = (form.get("fa_name") or "").strip()
    description = (form.get("description") or "").strip()
    display_format = (form.get("display_format") or "").strip()
    linear_display = (form.get("linear_display") or "").strip()
    columns = _to_int(form.get("columns"), 1)
    priority = _to_int(form.get("priority"), 1500)
    visible_in_schema = form.get("visible_in_schema") == "on"

    feature = TechnicalFeature(
        name=name or None, fa_name=fa_name or None, description=description or None,
        display_format=display_format or None, linear_display=linear_display or None,
        columns=max(columns, 1), priority=priority, visible_in_schema=visible_in_schema,
        d_value=_checkbox_bool(form, "d_value"), unit=_checkbox_bool(form, "unit"),
        s_value=_checkbox_bool(form, "s_value"), e_value=_checkbox_bool(form, "e_value"),
        e_value1=_checkbox_bool(form, "e_value1"),
        b_value=_checkbox_bool(form, "b_value"),
        min_value=_checkbox_bool(form, "min_value"), min_unit=_checkbox_bool(form, "min_unit"),
        max_value=_checkbox_bool(form, "max_value"), max_unit=_checkbox_bool(form, "max_unit"),
        x_value=_checkbox_bool(form, "x_value"), x_unit=_checkbox_bool(form, "x_unit"),
        y_value=_checkbox_bool(form, "y_value"), y_unit=_checkbox_bool(form, "y_unit"),
        z_value=_checkbox_bool(form, "z_value"), z_unit=_checkbox_bool(form, "z_unit"),
        created_by_user_id=current_user.id,
    )
    db.add(feature)
    await db.flush()
    await _assign_feature_enums(db, feature, form.getlist("feature_enum_ids"), form.getlist("feature_enum1_ids"))
    db.add(Log(
        record_id=feature.id, table_name="technical_features",
        description=f"ایجاد ویژگی فنی: {fa_name or name}", created_by_user_id=current_user.id, type="Create",
    ))
    return RedirectResponse(url="/administration/technical-features", status_code=303)


async def _assign_feature_enums(db: AsyncSession, feature: TechnicalFeature, enum_ids, enum1_ids):
    enum_ids = enum_ids or []
    enum1_ids = enum1_ids or []
    existing = (await db.execute(
        select(TechnicalFeatureEnum).where(
            TechnicalFeatureEnum.id.in_([uuid.UUID(x) for x in enum_ids + enum1_ids]),
            TechnicalFeatureEnum.is_removed == False,
        )
    )).scalars().all()
    for e in existing:
        if e.id in {uuid.UUID(x) for x in enum_ids}:
            e.technical_feature_id = feature.id
        else:
            e.technical_feature_id = None
        if e.id in {uuid.UUID(x) for x in enum1_ids}:
            e.technical_feature1_id = feature.id
        else:
            e.technical_feature1_id = None


@router.get("/technical-features/{feature_id}/edit", response_class=HTMLResponse)
async def admin_technical_feature_edit(
    request: Request,
    feature_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    feature = await db.get(TechnicalFeature, uuid.UUID(feature_id))
    if feature is None or feature.is_removed:
        raise HTTPException(status_code=404, detail="Technical feature not found")
    enums = (await db.execute(
        select(TechnicalFeatureEnum).where(TechnicalFeatureEnum.is_removed == False).order_by(TechnicalFeatureEnum.persian_name)
    )).scalars().all()
    return templates.TemplateResponse("admin/technical_feature_form.html", {
        "request": request, "current_user": current_user, "feature": feature, "enums": enums,
    })


@router.post("/technical-features/{feature_id}/edit", response_class=HTMLResponse)
async def admin_technical_feature_edit_submit(
    request: Request,
    feature_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    feature = await db.get(TechnicalFeature, uuid.UUID(feature_id))
    if feature is None or feature.is_removed:
        raise HTTPException(status_code=404, detail="Technical feature not found")

    form = await request.form()
    feature.name = (form.get("name") or "").strip() or None
    feature.fa_name = (form.get("fa_name") or "").strip() or None
    feature.description = (form.get("description") or "").strip() or None
    feature.display_format = (form.get("display_format") or "").strip() or None
    feature.linear_display = (form.get("linear_display") or "").strip() or None
    feature.columns = max(_to_int(form.get("columns"), 1), 1)
    feature.priority = _to_int(form.get("priority"), 1500)
    feature.visible_in_schema = form.get("visible_in_schema") == "on"

    feature.d_value = _checkbox_bool(form, "d_value")
    feature.unit = _checkbox_bool(form, "unit")
    feature.s_value = _checkbox_bool(form, "s_value")
    feature.e_value = _checkbox_bool(form, "e_value")
    feature.e_value1 = _checkbox_bool(form, "e_value1")
    feature.b_value = _checkbox_bool(form, "b_value")
    feature.min_value = _checkbox_bool(form, "min_value")
    feature.min_unit = _checkbox_bool(form, "min_unit")
    feature.max_value = _checkbox_bool(form, "max_value")
    feature.max_unit = _checkbox_bool(form, "max_unit")
    feature.x_value = _checkbox_bool(form, "x_value")
    feature.x_unit = _checkbox_bool(form, "x_unit")
    feature.y_value = _checkbox_bool(form, "y_value")
    feature.y_unit = _checkbox_bool(form, "y_unit")
    feature.z_value = _checkbox_bool(form, "z_value")
    feature.z_unit = _checkbox_bool(form, "z_unit")
    await _assign_feature_enums(db, feature, form.getlist("feature_enum_ids"), form.getlist("feature_enum1_ids"))
    await db.flush()

    db.add(Log(
        record_id=feature.id, table_name="technical_features",
        description=f"ویرایش ویژگی فنی: {feature.fa_name or feature.name}", created_by_user_id=current_user.id, type="Update",
    ))
    return RedirectResponse(url="/administration/technical-features", status_code=303)


@router.get("/technical-features/{feature_id}", response_class=HTMLResponse)
async def admin_technical_feature_detail(
    request: Request,
    feature_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    feature = await db.get(TechnicalFeature, uuid.UUID(feature_id))
    if feature is None or feature.is_removed:
        raise HTTPException(status_code=404, detail="Technical feature not found")
    return templates.TemplateResponse("admin/technical_feature_detail.html", {
        "request": request, "current_user": current_user, "feature": feature,
    })


@router.post("/technical-features/{feature_id}/delete", response_class=HTMLResponse)
async def admin_technical_feature_delete(
    request: Request,
    feature_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    feature = await db.get(TechnicalFeature, uuid.UUID(feature_id))
    if feature is None or feature.is_removed:
        raise HTTPException(status_code=404, detail="Technical feature not found")
    name = feature.name or feature.fa_name or ""
    feature.is_removed = True
    await db.flush()
    db.add(Log(
        record_id=feature.id, table_name="technical_features",
        description=f"حذف ویژگی فنی: {name}", created_by_user_id=current_user.id, type="Delete",
    ))
    return RedirectResponse(url="/administration/technical-features", status_code=303)


# ── Category Options ──

@router.get("/category-options", response_class=HTMLResponse)
async def admin_category_options(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    items = (await db.execute(
        select(CategoryOption).where(CategoryOption.is_removed == False).order_by(CategoryOption.name)
    )).scalars().all()
    return templates.TemplateResponse("admin/category_options.html", {
        "request": request, "current_user": current_user, "items": items,
    })


@router.get("/category-options/new", response_class=HTMLResponse)
async def admin_category_option_create(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    return templates.TemplateResponse("admin/category_option_form.html", {
        "request": request, "current_user": current_user, "category_option": None,
    })


@router.post("/category-options/new", response_class=HTMLResponse)
async def admin_category_option_create_submit(
    request: Request,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    form = await request.form()
    name = (form.get("name") or "").strip()
    try:
        co = CategoryOption(name=name, created_by_user_id=current_user.id)
        db.add(co)
        await db.flush()
    except Exception:
        raise HTTPException(status_code=400, detail="این گزینه قبلاً تعریف شده")
    db.add(Log(
        record_id=co.id, table_name="category_options",
        description=f"ایجاد گزینه دسته‌بندی: {name}", created_by_user_id=current_user.id, type="Create",
    ))
    return RedirectResponse(url="/administration/category-options", status_code=303)


@router.get("/category-options/{co_id}/edit", response_class=HTMLResponse)
async def admin_category_option_edit(
    request: Request,
    co_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    co = await db.get(CategoryOption, uuid.UUID(co_id))
    if co is None or co.is_removed:
        raise HTTPException(status_code=404, detail="Category option not found")
    return templates.TemplateResponse("admin/category_option_form.html", {
        "request": request, "current_user": current_user, "category_option": co,
    })


@router.post("/category-options/{co_id}/edit", response_class=HTMLResponse)
async def admin_category_option_edit_submit(
    request: Request,
    co_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    co = await db.get(CategoryOption, uuid.UUID(co_id))
    if co is None or co.is_removed:
        raise HTTPException(status_code=404, detail="Category option not found")

    form = await request.form()
    name = (form.get("name") or "").strip()
    if not name:
        raise HTTPException(status_code=400, detail="نام الزامی است")
    co.name = name
    await db.flush()

    db.add(Log(
        record_id=co.id, table_name="category_options",
        description=f"ویرایش گزینه دسته‌بندی: {co.name}", created_by_user_id=current_user.id, type="Update",
    ))
    return RedirectResponse(url="/administration/category-options", status_code=303)


@router.get("/category-options/{co_id}", response_class=HTMLResponse)
async def admin_category_option_detail(
    request: Request,
    co_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    co = await db.get(CategoryOption, uuid.UUID(co_id))
    if co is None or co.is_removed:
        raise HTTPException(status_code=404, detail="Category option not found")
    return templates.TemplateResponse("admin/category_option_detail.html", {
        "request": request, "current_user": current_user, "category_option": co,
    })


@router.post("/category-options/{co_id}/delete", response_class=HTMLResponse)
async def admin_category_option_delete(
    request: Request,
    co_id: str,
    current_user: User = Depends(require_any_role("Admin", "Product Manager")),
    db: AsyncSession = Depends(get_db),
):
    co = await db.get(CategoryOption, uuid.UUID(co_id))
    if co is None or co.is_removed:
        raise HTTPException(status_code=404, detail="Category option not found")
    name = co.name or ""
    co.is_removed = True
    await db.flush()
    db.add(Log(
        record_id=co.id, table_name="category_options",
        description=f"حذف گزینه دسته‌بندی: {name}", created_by_user_id=current_user.id, type="Delete",
    ))
    return RedirectResponse(url="/administration/category-options", status_code=303)


def _to_int(value, default=0):
    try:
        return int(value)
    except (TypeError, ValueError):
        return default


def _checkbox_bool(form, key):
    return form.get(key) == "on"


# ── Role Claims ──

@router.get("/role-claims", response_class=HTMLResponse)
async def admin_role_claims(
    request: Request,
    page: int = Query(1),
    current_user: User = Depends(get_admin_user),
    db: AsyncSession = Depends(get_db),
):
    items, total = await identity_service.get_role_claims_with_relations(db, page, 50)
    return templates.TemplateResponse("admin/role_claims.html", {
        "request": request, "current_user": current_user,
        "items": items, "total": total, "page": page, "total_pages": (total + 49) // 50,
    })


@router.get("/role-claims/new", response_class=HTMLResponse)
async def admin_role_claim_create(
    request: Request,
    current_user: User = Depends(get_admin_user),
    db: AsyncSession = Depends(get_db),
):
    roles = (await db.execute(select(Role).where(Role.is_removed == False))).scalars().all()
    from app.models.enums import OperationType
    return templates.TemplateResponse("admin/role_claim_form.html", {
        "request": request, "current_user": current_user,
        "rc": None, "roles": roles, "operation_types": list(OperationType),
    })


@router.post("/role-claims/new", response_class=HTMLResponse)
async def admin_role_claim_create_submit(
    request: Request,
    current_user: User = Depends(get_admin_user),
    db: AsyncSession = Depends(get_db),
):
    form = await request.form()
    roles = (await db.execute(select(Role).where(Role.is_removed == False))).scalars().all()
    from app.models.enums import OperationType
    try:
        rc = await identity_service.create_role_claim(db, dict(form), current_user.id)
        return RedirectResponse(url="/administration/role-claims", status_code=303)
    except ValueError as e:
        return templates.TemplateResponse("admin/role_claim_form.html", {
            "request": request, "current_user": current_user,
            "rc": None, "roles": roles, "operation_types": list(OperationType),
            "error": str(e), "form_data": dict(form),
        })


@router.get("/role-claims/{claim_id}/edit", response_class=HTMLResponse)
async def admin_role_claim_edit(
    request: Request,
    claim_id: str,
    current_user: User = Depends(get_admin_user),
    db: AsyncSession = Depends(get_db),
):
    rc = await identity_service.get_role_claim_by_id(db, uuid.UUID(claim_id))
    if not rc:
        raise HTTPException(status_code=404, detail="Role claim not found")
    roles = (await db.execute(select(Role).where(Role.is_removed == False))).scalars().all()
    from app.models.enums import OperationType
    return templates.TemplateResponse("admin/role_claim_form.html", {
        "request": request, "current_user": current_user,
        "rc": rc, "roles": roles, "operation_types": list(OperationType),
    })


@router.post("/role-claims/{claim_id}/edit", response_class=HTMLResponse)
async def admin_role_claim_edit_submit(
    request: Request,
    claim_id: str,
    current_user: User = Depends(get_admin_user),
    db: AsyncSession = Depends(get_db),
):
    rc = await identity_service.get_role_claim_by_id(db, uuid.UUID(claim_id))
    if not rc:
        raise HTTPException(status_code=404, detail="Role claim not found")
    form = await request.form()
    roles = (await db.execute(select(Role).where(Role.is_removed == False))).scalars().all()
    from app.models.enums import OperationType
    try:
        await identity_service.update_role_claim(db, rc, dict(form), current_user.id)
        return RedirectResponse(url="/administration/role-claims", status_code=303)
    except ValueError as e:
        return templates.TemplateResponse("admin/role_claim_form.html", {
            "request": request, "current_user": current_user,
            "rc": rc, "roles": roles, "operation_types": list(OperationType),
            "error": str(e), "form_data": dict(form),
        })


@router.get("/role-claims/{claim_id}", response_class=HTMLResponse)
async def admin_role_claim_detail(
    request: Request,
    claim_id: str,
    current_user: User = Depends(get_admin_user),
    db: AsyncSession = Depends(get_db),
):
    rc = await identity_service.get_role_claim_by_id(db, uuid.UUID(claim_id))
    if not rc:
        raise HTTPException(status_code=404, detail="Role claim not found")
    return templates.TemplateResponse("admin/role_claim_detail.html", {
        "request": request, "current_user": current_user, "rc": rc,
    })


@router.post("/role-claims/{claim_id}/delete", response_class=HTMLResponse)
async def admin_role_claim_delete(
    request: Request,
    claim_id: str,
    current_user: User = Depends(get_admin_user),
    db: AsyncSession = Depends(get_db),
):
    rc = await identity_service.get_role_claim_by_id(db, uuid.UUID(claim_id))
    if not rc:
        raise HTTPException(status_code=404, detail="Role claim not found")
    await identity_service.soft_delete_role_claim(db, rc, current_user.id)
    return RedirectResponse(url="/administration/role-claims", status_code=303)


# ── User Roles ──

@router.get("/user-roles", response_class=HTMLResponse)
async def admin_user_roles(
    request: Request,
    page: int = Query(1),
    current_user: User = Depends(get_admin_user),
    db: AsyncSession = Depends(get_db),
):
    items, total = await identity_service.get_user_roles_with_relations(db, page, 50)
    return templates.TemplateResponse("admin/user_roles.html", {
        "request": request, "current_user": current_user,
        "items": items, "total": total, "page": page, "total_pages": (total + 49) // 50,
    })


@router.get("/user-roles/new", response_class=HTMLResponse)
async def admin_user_role_create(
    request: Request,
    current_user: User = Depends(get_admin_user),
    db: AsyncSession = Depends(get_db),
):
    users = (await db.execute(select(User).where(User.is_removed == False).order_by(User.first_name))).scalars().all()
    roles = (await db.execute(select(Role).where(Role.is_removed == False))).scalars().all()
    return templates.TemplateResponse("admin/user_role_form.html", {
        "request": request, "current_user": current_user,
        "ur": None, "users": users, "roles": roles,
    })


@router.post("/user-roles/new", response_class=HTMLResponse)
async def admin_user_role_create_submit(
    request: Request,
    current_user: User = Depends(get_admin_user),
    db: AsyncSession = Depends(get_db),
):
    form = await request.form()
    users_list = (await db.execute(select(User).where(User.is_removed == False).order_by(User.first_name))).scalars().all()
    roles = (await db.execute(select(Role).where(Role.is_removed == False))).scalars().all()
    try:
        ur = await identity_service.create_user_role(db, dict(form), current_user.id)
        return RedirectResponse(url="/administration/user-roles", status_code=303)
    except ValueError as e:
        return templates.TemplateResponse("admin/user_role_form.html", {
            "request": request, "current_user": current_user,
            "ur": None, "users": users_list, "roles": roles,
            "error": str(e), "form_data": dict(form),
        })


@router.get("/user-roles/{ur_id}/edit", response_class=HTMLResponse)
async def admin_user_role_edit(
    request: Request,
    ur_id: str,
    current_user: User = Depends(get_admin_user),
    db: AsyncSession = Depends(get_db),
):
    ur = await identity_service.get_user_role_by_id(db, uuid.UUID(ur_id))
    if not ur:
        raise HTTPException(status_code=404, detail="User role not found")
    users_list = (await db.execute(select(User).where(User.is_removed == False).order_by(User.first_name))).scalars().all()
    roles = (await db.execute(select(Role).where(Role.is_removed == False))).scalars().all()
    return templates.TemplateResponse("admin/user_role_form.html", {
        "request": request, "current_user": current_user,
        "ur": ur, "users": users_list, "roles": roles,
    })


@router.post("/user-roles/{ur_id}/edit", response_class=HTMLResponse)
async def admin_user_role_edit_submit(
    request: Request,
    ur_id: str,
    current_user: User = Depends(get_admin_user),
    db: AsyncSession = Depends(get_db),
):
    ur = await identity_service.get_user_role_by_id(db, uuid.UUID(ur_id))
    if not ur:
        raise HTTPException(status_code=404, detail="User role not found")
    form = await request.form()
    try:
        await identity_service.update_user_role(db, ur, dict(form), current_user.id)
        return RedirectResponse(url="/administration/user-roles", status_code=303)
    except ValueError as e:
        users_list = (await db.execute(select(User).where(User.is_removed == False).order_by(User.first_name))).scalars().all()
        roles = (await db.execute(select(Role).where(Role.is_removed == False))).scalars().all()
        return templates.TemplateResponse("admin/user_role_form.html", {
            "request": request, "current_user": current_user,
            "ur": ur, "users": users_list, "roles": roles,
            "error": str(e), "form_data": dict(form),
        })


@router.get("/user-roles/{ur_id}", response_class=HTMLResponse)
async def admin_user_role_detail(
    request: Request,
    ur_id: str,
    current_user: User = Depends(get_admin_user),
    db: AsyncSession = Depends(get_db),
):
    ur = await identity_service.get_user_role_by_id(db, uuid.UUID(ur_id))
    if not ur:
        raise HTTPException(status_code=404, detail="User role not found")
    return templates.TemplateResponse("admin/user_role_detail.html", {
        "request": request, "current_user": current_user, "ur": ur,
    })


@router.post("/user-roles/{ur_id}/delete", response_class=HTMLResponse)
async def admin_user_role_delete(
    request: Request,
    ur_id: str,
    current_user: User = Depends(get_admin_user),
    db: AsyncSession = Depends(get_db),
):
    ur = await identity_service.get_user_role_by_id(db, uuid.UUID(ur_id))
    if not ur:
        raise HTTPException(status_code=404, detail="User role not found")
    await identity_service.soft_delete_user_role(db, ur, current_user.id)
    return RedirectResponse(url="/administration/user-roles", status_code=303)


# ── Technical Feature Config (for dynamic modal form) ──

FIELD_CONFIG = [
    ("d_value", "D Value", "number"),
    ("unit", "Unit", "text"),
    ("s_value", "S Value", "text"),
    ("min_value", "Min Value", "number"),
    ("min_unit", "Min Unit", "text"),
    ("max_value", "Max Value", "number"),
    ("max_unit", "Max Unit", "text"),
    ("x_value", "X Value", "number"),
    ("x_unit", "X Unit", "text"),
    ("y_value", "Y Value", "number"),
    ("y_unit", "Y Unit", "text"),
    ("z_value", "Z Value", "number"),
    ("z_unit", "Z Unit", "text"),
]

@router.get("/technical-features/{feature_id}/config")
async def admin_technical_feature_config(
    feature_id: str,
    current_user: User = Depends(get_admin_user),
    db: AsyncSession = Depends(get_db),
):
    try:
        fid = uuid.UUID(feature_id)
    except ValueError:
        raise HTTPException(status_code=400, detail="Invalid feature ID")

    stmt = (
        select(TechnicalFeature)
        .options(
            selectinload(TechnicalFeature.technical_feature_enums),
            selectinload(TechnicalFeature.technical_feature_enums1),
        )
        .where(TechnicalFeature.id == fid)
    )
    result = await db.execute(stmt)
    tf = result.unique().scalar_one_or_none()
    if not tf:
        raise HTTPException(status_code=404, detail="Technical feature not found")

    fields = []
    for prop_name, label, input_type in FIELD_CONFIG:
        val = getattr(tf, prop_name, None)
        if val:
            field = {"name": prop_name, "label": label, "type": input_type}
            fields.append(field)
    # Handle EValue and EValue1 (the bool props are `e_value`, `e_value1`)
    if tf.e_value:
        fields.append({
            "name": "EValue",
            "label": "E Value",
            "type": "select",
            "options": [{"id": str(e.id), "text": e.persian_name} for e in tf.technical_feature_enums if not e.is_removed],
        })
    if tf.e_value1:
        fields.append({
            "name": "EValue1",
            "label": "E Value1",
            "type": "select",
            "options": [{"id": str(e.id), "text": e.persian_name} for e in tf.technical_feature_enums1 if not e.is_removed],
        })
    # BValue
    if tf.b_value:
        fields.append({
            "name": "BValue",
            "label": "B Value",
            "type": "select",
            "options": [{"id": "", "text": "انتخاب"}, {"id": "True", "text": "True"}, {"id": "False", "text": "False"}],
        })

    return JSONResponse({
        "fa_name": tf.fa_name,
        "fields": fields,
    })